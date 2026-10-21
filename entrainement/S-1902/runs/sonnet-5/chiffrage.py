#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » à l'aveugle — dossier S-1902 (Bibliothèque Georgette-Lepage, Brossard).

Script autonome et déterministe (stdlib + openpyxl). Lit uniquement le kit situé dans le dossier du script :
  - releve.xlsx (feuilles « Par libellé », « Par feuille »)            -> quantités sur plan
  - rapport-releve.md (tableau famille / portées)                        -> portées (INSTALLER, ENLEVER, ...)
  - entree/reference-Estimation_2.xlsx (Détail, Matériel Lumen)          -> quantités hors plan, heures « à valider »
  - catalogue-prix-temps-dupuis.csv (colonnes prix_a_utiliser / temps_a_utiliser = dernier Dupuis)
  - METHODE.md (taux, facteur, lots fournisseur) : valeurs reprises en constantes ci-dessous, avec contrôle de présence
Écrit chiffrage.json. Toute incohérence lève une exception (aucun repli silencieux).
"""
import csv
import json
import math
import os
import re
import sys
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
P_RELEVE = os.path.join(ICI, "releve.xlsx")
P_RAPPORT = os.path.join(ICI, "rapport-releve.md")
P_CAT = os.path.join(ICI, "catalogue-prix-temps-dupuis.csv")
P_EST = os.path.join(ICI, "entree", "reference-Estimation_2.xlsx")
P_METHODE = os.path.join(ICI, "METHODE.md")
P_OUT = os.path.join(ICI, "chiffrage.json")

# --- Paramètres de METHODE.md ------------------------------------------------------------------
TAUX = 85.0            # $/h, dossiers août-octobre 2026 (S-1797, S-1835, S-1854, S-1881)
FACTEUR = 1.244        # unique couple observé (S-1844)
LOTS = {               # prix de lot confirmés par Francis le 2026-10-20 (METHODE.md §2)
    "eclairage": 11893.0,
    "distribution": 6200.0,   # NF 225 A 42 cct 3 400 + EDB 1P 95 x20 (1 900) + EJB 3P 300 x3 (900)
    "JCI": 7400.0,
    "detecteurs JCI": 600.0,
    "GPR": 1250.0,
}
LOT_DIST_PANNEAU = 3400.0
LOT_DIST_DISJ = 1900.0 + 900.0
M_PI = 3.28084  # pieds par mètre


def pi(m):
    """mètres -> pieds, arrondi vers le haut."""
    return int(math.ceil(round(m * M_PI, 6)))


# --- Contrôle que les paramètres codés correspondent bien à METHODE.md -------------------------
def verifier_methode():
    txt = open(P_METHODE, encoding="utf-8").read()
    attendus = ["11 893", "6 200", "7 400", "600 $", "1 250", "85 $/h", "1,244", "3 400"]
    manquants = [a for a in attendus if a not in txt]
    if manquants:
        raise RuntimeError("METHODE.md ne contient pas : %s" % manquants)
    if abs(LOT_DIST_PANNEAU + LOT_DIST_DISJ - LOTS["distribution"]) > 0.005:
        raise RuntimeError("lot distribution incohérent")


# --- Lecture du kit ----------------------------------------------------------------------------
def lire_catalogue():
    with open(P_CAT, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    idx = {}
    for r in rows:
        idx.setdefault(r["description"], []).append(r)
    return idx


def lire_rapport():
    """famille -> (total, {portee: n})"""
    out = {}
    for ligne in open(P_RAPPORT, encoding="utf-8"):
        m = re.match(r"^\|\s*([A-Z0-9_ ]+?)\s*\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*$", ligne)
        if not m or m.group(1) == "famille":
            continue
        portees = {}
        for part in m.group(3).split(","):
            mm = re.match(r"\s*([A-ZÉ]+)\s+(\d+)\s*$", part)
            if not mm:
                raise RuntimeError("portée illisible : %r" % ligne)
            portees[mm.group(1)] = int(mm.group(2))
        tot = int(m.group(2))
        if sum(portees.values()) != tot:
            raise RuntimeError("portées != total pour %s" % m.group(1))
        out[m.group(1).strip()] = (tot, portees)
    if len(out) != 31:
        raise RuntimeError("31 familles attendues dans rapport-releve.md, %d lues" % len(out))
    return out


def lire_releve():
    wb = openpyxl.load_workbook(P_RELEVE, data_only=True)
    par_lib = {}
    for r in wb["Par libellé"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            par_lib[r[0]] = int(r[2])
    par_feuille = {}
    for r in wb["Par feuille"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            par_feuille[(r[0], r[2])] = int(r[3])
    return par_lib, par_feuille


def lire_detail():
    wb = openpyxl.load_workbook(P_EST, data_only=False)
    ws = wb["Détail"]
    det = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        code = r[0]
        if code and re.match(r"^[A-H]-\d\d$", str(code)):
            det[code] = {"desc": r[1], "q": r[3], "u": r[4], "h": r[7]}
    lumen = {}
    for r in wb["Matériel Lumen"].iter_rows(min_row=6, values_only=True):
        if r[0] and r[2] is not None and r[4] is not None:
            lumen[r[0]] = {"desc": r[1], "q": r[2], "u": r[3], "p": r[4]}
    return det, lumen


# --- Construction des lignes -------------------------------------------------------------------
def main():
    verifier_methode()
    cat = lire_catalogue()
    rap = lire_rapport()
    par_lib, par_feuille = lire_releve()
    det, lumen = lire_detail()

    for fam, (tot, _) in rap.items():
        if par_lib.get(fam) != tot:
            raise RuntimeError("releve.xlsx (%s=%s) != rapport-releve.md (%s)" % (fam, par_lib.get(fam), tot))

    def plan(fam, portee):
        return rap[fam][1].get(portee, 0)

    def feuille(fam, sh):
        return par_feuille.get((sh, fam), 0)

    # Contrôle : les familles « nouveau travail » s'additionnent bien sur E104 + E105 (= INSTALLER)
    for fam in ["PRISE DOUBLE 15A", "PRISE DDFT", "SORTIE DATA", "PRISE USB", "MONUMENT PLANCHER",
                "LECTEUR CARTE", "BOITE JONCTION", "THERMOSTAT"]:
        s = feuille(fam, "E104") + feuille(fam, "E105")
        if s != plan(fam, "INSTALLER"):
            raise RuntimeError("E104+E105 != INSTALLER pour %s (%s vs %s)" % (fam, s, plan(fam, "INSTALLER")))

    def dq(code):
        v = det[code]["q"]
        if v is None:
            raise RuntimeError("Détail %s sans quantité" % code)
        return v

    def dh(code):
        v = det[code]["h"]
        if v is None:
            raise RuntimeError("Détail %s sans heures" % code)
        return v

    lignes = []

    def cat_row(label):
        rs = cat.get(label)
        if not rs or len(rs) != 1:
            raise RuntimeError("libellé catalogue absent ou ambigu : %r" % label)
        r = rs[0]
        return r, float(r["prix_a_utiliser"]), float(r["temps_a_utiliser"])

    def L(code, label, q, art, desc, p=None, note="", u=None):
        """Ligne issue du catalogue Dupuis (PRIX À UTILISER / TEMPS À UTILISER, dernier Dupuis)."""
        if q <= 0:
            raise RuntimeError("quantité nulle pour %s" % code)
        r, p0, t0 = cat_row(label)
        src = "Catalogue Dupuis « %s » (dernier Dupuis %s, %s; prix %.4f $, temps %.4f h)" % (
            label, r["date_dernier"], r["projet_dernier"].strip(), p0, t0)
        if p is not None and abs(p - p0) > 1e-9:
            src += " — prix de ligne forcé à %.2f $ : %s" % (p, note)
        elif note:
            src += " — " + note
        lignes.append({"code": code, "desc": desc, "q": q, "u": u or ("pi" if r["um"] == "P" else "u"),
                       "p": p0 if p is None else p, "h": t0, "art": art, "src": src})

    def J(code, desc, q, u, h, art, note, p=0.0, src_extra=""):
        """Ligne de jugement (heures de la feuille Détail d'Estimation_2, à valider)."""
        lignes.append({"code": code, "desc": desc, "q": q, "u": u, "p": p, "h": h, "art": art,
                       "src": "jugement (Estimation_2) — %s%s" % (note, src_extra)})

    def LUM(code, lcode, desc, q, art, h, note):
        """Famille absente du catalogue Dupuis : prix Lumen (feuille Matériel Lumen) + temps par jugement."""
        lu = lumen[lcode]
        lignes.append({"code": code, "desc": desc, "q": q, "u": "u", "p": float(lu["p"]), "h": h, "art": art,
                       "src": "prix Lumen %s « %s » %.4f $ (feuille Matériel Lumen, 2026-10-07, régulier) + temps par "
                              "jugement %.2f h — famille absente du catalogue Dupuis; %s" % (lcode, lu["desc"], lu["p"], h, note)})

    # ============ A. DÉMANTÈLEMENT / RELOCALISATION =============================================
    n_lum_enl = plan("ENCASTRE", "ENLEVER") + plan("LED 1X4", "ENLEVER") + plan("LED 2X2", "ENLEVER")
    L("A-01", "EQUIPEMENT A ENLEVER", n_lum_enl, "5.19",
      "Démantèlement luminaires existants RDC + 2e (E.A.E : encastrés %d, LED 1x4 %d, LED 2x2 %d) — absorbe A-02" % (
          plan("ENCASTRE", "ENLEVER"), plan("LED 1X4", "ENLEVER"), plan("LED 2X2", "ENLEVER")),
      note="luminaire enlevé 0,5 h (METHODE.md)")
    n_prises_enl = plan("PRISE DOUBLE 15A", "ENLEVER") + plan("PRISE DDFT", "ENLEVER") + plan("SORTIE DATA", "ENLEVER")
    L("A-03", "ENLEVER PRISE", n_prises_enl, "5.19",
      "Démantèlement services RDC + 2e (prises %d, DDFT %d, sorties data %d) E.A.E — absorbe A-04" % (
          plan("PRISE DOUBLE 15A", "ENLEVER"), plan("PRISE DDFT", "ENLEVER"), plan("SORTIE DATA", "ENLEVER")),
      note="prise enlevée 0,35 h (METHODE.md)")
    L("A-03a", "ENLEVER INTERRUPTEUR", plan("INT SIMPLE", "ENLEVER"), "5.19",
      "Démantèlement interrupteur existant (E.A.E)")
    n_div_enl = (plan("MONUMENT PLANCHER", "ENLEVER") + plan("BOITE JONCTION", "ENLEVER") + plan("BOITIER C", "ENLEVER")
                 + plan("DETECTEUR FUMEE", "ENLEVER") + plan("DETECTEUR MOUVEMENT", "ENLEVER") + plan("HP ENCASTRE PLAFOND", "ENLEVER"))
    L("A-03b", "EQUIPEMENT A ENLEVER", n_div_enl, "5.19",
      "Démantèlement monuments plancher %d, boîtes de jonction %d, boîtiers C %d, détecteur fumée %d, détecteur mouvement %d, "
      "haut-parleur encastré %d (E.A.E)" % (plan("MONUMENT PLANCHER", "ENLEVER"), plan("BOITE JONCTION", "ENLEVER"),
                                            plan("BOITIER C", "ENLEVER"), plan("DETECTEUR FUMEE", "ENLEVER"),
                                            plan("DETECTEUR MOUVEMENT", "ENLEVER"), plan("HP ENCASTRE PLAFOND", "ENLEVER")),
      note="absorbe E-03 (boîtes de plancher 2e à refaire)")
    n_reloc_serv = plan("PRISE DOUBLE 15A", "RELOCALISER") + plan("PRISE DDFT", "RELOCALISER") + plan("SORTIE DATA", "RELOCALISER")
    L("A-05", "RELO PRISE", n_reloc_serv, "3.22",
      "Relocalisation équipements services (E.A.R) : prises %d, DDFT %d, sorties data %d" % (
          plan("PRISE DOUBLE 15A", "RELOCALISER"), plan("PRISE DDFT", "RELOCALISER"), plan("SORTIE DATA", "RELOCALISER")),
      note="relocalisation de prise 1,5 h (METHODE.md)")
    L("A-06", "RELO ENCASTRE", plan("ENCASTRE", "RELOCALISER"), "3.22",
      "Relocalisation luminaires encastrés (E.A.R) RDC + 2e")
    q_a07 = int(round(dq("A-07") * dh("A-07") / cat_row("RELO INT")[2]))
    L("A-07", "RELO INT", q_a07, "3.22",
      "Relocalisation interrupteurs maîtres du hall (note spécifique 1 E-201)",
      note="Détail A-07 = 6 h d'hypothèse = %d interrupteurs x 1,5 h" % q_a07)
    L("A-08", "RELO ENCASTRE", dq("A-08"), "3.22",
      "FabLab 210 : retrait temporaire + réinstallation luminaires linéaires existants",
      note="quantité Détail A-08 (comptée sur plan E-202 note 1); pas de libellé « linéaire » : RELO ENCASTRE par analogie")
    L("A-09", "RELO ENCASTRE", dq("A-09"), "3.22",
      "FabLab 210 : retrait temporaire + réinstallation encastrés existants", note="quantité Détail A-09")
    L("A-10", "DISCONNECTER ET RECON", plan("MOTEUR MONOPHASE", "RELOCALISER"), "3.23",
      "Toit : débrancher, entreposer, rebrancher équipements de ventilation (E.A.R, moteurs monophasés)",
      note="Détail A-10 = 4 h par moteur; catalogue 8 h retenu")
    J("A-11", det["A-11"]["desc"], dq("A-11"), "lot", dh("A-11"), "3.22", "Détail A-11, %s h à valider (hypothèse)" % dh("A-11"))

    # ============ B. DISTRIBUTION / PANNEAUX ====================================================
    L("B-01", "PANN PD 400A 347/600V 42CCT 35KA C/A", dq("B-01"), "3.24",
      "Panneau PP-3 NEUF 225A 347/600V 3Ø 4F 42 cct (NF Schneider) — pose, temps seulement", p=0.0,
      note="panneau 8 h (METHODE.md); matériel dans le lot distribution (NF 225 A 42 cct 3 400 $)")
    L("B-01b", "ENL PANN 42CCT 120/208V", 1, "3.24", "Démantèlement de l'ancien PP-3 (42 cct) + transfert des charges")
    L("B-02", "DISJ15A1P", dq("B-02"), "3.25", "Disjoncteurs NF 1P-15A pour circuits existants transférés dans PP-3",
      note="disjoncteur 1P 0,2 h (METHODE.md); matériel dans le lot distribution (EDB 1P)")
    for code in ["B-03", "B-04", "B-05"]:
        L(code, "DISJ 100A 3P 120/240V 22 K", dq(code), "3.25", det[code]["desc"], p=0.0,
          note="disjoncteur 3P 1 h (METHODE.md); seul le temps est pris : matériel dans le lot distribution (EJB 3P)")
    L("B-06", "DISJ15A1P", dq("B-06"), "3.25", det["B-06"]["desc"], note="matériel dans le lot distribution")
    L("B-07", "DISJ15A1P", dq("B-07"), "3.25", det["B-07"]["desc"], note="matériel dans le lot distribution")
    L("B-08", "DISJ15A1P", dq("B-08"), "3.25", det["B-08"]["desc"],
      note="20 A 1P : même temps que 15 A 1P; matériel dans le lot distribution")
    J("B-09", det["B-09"]["desc"], dq("B-09"), "u", dh("B-09"), "3.24", "Détail B-09, %s h/u à valider" % dh("B-09"))
    lignes.append({"code": "LOT-B1", "desc": "PRIX DE LOT distribution — panneau NF 225 A 42 cct", "q": 1, "u": "lot",
                   "p": LOT_DIST_PANNEAU, "h": 0.0, "art": "3.24",
                   "src": "lot fournisseur confirmé par Francis le 2026-10-20 (METHODE.md §2) — lot distribution 6 200 $ : part panneau"})
    lignes.append({"code": "LOT-B2", "desc": "PRIX DE LOT distribution — disjoncteurs EDB 1P x20 + EJB 3P x3", "q": 1, "u": "lot",
                   "p": LOT_DIST_DISJ, "h": 0.0, "art": "3.25",
                   "src": "lot fournisseur confirmé par Francis le 2026-10-20 (METHODE.md §2) — lot distribution 6 200 $ : part disjoncteurs (95 x20 + 300 x3)"})

    # ============ C. RACCORDEMENTS MÉCANIQUES ===================================================
    n_tri_inst = plan("MOTEUR TRIPHASE", "INSTALLER")
    n_mono_inst = plan("MOTEUR MONOPHASE", "INSTALLER")
    if n_tri_inst != 4 or n_mono_inst != 3 or dq("C-01") + dq("C-02") + dq("C-03") != n_tri_inst or dq("C-04") != n_mono_inst:
        raise RuntimeError("moteurs releve != Détail (UTA+VE+2 CH = 4 triphasés, 3 HUM monophasés)")
    L("C-01", "UNITER UC AU TOIT", dq("C-01"), "3.23", "UTA-01 York 575V/3Ø au toit — raccordement",
      note="UTA 1 h (METHODE.md)")
    L("C-02", "MOTEUR", dq("C-02"), "3.23", "VE-02N Plymovent 575V/3Ø c/a VFD au toit — raccordement",
      note="moteur 0,75 h (METHODE.md)")
    L("C-03", "MOTEUR", dq("C-03"), "3.26",
      "Chaudières CH-01 / CH-02 575-600V/3Ø — tiroirs CCM-1, raccord final", note="moteur triphasé; MCC = art. 3.26")
    L("C-04", "MOTEUR", dq("C-04"), "3.23", "Humidificateurs HUM-1/2/3 120V/1Ø — raccordement")
    L("C-04b", "SECTIONNEUR POR ENTRETIEN HUBHBL13R22", dq("C-04"), "3.23",
      "Sectionneur 30A sans fusible par humidificateur (Hubbell)",
      note="équivalent du HBL1372 de la feuille Lumen (117,84 $); prix Dupuis retenu")
    ft_uta = pi(dq("C-05"))
    L("C-05", "CONDUIT 1 04C08", ft_uta, "3.23",
      "Conduit EMT 1\" + 4 #8 (alim. UTA-01), %s m — absorbe C-06" % dq("C-05"),
      note="conduit + conducteurs en assemblage; longueur hypothèse Détail C-05")
    ft_ve = pi(dq("C-05"))  # VE-02N : 35 m, même longueur que l'UTA (Détail C-02 : « longueur estimée 35 m »)
    m_hum = dq("C-07") - dq("C-05")
    if abs(m_hum - 3 * 15) > 1e-9:
        raise RuntimeError("C-07 - C-05 != 3 x 15 m")
    L("C-07", "CONDUIT 3/4 04C12", ft_ve, "3.23", "Conduit EMT 3/4\" + 4 #12 (VE-02N), 35 m — absorbe une partie de C-08",
      note="longueur hypothèse Détail C-02 / C-07")
    L("C-07b", "CONDUIT 3/4 03C12", pi(m_hum), "3.23", "Conduit EMT 3/4\" + 3 #12 (HUM-1/2/3), 3 x 15 m — absorbe le reste de C-08",
      note="longueur hypothèse Détail C-04 / C-07")
    L("C-09", "COND FLEX ETANCHE 3/4PO ACIER BOBINE STD 150M", pi(lumen["CONCONSEAL3/4"]["q"]), "3.23",
      "Flexible étanche (LFMC) 3/4\" — raccords finaux équipements mécaniques",
      note="quantité = feuille Matériel Lumen (%s m)" % lumen["CONCONSEAL3/4"]["q"])
    L("C-09b", "CONN CONSEAL 3/4PO 90D ALLIAGE ZINC", int(lumen["THO5333"]["q"]), "3.23",
      "Connecteurs étanches 3/4\" pour LFMC", note="quantité = feuille Matériel Lumen (%s)" % lumen["THO5333"]["q"])
    L("C-10", "12AWG 3C ARME ALUM 600'", pi(lumen["CAB12TECK3C600V"]["q"]), "3.26",
      "Câble armé 3C #12 600V (chaudières), %s m" % lumen["CAB12TECK3C600V"]["q"],
      note="équivalent de TECK90 3C #12 (feuille Lumen); quantité Lumen")
    LUM("C-10b", "IBETEK75120", "Connecteurs Star Teck 3/4\" (chaudières)", int(lumen["IBETEK75120"]["q"]), "3.26", 0.25,
        "quantité = feuille Matériel Lumen")

    # ============ D. SERVICES RDC (E-104) =======================================================
    L("D-01", "PRISE DUPLEX 15A 120V", feuille("PRISE DOUBLE 15A", "E104"), "3.28",
      "Prises doubles 15A neuves RDC (E-104) — absorbe les composites D-01, D-03, D-04, D-05, D-06",
      note="prise 1,175 h, 64,81 $ (METHODE.md)")
    L("D-02", "PRISE USB 125V 15 AMP", feuille("PRISE USB", "E104"), "3.27", "Prises doubles USB Leviton Decora neuves RDC (E-104)",
      note="USB 1,125 h, 77,45 $ (METHODE.md)")
    L("D-07", "PRISE DUPLEX 20A 125V GFI", feuille("PRISE DDFT", "E104"), "3.30",
      "Prises DDFT neuves RDC (E-104) — symbole DDFT/MALT isolée non départagé",
      note="DDFT 1,3 h, 65,38 $ (METHODE.md)")
    L("D-08", "TEL/DATA 3/4", feuille("SORTIE DATA", "E104"), "3.33", "Sorties data neuves RDC (E-104), boîte + conduit 3/4\"",
      note="tel/data 3/4 22,32 $ (METHODE.md); câble data hors prix (OPT-1)")
    J("D-09", det["D-09"]["desc"], dq("D-09"), "lot", dh("D-09"), "3.31", "Détail D-09, %s h à valider" % dh("D-09"))
    q_jb = feuille("BOITE JONCTION", "E104")
    L("D-10", "BOITE CARRE 2-1/8PO KO 1", q_jb, "3.31", "Boîtes de jonction neuves RDC (E-104) — boîte carrée 4\"")
    L("D-10b", "COUV CARRE 4PO VIERGE", q_jb, "3.31", "Couvercles carrés 4\" pour boîtes de jonction RDC")
    L("D-11", "BOUTON HANDICAPE", dq("D-11"), "3.31", "Bouton d'appel — boîte (système fourni par d'autres)")
    L("D-12", "PERCEMENT PLANCHER", dq("D-12"), "3.31", "Percement de dalle carotté — conduits sous dalle RDC (note 5 E-104)",
      note="percement 150 $ / 1 h (METHODE.md); quantité Détail D-12")
    L("D-12b", "COUPE FEU TUBE 10.1 FL O", dq("D-12"), "3.31", "Coupe-feu aux percements de dalle RDC (CAN/ULC-S115)")
    L("D-13", "CONDUIT 3/4 05C12", pi(dq("D-13")), "3.31",
      "Conduit + 5 #12 sous dalle RDC (16 courses), %s m — absorbe D-15" % dq("D-13"),
      note="Détail prévoit EMT 1\" : 3/4\" 05C12 retenu (même nombre de conducteurs); longueur hypothèse Détail")
    L("D-14", "CONDUIT VIDE 3/4", pi(dq("D-14")), "3.31", "Conduit vide 3/4\" data sous dalle RDC (11 courses), %s m" % dq("D-14"),
      note="conduit vide 1,41 $/pi (METHODE.md); longueur hypothèse Détail")
    L("D-16", "BOITE CARRE 2-1/8PO KO 1", 2 * dq("D-12"), "3.31", "Boîtes de tirage carrées 4\" — 2 par course sous dalle RDC",
      note="2 par course (feuille Matériel Lumen : « 23 courses x 2 »)")
    L("D-16b", "COUV CARRE 4PO VIERGE", 2 * dq("D-12"), "3.31", "Couvercles de boîtes de tirage RDC")
    L("N-02", "THERMOSTAT", plan("THERMOSTAT", "INSTALLER"), "3.23", "Thermostats neufs RDC (E-104)",
      note="thermostat 35 $ (METHODE.md)")
    L("N-03", "LECTEUR DE CARTE", plan("LECTEUR CARTE", "INSTALLER"), "3.31", "Lecteurs de carte neufs RDC (E-104)",
      note="lecteur de carte 15 $ (METHODE.md); « LECTEUR CARTE » existe aussi (0,35 h)")
    L("N-01", "SCAN DALLE BETON", 1, "3.32", "Scan de dalle béton (GPR) avant percements",
      note="scan dalle 1 250 $ (METHODE.md) = lot GPR 1 250 $ confirmé 2026-10-20 : une seule fois")

    # ============ E. SERVICES ÉTAGE (E-105) =====================================================
    L("E-01", "PRISE DUPLEX 15A 120V", feuille("PRISE DOUBLE 15A", "E105"), "3.28",
      "Prises doubles 15A neuves étage (E-105) — absorbe E-01, E-04, E-05 composites")
    L("E-01b", "PRISE DUPLEX 20A 125V GFI", feuille("PRISE DDFT", "E105"), "3.30", "Prises DDFT neuves étage (E-105)")
    L("E-01c", "TEL/DATA 3/4", feuille("SORTIE DATA", "E105"), "3.33", "Sorties data neuves étage (E-105)")
    L("E-02", "MONUMENT PLANCHER", feuille("MONUMENT PLANCHER", "E105"), "3.34",
      "Monuments de plancher neufs étage (E-105) — absorbe E-03 refaites", note="monument 250 $ (METHODE.md)")
    L("E-06", "BOUTON HANDICAPE", dq("E-06"), "3.31", "Boutons d'appel (2) + panique (1) — boîte")
    L("E-07", "PERCEMENT PLANCHER", dq("E-07"), "3.31", "Percement de dalle carotté — conduits sous dalle étage (note 3 E-105)")
    L("E-07b", "COUPE FEU TUBE 10.1 FL O", dq("E-07"), "3.31", "Coupe-feu aux percements de dalle étage")
    L("E-08", "CONDUIT 3/4 05C12", pi(dq("E-08")), "3.31",
      "Conduit + 5 #12 sous dalle étage (7 courses), %s m — absorbe E-10" % dq("E-08"),
      note="Détail prévoit EMT 1\" : 3/4\" 05C12 retenu; longueur hypothèse Détail")
    L("E-09", "CONDUIT VIDE 3/4", pi(dq("E-09")), "3.31", "Conduit vide 3/4\" data sous dalle étage (7 courses), %s m" % dq("E-09"))
    L("E-11", "BOITE CARRE 2-1/8PO KO 1", 2 * dq("E-07"), "3.31", "Boîtes de tirage carrées 4\" étage — 2 par course")
    L("E-11b", "COUV CARRE 4PO VIERGE", 2 * dq("E-07"), "3.31", "Couvercles de boîtes de tirage étage")

    # ============ F. ÉCLAIRAGE (E-203) — 0 $ en ligne + lot =====================================
    n = {k: par_lib[k] for k in ["LUMINAIRE A1", "LUMINAIRE C1", "LUMINAIRE C1_1", "LUMINAIRE C1_3", "LUMINAIRE E", "ENSEIGNE SORTIE Z6"]}
    L("F-01", "FIXTURE TYPE A1", n["LUMINAIRE A1"], "3.35", "Luminaire A1 Cooper 2'x4' encastré 347V — pose seulement", p=0.0,
      note="luminaire 0,75 h (METHODE.md); matériel dans le lot éclairage")
    L("F-02", "FIXTURE TYPE C1", n["LUMINAIRE C1"], "3.35", "Luminaire C1 Axis SL1 module 4' — pose seulement", p=0.0,
      note="matériel dans le lot éclairage")
    L("F-03", "FIXTURE TYPE C1", n["LUMINAIRE C1_1"], "3.35", "Luminaire C1_1 Axis SL1 module 1' — pose seulement", p=0.0,
      note="pas de libellé C1_1 : FIXTURE TYPE C1 (0,75 h; Détail 0,5 h); matériel dans le lot éclairage")
    L("F-04", "FIXTURE TYPE C1", n["LUMINAIRE C1_3"], "3.35", "Luminaire C1_3 Axis SL1 module 3' — pose seulement", p=0.0,
      note="pas de libellé C1_3 : FIXTURE TYPE C1 (0,75 h; Détail 0,6 h); matériel dans le lot éclairage")
    J("F-05", det["F-05"]["desc"], dq("F-05"), "rangée", dh("F-05"), "3.35", "Détail F-05, %s h/rangée à valider" % dh("F-05"),
      src_extra="; trousses de suspension dans le lot éclairage")
    L("F-06", "FIXTURE TYPE E", n["LUMINAIRE E"], "3.35", "Luminaire E GreenCreative cylindre suspendu — pose seulement", p=0.0,
      note="3 étiquettes au relevé (Détail : 4; lot : 3); matériel dans le lot éclairage")
    L("F-07", "ENSEIGNE SORTIE Z6", n["ENSEIGNE SORTIE Z6"], "3.35", "Enseigne de sortie Z6 Stanpro — pose seulement", p=0.0,
      note="enseigne 1 h (METHODE.md); matériel dans le lot éclairage")
    L("F-08", "CONDUIT 3/4 03C12", pi(dq("F-08")), "3.35",
      "Conduit EMT 3/4\" + 3 #12 — raccordement luminaires sur circuits existants, %s m — absorbe F-09" % dq("F-08"),
      note="conduit 3/4 03C12 4,24 $/pi, 0,113 h/pi (METHODE.md); longueur hypothèse Détail (8 m x 35 points)")
    L("F-10", "BOITE OCTOG 1-1/2PO KO 1", 35, "3.35", "Boîtes de jonction octogonales éclairage (35 points)",
      note="35 points = base de la feuille Matériel Lumen (IBE54151L)")
    L("F-10b", "QUINCAILLERIE", 1, "3.35", "Quincaillerie / chaînes de retenue à la dalle (devis 16.5)",
      note="forfait catalogue (classe Jugement / forfait)")
    L("N-04", "INTERRUPTEUR 15A 120V", plan("INT SIMPLE", "INSTALLER"), "3.35", "Interrupteurs simples neufs RDC (E-203)",
      note="interrupteur 75,47 $ (METHODE.md)")
    lignes.append({"code": "LOT-F", "desc": "PRIX DE LOT éclairage (A1 x11, SL1 4' x12, 3' x15, 1' x6, trousses, E x3, Z6 x3)",
                   "q": 1, "u": "lot", "p": LOTS["eclairage"], "h": 0.0, "art": "3.35",
                   "src": "lot fournisseur confirmé par Francis le 2026-10-20 (METHODE.md §2)"})

    # ============ G. ALARME INCENDIE ============================================================
    L("G-01", "DETECTEUR FUMEE", plan("DETECTEUR FUMEE", "INSTALLER"), "3.36",
      "Détecteurs de fumée neufs RDC (JCI) — pose seulement", p=0.0,
      note="détecteur 0,5 h (METHODE.md); matériel dans le lot « 2 détecteurs JCI 600 $ »")
    L("G-01b", "RELO DETECTEUR A FUME", plan("DETECTEUR FUMEE", "RELOCALISER"), "3.36", "Relocalisation détecteur de fumée existant")
    L("G-02", "CONDUIT VIDE 3/4", pi(dq("G-02")), "3.36", "Conduit 3/4\" boucle AI, %s m" % dq("G-02"),
      note="longueur hypothèse Détail G-02")
    L("G-03", "COND ALARME 3/16 FT4", pi(dq("G-03")), "3.36", "Câble AI FT4 2c #16 (boucle adressable), %s m" % dq("G-03"),
      note="longueur hypothèse Détail G-03")
    L("G-04", "RELO PANN ANNOCIATEUR", dq("G-04"), "3.37", "Relocalisation panneau annonciateur Simplex",
      note="annonciateur relocalisé 8 h (METHODE.md)")
    J("G-05", det["G-05"]["desc"], dq("G-05"), "lot", dh("G-05"), "3.36", "Détail G-05, %s h à valider (allocation)" % dh("G-05"))
    J("G-06", det["G-06"]["desc"], dq("G-06"), "lot", dh("G-06"), "3.38", "Détail G-06, %s h à valider" % dh("G-06"))
    lignes.append({"code": "G-07", "desc": "Sous-traitance JCI/Simplex : bypass, programmation 4007ES, vérification S-537, certificat",
                   "q": 1, "u": "lot", "p": LOTS["JCI"], "h": 0.0, "art": "3.38",
                   "src": "lot fournisseur JCI confirmé par Francis le 2026-10-20 (METHODE.md §2)"})
    lignes.append({"code": "LOT-DET", "desc": "Lot JCI : 2 détecteurs de fumée adressables + bases", "q": 1, "u": "lot",
                   "p": LOTS["detecteurs JCI"], "h": 0.0, "art": "3.36",
                   "src": "lot fournisseur confirmé par Francis le 2026-10-20 (METHODE.md §2)"})

    # ============ H. FRAIS DE CHANTIER ==========================================================
    L("H-01", "NACELLE", dq("H-01"), "6.1", "Mobilisation / démobilisation, plateforme (nacelle)",
      note="nacelle 1 250 $ / 16 h (METHODE.md) = Détail H-01 16 h")
    J("H-02", det["H-02"]["desc"], dq("H-02"), "lot", dh("H-02"), "6.1", "Détail H-02, %s h à valider" % dh("H-02"))
    J("H-03", det["H-03"]["desc"], dq("H-03"), "lot", dh("H-03"), "6.1", "Détail H-03, %s h à valider" % dh("H-03"))
    J("H-04", det["H-04"]["desc"], dq("H-04"), "lot", dh("H-04"), "6.1", "Détail H-04, %s h à valider" % dh("H-04"))

    # ============ Familles absentes du catalogue Dupuis (Lumen + jugement) ======================
    LUM("N-05", "GRE430", "Corde de tirage Poly Line 430 (conduits vides data)", int(lumen["GRE430"]["q"]), "3.31", 0.5,
        "quantité = feuille Matériel Lumen (devis 35.10)")

    # ---- Validation des lignes ----
    codes = [l["code"] for l in lignes]
    if len(codes) != len(set(codes)):
        raise RuntimeError("codes de ligne dupliqués")
    for l in lignes:
        for k in ("code", "desc", "q", "u", "p", "h", "art", "src"):
            if l.get(k) in (None, ""):
                raise RuntimeError("champ %s vide pour %s" % (k, l["code"]))

    # ---- Totaux ----
    materiel = round(sum(l["q"] * l["p"] for l in lignes), 2)
    h_rel = round(sum(l["q"] * l["h"] for l in lignes), 4)
    h_ch = int(math.ceil(round(h_rel / 10.0, 9))) * 10
    mo = round(h_ch * TAUX, 2)
    total = round(materiel + mo, 2)
    pv = int((Decimal(str(total)) * Decimal(str(FACTEUR)) / Decimal(1000)).quantize(Decimal(1), rounding=ROUND_HALF_UP)) * 1000

    # Lots : montant par lot (les lignes de lot sont dans `lignes`; GPR = ligne N-01 au prix catalogue 1 250 $)
    lots_json = dict(LOTS)
    somme_lots_lignes = sum(l["p"] for l in lignes if l["code"] in ("LOT-F", "LOT-B1", "LOT-B2", "G-07", "LOT-DET", "N-01"))
    if abs(somme_lots_lignes - sum(LOTS.values())) > 0.005:
        raise RuntimeError("lots dans les lignes != lots confirmés")

    # ---- Ventilation par article du bordereau ----
    arts = {}
    for l in lignes:
        a = arts.setdefault(l["art"], {"materiel": 0.0, "heures": 0.0})
        a["materiel"] += l["q"] * l["p"]
        a["heures"] += l["q"] * l["h"]
    coef = h_ch / h_rel  # les heures chargées (arrondi au 10) sont réparties au prorata des heures relevées
    par_article = {}
    for art in sorted(arts, key=lambda s: [int(x) for x in s.split(".")]):
        m = round(arts[art]["materiel"], 2)
        h = round(arts[art]["heures"], 4)
        par_article[art] = {"materiel": m, "heures": h, "total": round(m + h * coef * TAUX, 2)}
    ecart = round(total - sum(v["total"] for v in par_article.values()), 2)
    if ecart:  # résidu de centimes sur l'article le plus gros
        gros = max(par_article, key=lambda k: par_article[k]["total"])
        par_article[gros]["total"] = round(par_article[gros]["total"] + ecart, 2)

    sortie = {
        "taux": TAUX,
        "materiel": materiel,
        "heures_relevees": h_rel,
        "heures_chargees": h_ch,
        "main_doeuvre": mo,
        "total_acceo": total,
        "facteur": FACTEUR,
        "prix_vente": pv,
        "lignes": lignes,
        "par_article": par_article,
        "lots": lots_json,
    }
    with open(P_OUT, "w", encoding="utf-8") as f:
        json.dump(sortie, f, ensure_ascii=False, indent=2)

    # ---- Contrôles de cohérence (échec = exit code non nul) ----
    chk = json.load(open(P_OUT, encoding="utf-8"))
    assert abs(chk["materiel"] + chk["heures_chargees"] * chk["taux"] - chk["total_acceo"]) < 0.005, "total ACCEO incohérent"
    assert abs(chk["main_doeuvre"] - chk["heures_chargees"] * chk["taux"]) < 0.005, "MO incohérente"
    assert chk["heures_chargees"] % 10 == 0 and chk["heures_chargees"] >= chk["heures_relevees"] > chk["heures_chargees"] - 10
    assert chk["prix_vente"] % 1000 == 0 and abs(chk["prix_vente"] - chk["total_acceo"] * chk["facteur"]) <= 500.0001
    assert abs(sum(v["materiel"] for v in chk["par_article"].values()) - chk["materiel"]) < 0.02
    assert abs(sum(v["total"] for v in chk["par_article"].values()) - chk["total_acceo"]) < 0.02
    assert abs(sum(l["q"] * l["p"] for l in chk["lignes"]) - chk["materiel"]) < 0.02
    print("OK  lignes=%d  materiel=%.2f  h_rel=%.2f  h_ch=%d  MO=%.2f  total_ACCEO=%.2f  prix_vente=%d" % (
        len(lignes), materiel, h_rel, h_ch, mo, total, pv))
    part_mo = mo / total
    print("part MO = %.1f %%  ;  h par 100 k$ de prix = %.0f" % (100 * part_mo, h_ch / (pv / 100000.0)))


if __name__ == "__main__":
    main()
    sys.exit(0)
