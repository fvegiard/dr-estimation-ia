#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » à l'aveugle — dossier S-1902 (Bibliothèque Georgette-Lepage, Brossard).

Déterministe, stdlib + openpyxl. Lit le kit (dossier du script) et écrit chiffrage.json.
Sources :
  - releve.xlsx (totaux par libellé) + rapport-releve.md (ventilation par portée) : quantités sur plan
  - entree/reference-Estimation_2.xlsx, feuille Détail : quantités hors plan + heures « à valider » (jugement)
  - entree/reference-Estimation_2.xlsx, feuille Matériel Lumen : prix réguliers pour familles absentes du catalogue
  - catalogue-prix-temps-dupuis.csv : colonnes prix_a_utiliser / temps_a_utiliser (dernier Dupuis), jamais la médiane
  - METHODE.md : lots confirmés, taux, facteur
"""
import csv
import json
import math
import os
import re
import sys

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
P = lambda *a: os.path.join(ICI, *a)

TAUX = 85.0          # $/h — METHODE.md §5 : dossiers Dupuis août–octobre 2026
FACTEUR = 1.244      # METHODE.md §7 : unique couple observé (S-1844)
PI_PAR_M = 3.28084   # conversion des longueurs de la feuille Détail (m) vers l'unité Dupuis (pi)

# Lots fournisseur confirmés (METHODE.md §2, prix confirmés par Francis le 2026-10-20)
LOTS = {
    "Lot éclairage (A1, Axis SL1 4'/3'/1', trousses, E, Z6)": 11893.0,
    "Lot distribution (NF 225 A 42 cct + EDB 1P x20 + EJB 3P x3)": 6200.0,
    "JCI / Simplex (bypass, programmation, vérification ULC-S537, certificat)": 7400.0,
    "Détecteurs de fumée JCI (2)": 600.0,
    "GPR / scan de dalle": 1250.0,
}
# Ventilation des lots par article du bordereau (le lot distribution est scindé panneau / disjoncteurs
# selon le détail confirmé : panneau NF 3 400 $ ; EDB 95 $ x20 + EJB 300 $ x3 = 2 800 $)
LOTS_ART = [
    ("Lot éclairage (A1, Axis SL1 4'/3'/1', trousses, E, Z6)", "3.35", 11893.0),
    ("Lot distribution (NF 225 A 42 cct + EDB 1P x20 + EJB 3P x3)", "3.24", 3400.0),
    ("Lot distribution (NF 225 A 42 cct + EDB 1P x20 + EJB 3P x3)", "3.25", 95.0 * 20 + 300.0 * 3),
    ("JCI / Simplex (bypass, programmation, vérification ULC-S537, certificat)", "3.38", 7400.0),
    ("Détecteurs de fumée JCI (2)", "3.36", 600.0),
    ("GPR / scan de dalle", "3.32", 1250.0),
]


def lire_catalogue():
    cat = {}
    doublons = set()
    with open(P("catalogue-prix-temps-dupuis.csv"), encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            d = r["description"]
            if d in cat:
                doublons.add(d)
            cat[d] = {"um": r["um"], "prix": float(r["prix_a_utiliser"]), "temps": float(r["temps_a_utiliser"])}
    return cat, doublons


def lire_detail():
    wb = openpyxl.load_workbook(P("entree", "reference-Estimation_2.xlsx"), data_only=False)
    det = {}
    for r in wb["Détail"].iter_rows(min_row=6, values_only=True):
        code = r[0]
        if isinstance(code, str) and re.fullmatch(r"[A-H]-\d{2}", code):
            det[code] = {"desc": r[1], "q": r[3], "u": r[4], "h": r[7], "note": r[11]}
    lumen = {}
    for r in wb["Matériel Lumen"].iter_rows(min_row=5, values_only=True):
        if r[0] and isinstance(r[4], (int, float)):
            lumen[r[0]] = {"desc": r[1], "q": r[2], "u": r[3], "prix": float(r[4])}
    return det, lumen


def lire_releve():
    wb = openpyxl.load_workbook(P("releve.xlsx"), data_only=True)
    tot = {}
    for r in wb["Par libellé"].iter_rows(min_row=2, values_only=True):
        if r[0]:
            tot[r[0]] = int(r[2])
    # ventilation par portée : tableau « famille | total | portées » de rapport-releve.md
    portee = {}
    with open(P("rapport-releve.md"), encoding="utf-8") as f:
        for ligne in f:
            m = re.match(r"\|\s*([A-Z0-9_ ]+?)\s*\|\s*(\d+)\s*\|\s*([A-Z0-9 ,]+?)\s*\|\s*$", ligne)
            if not m:
                continue
            fam, total, txt = m.group(1), int(m.group(2)), m.group(3)
            d = {k: int(v) for k, v in re.findall(r"([A-Z]+)\s+(\d+)", txt)}
            if sum(d.values()) != total:
                sys.exit(f"ERREUR portées {fam}: {d} != {total}")
            portee[fam] = d
    for fam, d in portee.items():
        if tot.get(fam) != sum(d.values()):
            sys.exit(f"ERREUR relevé {fam}: releve.xlsx={tot.get(fam)} rapport={sum(d.values())}")
    if set(tot) != set(portee):
        sys.exit(f"ERREUR familles relevé/rapport : {set(tot) ^ set(portee)}")
    return tot, portee


def main():
    cat, doublons = lire_catalogue()
    det, lumen = lire_detail()
    tot, portee = lire_releve()

    def rel(fam, port):
        return portee[fam].get(port, 0)

    def dq(code):
        return det[code]["q"]

    def m2pi(m):
        return int(round(m * PI_PAR_M))

    lignes = []

    def cat_l(code, desc, q, lib, art, lot=False, qsrc=""):
        if lib not in cat:
            sys.exit(f"ERREUR libellé absent du catalogue : {lib!r}")
        if lib in doublons:
            sys.exit(f"ERREUR libellé en double dans le catalogue : {lib!r}")
        c = cat[lib]
        u = "pi" if c["um"] == "P" else "u"
        p = 0.0 if lot else c["prix"]
        src = f"Dupuis « {lib} » (prix {c['prix']} / temps {c['temps']} à utiliser)"
        if lot:
            src += " — 0 $ en ligne, prix au lot"
        if qsrc:
            src += f" ; qté : {qsrc}"
        lignes.append({"code": code, "desc": desc, "q": q, "u": u, "p": p, "h": c["temps"], "art": art, "src": src})

    def jug_l(code, art, q=None, desc=None, qsrc="Détail"):
        d = det[code]
        q = d["q"] if q is None else q
        lignes.append({"code": code, "desc": desc or d["desc"], "q": q, "u": d["u"], "p": 0.0, "h": float(d["h"]),
                       "art": art, "src": f"jugement (Estimation_2, Détail {code} : {d['h']} h/{d['u']}) ; qté : {qsrc}"})

    def lumen_l(code, desc, lcode, h, art, qsrc):
        L = lumen[lcode]
        lignes.append({"code": code, "desc": desc, "q": L["q"], "u": L["u"], "p": L["prix"], "h": h, "art": art,
                       "src": f"Lumen {lcode} « {L['desc']} » {L['prix']} $/{L['u']} (famille absente du catalogue Dupuis) ; "
                              f"temps : jugement {h} h ; qté : {qsrc}"})

    # ---------------- A — Démantèlement / relocalisation (3.22, 3.23, 5.19) ----------------
    cat_l("A-01", "Luminaires encastrés existants enlevés (E.A.E)", rel("ENCASTRE", "ENLEVER"),
          "ENLEVER FIXTURE ENCAST", "3.22", qsrc="relevé ENCASTRE / ENLEVER")
    cat_l("A-01", "Luminaires LED 1x4 existants enlevés (E.A.E)", rel("LED 1X4", "ENLEVER"),
          "ENLEVER FIXTURE 1X4", "3.22", qsrc="relevé LED 1X4 / ENLEVER")
    cat_l("A-01", "Luminaires LED 2x2 existants enlevés (E.A.E)", rel("LED 2X2", "ENLEVER"),
          "ENLEVER FIXTURE 2X2", "3.22", qsrc="relevé LED 2X2 / ENLEVER")
    cat_l("A-03", "Prises existantes enlevées (doubles 15 A + DDFT, E.A.E)",
          rel("PRISE DOUBLE 15A", "ENLEVER") + rel("PRISE DDFT", "ENLEVER"),
          "ENLEVER PRISE", "3.22", qsrc="relevé PRISE DOUBLE 15A + PRISE DDFT / ENLEVER")
    cat_l("A-03", "Sorties data existantes enlevées (E.A.E)", rel("SORTIE DATA", "ENLEVER"),
          "ENLEVER PRISE", "3.22", qsrc="relevé SORTIE DATA / ENLEVER")
    cat_l("A-03", "Interrupteur existant enlevé (E.A.E)", rel("INT SIMPLE", "ENLEVER"),
          "ENLEVER INTERRUPTEUR", "3.22", qsrc="relevé INT SIMPLE / ENLEVER")
    fams_div = ["MONUMENT PLANCHER", "BOITE JONCTION", "BOITIER C", "HP ENCASTRE PLAFOND", "DETECTEUR FUMEE",
                "DETECTEUR MOUVEMENT"]
    cat_l("A-03", "Équipements divers enlevés (monuments, JB, boîtiers C, HP, détecteurs)",
          sum(rel(f, "ENLEVER") for f in fams_div), "EQUIPEMENT A ENLEVER", "3.22",
          qsrc="relevé / ENLEVER : " + " + ".join(f"{f} {rel(f, 'ENLEVER')}" for f in fams_div))
    cat_l("A-05", "Relocalisation prises / sorties data (E.A.R)",
          rel("PRISE DOUBLE 15A", "RELOCALISER") + rel("PRISE DDFT", "RELOCALISER") + rel("SORTIE DATA", "RELOCALISER"),
          "RELO PRISE", "3.22", qsrc="relevé PRISE DOUBLE 15A + PRISE DDFT + SORTIE DATA / RELOCALISER")
    cat_l("A-06", "Relocalisation luminaires encastrés (E.A.R)", rel("ENCASTRE", "RELOCALISER"),
          "RELO ENCASTRE", "3.22", qsrc="relevé ENCASTRE / RELOCALISER")
    jug_l("A-07", "3.22")
    cat_l("A-08", "FabLab 210 : retrait temporaire + réinstallation luminaires linéaires existants", dq("A-08"),
          "RELO ENCASTRE", "3.22", qsrc="Détail A-08 (hors relevé, note 1 E-202)")
    cat_l("A-09", "FabLab 210 : retrait temporaire + réinstallation encastrés existants", dq("A-09"),
          "RELO ENCASTRE", "3.22", qsrc="Détail A-09 (hors relevé, note 1 E-202)")
    jug_l("A-10", "3.23", q=rel("MOTEUR MONOPHASE", "RELOCALISER"), qsrc="relevé MOTEUR MONOPHASE / RELOCALISER (E-103)")
    jug_l("A-11", "5.19")

    # ---------------- B — Distribution (3.24, 3.25) ----------------
    cat_l("B-01", "Panneau PP-3 neuf 225 A 347/600 V 42 cct NF (fourniture au lot distribution)", dq("B-01"),
          "PPM-N-01 PANN 225A 347/6", "3.24", lot=True, qsrc="Détail B-01")
    cat_l("B-01", "Enlèvement de l'ancien PP-3 (inclus à B-01 : démantèlement + transfert)", 1,
          "ENL PANN 42CCT 120/208V", "3.24", qsrc="Détail B-01 (note : incl. démantèlement ancien PP-3)")
    for code in ("B-02", "B-06", "B-07"):
        cat_l(code, det[code]["desc"], dq(code), "DISJ15A1P", "3.25", lot=True, qsrc=f"Détail {code}")
    cat_l("B-08", det["B-08"]["desc"], dq("B-08"), "DISJ 20A 1P 120V 10KA", "3.25", lot=True, qsrc="Détail B-08")
    for code in ("B-03", "B-04", "B-05"):
        cat_l(code, det[code]["desc"], dq(code), "DISJ 100A 3P 120/240V 22 K", "3.25", lot=True, qsrc=f"Détail {code}")
    jug_l("B-09", "3.24")

    # ---------------- C — Raccordements mécaniques (3.23, 3.26) ----------------
    cat_l("C-01", "Raccordement UTA-01 au toit (575 V 3Ø)", dq("C-01"), "UNITER UC AU TOIT", "3.23", qsrc="Détail C-01")
    cat_l("C-05", "Conduit EMT 1\" 4 #8 RW90 (alim. UTA-01) — C-05 + C-06", m2pi(dq("C-05")),
          "CONDUIT 1 04C08", "3.23", qsrc=f"Détail C-05 {dq('C-05')} m x {PI_PAR_M} (fils C-06 compris dans l'assemblage)")
    cat_l("C-02", "Raccordement VE-02N au toit (575 V 3Ø, VFD)", dq("C-02"), "MOTEUR", "3.23", qsrc="Détail C-02")
    cat_l("C-07", "Conduit EMT 3/4\" 4 #12 (VE-02N) — part C-07/C-08", m2pi(35),
          "CONDUIT 3/4 04C12", "3.23", qsrc="Détail C-07 note « 35 m + 3 x 15 m » : 35 m VE-02N")
    cat_l("C-04", "Raccordement humidificateurs HUM-1/2/3 (120 V 1Ø)", dq("C-04"), "HUMIDIFICATEUR", "3.23",
          qsrc="Détail C-04 = relevé MOTEUR MONOPHASE / INSTALLER (3)")
    cat_l("C-04", "Sectionneur 30 A sans fusible HUM-1/2/3", dq("C-04"),
          "SECTIONNEUR POR ENTRETIEN HUBHBL13R22", "3.23", qsrc="Détail C-04")
    cat_l("C-07", "Conduit EMT 3/4\" 3 #12 (HUM x3) — part C-07/C-08", m2pi(3 * 15),
          "CONDUIT 3/4 03C12", "3.23", qsrc="Détail C-07 note « 35 m + 3 x 15 m » : 45 m HUM")
    cat_l("C-03", "Chaudières CH-01/CH-02 : réutiliser tiroirs CCM-1 + câbles, raccord final", dq("C-03"),
          "MOTEUR", "3.26", qsrc="Détail C-03")
    cat_l("N-01", "Thermostats neufs", rel("THERMOSTAT", "INSTALLER"), "THERMOSTAT", "3.23",
          qsrc="relevé THERMOSTAT / INSTALLER")

    # ---------------- D/E — Services RDC + étage (3.27–3.34) ----------------
    n_ddft_fontaine = dq("D-07")
    cat_l("D-07", "Prise DDFT fontaine (P2-61)", n_ddft_fontaine, "PRISE DUPLEX 20A 125V GFI", "3.30",
          qsrc="Détail D-07 (1 des 22 « PRISE DDFT » INSTALLER du relevé)")
    cat_l("D-01", "Prises doubles 20 A Decora / M.A.L.T. isolée (postes, barre pleine)",
          rel("PRISE DDFT", "INSTALLER") - n_ddft_fontaine, "PRISE DUPLEX 15/20A 120V", "3.29",
          qsrc="relevé PRISE DDFT / INSTALLER (22) − 1 DDFT fontaine (R-007 : barre pleine = M.A.L.T. isolée)")
    cat_l("D-06", "Prises doubles 15 A 120 V neuves (postes, tables, cabines, murales)",
          rel("PRISE DOUBLE 15A", "INSTALLER"), "PRISE DUPLEX 15A 120V", "3.28", qsrc="relevé PRISE DOUBLE 15A / INSTALLER")
    cat_l("D-02", "Prises USB Leviton Decora (tables)", rel("PRISE USB", "INSTALLER"), "PRISE USB 125V 15 AMP", "3.27",
          qsrc="relevé PRISE USB / INSTALLER")
    n_data = rel("SORTIE DATA", "INSTALLER")
    n_data_mur = dq("D-08")
    cat_l("D-08", "Sorties data murales/colonnes : boîte + conduit 3/4\" vers entreplafond", n_data_mur,
          "TEL/DATA 3/4", "3.33", qsrc="Détail D-08 (part des sorties data INSTALLER du relevé)")
    cat_l("D-01", "Sorties data postes/mobilier (plaque, câble réutilisé)", n_data - n_data_mur, "TEL/DATA", "3.33",
          qsrc=f"relevé SORTIE DATA / INSTALLER ({n_data}) − {n_data_mur} murales D-08")
    cat_l("D-11", "Boutons d'appel / panique : boîte + conduit vide (D-11 + E-06)", dq("D-11") + dq("E-06"),
          "TEL/DATA 3/4", "3.31", qsrc="Détail D-11 + E-06")
    jug_l("D-09", "3.31")
    jug_l("D-10", "3.31")
    cat_l("D-12", "Percements de dalle carottés + coupe-feu (D-12 + E-07)", dq("D-12") + dq("E-07"),
          "PERCEMENT PLANCHER", "3.31", qsrc="Détail D-12 + E-07 (1 par course)")
    cat_l("D-13", "Conduit sous dalle 5 #12 vers P-2/P-3/P-4/P-6 (D-13 + E-08, fils D-15/E-10 compris)",
          m2pi(dq("D-13") + dq("E-08")), "CONDUIT 3/4 05C12", "3.31",
          qsrc=f"Détail D-13 {dq('D-13')} m + E-08 {dq('E-08')} m, x {PI_PAR_M}")
    cat_l("D-14", "Conduit vide 3/4\" data sous dalle (D-14 + E-09)", m2pi(dq("D-14") + dq("E-09")),
          "CONDUIT VIDE 3/4", "3.33", qsrc=f"Détail D-14 {dq('D-14')} m + E-09 {dq('E-09')} m, x {PI_PAR_M}")
    cat_l("D-16", "Boîtes de tirage/jonction conduits sous dalle RDC (16 courses x 2)", 16 * 2,
          "BOITE CARRE 2-1/8PO KO 1", "3.31", qsrc="Détail D-16 (lot) → 2 boîtes/course x 16 courses D-13 (base Lumen IBE72171K)")
    cat_l("E-11", "Boîtes de tirage/jonction conduits sous dalle étage (7 courses x 2)", 7 * 2,
          "BOITE CARRE 2-1/8PO KO 1", "3.31", qsrc="Détail E-11 (lot) → 2 boîtes/course x 7 courses E-08")
    cat_l("N-02", "Scan de dalle (GPR) avant percements", 1, "SCAN DALLE BETON", "3.32", lot=True,
          qsrc="lot GPR confirmé METHODE.md")
    cat_l("E-02", "Monuments de plancher (poke-through) neufs", rel("MONUMENT PLANCHER", "INSTALLER"),
          "MONUMENT PLANCHER", "3.34", qsrc="relevé MONUMENT PLANCHER / INSTALLER (couvre E-02 + E-03)")
    cat_l("N-03", "Lecteurs de carte neufs (boîte + conduit)", rel("LECTEUR CARTE", "INSTALLER"),
          "LECTEUR DE CARTE", "3.31", qsrc="relevé LECTEUR CARTE / INSTALLER")

    # ---------------- F — Éclairage (3.35) ----------------
    cat_l("F-01", "Luminaire A1 Cooper 2'x4' encastré", rel("LUMINAIRE A1", "INSTALLER"), "FIXTURE TYPE A1", "3.35",
          lot=True, qsrc="relevé LUMINAIRE A1")
    cat_l("F-02", "Luminaire C1 Axis SL1 module 4'", rel("LUMINAIRE C1", "INSTALLER"), "FIXTURE TYPE C1", "3.35",
          lot=True, qsrc="relevé LUMINAIRE C1")
    cat_l("F-03", "Luminaire C1_1 Axis SL1 module 1'", rel("LUMINAIRE C1_1", "INSTALLER"), "FIXTURE TYPE C1", "3.35",
          lot=True, qsrc="relevé LUMINAIRE C1_1")
    cat_l("F-04", "Luminaire C1_3 Axis SL1 module 3'", rel("LUMINAIRE C1_3", "INSTALLER"), "FIXTURE TYPE C1", "3.35",
          lot=True, qsrc="relevé LUMINAIRE C1_3")
    jug_l("F-05", "3.35")
    cat_l("F-06", "Luminaire E GreenCreative cylindre suspendu", rel("LUMINAIRE E", "INSTALLER"), "FIXTURE TYPE E",
          "3.35", lot=True, qsrc="relevé LUMINAIRE E (3, = lot confirmé E x3)")
    cat_l("F-07", "Enseigne de sortie pictogramme Z6", rel("ENSEIGNE SORTIE Z6", "INSTALLER"), "ENSEIGNE SORTIE",
          "3.35", lot=True, qsrc="relevé ENSEIGNE SORTIE Z6")
    cat_l("N-04", "Interrupteurs d'éclairage neufs", rel("INT SIMPLE", "INSTALLER"), "INTERRUPTEUR 15A 120V", "3.35",
          qsrc="relevé INT SIMPLE / INSTALLER (E-203)")
    cat_l("F-08", "Conduit EMT 3/4\" 3 #12 raccordement luminaires (F-08 + F-09)", m2pi(dq("F-08")),
          "CONDUIT 3/4 03C12", "3.35", qsrc=f"Détail F-08 {dq('F-08')} m x {PI_PAR_M} (fils F-09 compris)")
    cat_l("F-10", "Boîtes de jonction éclairage (35 points)", lumen["IBE54151L"]["q"], "BOITE OCTOG 1-1/2PO KO 1",
          "3.35", qsrc="Détail F-10 (lot) → 35 points (base Lumen IBE54151L / F-08)")
    cat_l("F-10", "Tige filetée 3/8\" suspension rangées Axis à la dalle (devis 16.5)", lumen["VISTF3/8"]["q"],
          "TIGE FILETE 3/8PO", "3.35", qsrc="Détail F-10 → base Lumen VISTF3/8 (200 pi)")
    lumen_l("F-10", "Ancrages drop-in 3/8-16 (pqt 50) pour tiges", "VISDIA3/816BX50", 0.0, "3.35",
            "Lumen VISDIA3/816BX50 (2 pqt) ; pose comprise dans la tige filetée / F-05")

    # ---------------- G — Alarme incendie (3.36, 3.37, 3.38) ----------------
    cat_l("G-01", "Détecteurs de fumée adressables neufs (fournis JCI)", rel("DETECTEUR FUMEE", "INSTALLER"),
          "DETECTEUR A FUMEE", "3.36", lot=True, qsrc="relevé DETECTEUR FUMEE / INSTALLER (= Détail G-01)")
    cat_l("G-02", "Conduit 3/4\" + câble FT4 boucle AI (G-02 + G-03)", m2pi(dq("G-02")), "COND ALARME 3/16 FT4",
          "3.36", qsrc=f"Détail G-02 {dq('G-02')} m x {PI_PAR_M} (câble G-03 compris)")
    cat_l("G-05", "Relocalisation détecteur de fumée existant (E.A.R)", rel("DETECTEUR FUMEE", "RELOCALISER"),
          "RELO DETECTEUR A FUME", "3.36", qsrc="relevé DETECTEUR FUMEE / RELOCALISER")
    jug_l("G-05", "3.36")
    cat_l("G-04", "Relocalisation panneau annonciateur Simplex", rel("PANNEAU ANNONCIATEUR", "INSTALLER"),
          "RELO PANN ANNOCIATEUR", "3.37", qsrc="relevé PANNEAU ANNONCIATEUR (= Détail G-04)")
    jug_l("G-06", "3.38")
    jug_l("G-07", "3.38", desc=det["G-07"]["desc"] + " — montant au lot JCI")

    # ---------------- H — Frais de chantier (6.1) ----------------
    cat_l("H-01", "Mobilisation, plateforme élévatrice / nacelle", dq("H-01"), "NACELLE", "6.1", qsrc="Détail H-01")
    jug_l("H-02", "6.1")
    jug_l("H-03", "6.1")
    jug_l("H-04", "6.1")

    # ---------------- Totaux ----------------
    r2 = lambda x: round(x + 1e-9, 2)
    for l in lignes:
        if l["q"] is None or l["q"] < 0:
            sys.exit(f"ERREUR quantité invalide : {l}")
    mat_lignes = sum(l["q"] * l["p"] for l in lignes)
    mat_lots = sum(LOTS.values())
    if abs(sum(m for _, _, m in LOTS_ART) - mat_lots) > 0.005:
        sys.exit("ERREUR ventilation des lots")
    materiel = r2(mat_lignes + mat_lots)
    heures = round(sum(l["q"] * l["h"] for l in lignes), 4)
    heures_chargees = int(math.ceil(round(heures, 6) / 10.0) * 10)
    main_doeuvre = r2(heures_chargees * TAUX)
    total_acceo = r2(materiel + main_doeuvre)
    prix_vente = int(round(total_acceo * FACTEUR / 1000.0) * 1000)

    par_article = {}
    for l in lignes:
        a = par_article.setdefault(l["art"], {"materiel": 0.0, "heures": 0.0, "total": 0.0})
        a["materiel"] += l["q"] * l["p"]
        a["heures"] += l["q"] * l["h"]
    for _, art, m in LOTS_ART:
        a = par_article.setdefault(art, {"materiel": 0.0, "heures": 0.0, "total": 0.0})
        a["materiel"] += m
    for a in par_article.values():
        a["materiel"] = r2(a["materiel"])
        a["heures"] = round(a["heures"], 3)
        a["total"] = r2(a["materiel"] + a["heures"] * TAUX)
    par_article = dict(sorted(par_article.items(), key=lambda kv: [int(x) for x in kv[0].split(".")]))

    # contrôles de cohérence
    assert abs(materiel + heures_chargees * TAUX - total_acceo) < 0.005
    assert prix_vente == int(round(total_acceo * FACTEUR / 1000.0) * 1000)
    assert heures_chargees >= heures and heures_chargees - heures < 10
    assert abs(sum(a["materiel"] for a in par_article.values()) - materiel) < 0.05
    assert abs(sum(a["heures"] for a in par_article.values()) - heures) < 0.01

    out = {
        "taux": TAUX,
        "materiel": materiel,
        "heures_relevees": round(heures, 3),
        "heures_chargees": heures_chargees,
        "main_doeuvre": main_doeuvre,
        "total_acceo": total_acceo,
        "facteur": FACTEUR,
        "prix_vente": prix_vente,
        "lignes": [{**l, "q": l["q"], "p": round(l["p"], 4), "h": round(l["h"], 4)} for l in lignes],
        "par_article": par_article,
        "lots": LOTS,
    }
    with open(P("chiffrage.json"), "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)

    print(f"lignes={len(lignes)} materiel={materiel:.2f} (lignes {mat_lignes:.2f} + lots {mat_lots:.2f}) "
          f"heures={heures:.3f} -> {heures_chargees} h x {TAUX} = {main_doeuvre:.2f} ; total ACCEO={total_acceo:.2f} ; "
          f"x {FACTEUR} = {total_acceo * FACTEUR:.2f} -> prix de vente {prix_vente}")
    part_mo = main_doeuvre / total_acceo * 100
    print(f"contrôle : part MO {part_mo:.1f} % ; {heures_chargees / prix_vente * 100000:.0f} h / 100 k$ de prix")
    return 0


if __name__ == "__main__":
    sys.exit(main())
