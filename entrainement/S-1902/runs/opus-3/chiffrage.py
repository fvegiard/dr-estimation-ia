#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » à l'aveugle — S-1902 Bibliothèque Georgette-Lepage (Brossard).

Déterministe, autonome (stdlib + openpyxl). Lit le kit du dossier courant et écrit chiffrage.json.
- Quantités sur plan : releve.xlsx (feuille « Par libellé ») + portées du tableau de rapport-releve.md.
- Quantités hors plan et heures de jugement : entree/reference-Estimation_2.xlsx, feuille « Détail ».
- Prix / temps : catalogue-prix-temps-dupuis.csv, colonnes prix_a_utiliser / temps_a_utiliser (dernier Dupuis).
- Lots, taux et facteur : METHODE.md.
Convention des lignes : p = prix unitaire ($), h = heures unitaires; total ligne = q*p et q*h.
"""
import csv
import json
import math
import os
import re
import sys

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
P = lambda *a: os.path.join(ICI, *a)
M_TO_PI = 3.28084


def erreur(msg):
    sys.stderr.write("ERREUR : " + msg + "\n")
    sys.exit(1)


# ---------------------------------------------------------------- catalogue Dupuis
def lire_catalogue():
    cat = {}
    with open(P("catalogue-prix-temps-dupuis.csv"), encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            cle = r["description"].strip()
            cat.setdefault(cle, []).append(r)
    return cat


CAT = lire_catalogue()


def dupuis(libelle):
    """(prix, temps) dernier Dupuis pour un libellé EXACT; refuse un libellé absent ou ambigu."""
    rows = CAT.get(libelle)
    if not rows:
        erreur("libellé absent du catalogue : %r" % libelle)
    if len(rows) > 1:
        erreur("libellé ambigu dans le catalogue : %r (%d lignes)" % (libelle, len(rows)))
    r = rows[0]
    return float(r["prix_a_utiliser"]), float(r["temps_a_utiliser"]), r["um"]


# ---------------------------------------------------------------- relevé
def lire_releve():
    wb = openpyxl.load_workbook(P("releve.xlsx"), read_only=True, data_only=True)
    ws = wb["Par libellé"]
    qte = {}
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or row[0] is None:
            continue
        qte[str(row[0]).strip()] = int(row[2])
    wb.close()
    # portées : tableau « famille | total | portées » de rapport-releve.md
    portees = {}
    with open(P("rapport-releve.md"), encoding="utf-8") as f:
        for ligne in f:
            m = re.match(r"^\|\s*([A-Z0-9_ ]+?)\s*\|\s*(\d+)\s*\|\s*([A-Z ,0-9]+?)\s*\|\s*$", ligne)
            if not m:
                continue
            fam, tot, txt = m.group(1), int(m.group(2)), m.group(3)
            d = {}
            for morceau in txt.split(","):
                k, v = morceau.strip().rsplit(" ", 1)
                d[k] = int(v)
            if sum(d.values()) != tot:
                erreur("portées incohérentes pour %s" % fam)
            portees[fam] = d
    for fam, n in qte.items():
        if fam not in portees:
            erreur("famille sans portée dans rapport-releve.md : %s" % fam)
        if sum(portees[fam].values()) != n:
            erreur("relevé %s = %d mais portées = %d" % (fam, n, sum(portees[fam].values())))
    return qte, portees


QTE, PORTEE = lire_releve()


def rel(fam, portee):
    return PORTEE[fam].get(portee, 0)


# ---------------------------------------------------------------- Estimation_2 / Détail
def lire_detail():
    wb = openpyxl.load_workbook(P("entree", "reference-Estimation_2.xlsx"), read_only=True, data_only=False)
    ws = wb["Détail"]
    det = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        code = row[0]
        if isinstance(code, str) and re.match(r"^[A-H]-\d\d$", code):
            det[code] = {"desc": row[1], "q": row[3], "u": row[4], "h": row[7]}
    wb.close()
    return det


DET = lire_detail()


def d(code, champ):
    v = DET[code][champ]
    if v is None:
        erreur("Détail %s : champ %s vide" % (code, champ))
    return v


def pieds(code):
    """Longueur Détail en mètres convertie en pieds (arrondi au pied supérieur)."""
    if DET[code]["u"] != "m":
        erreur("Détail %s n'est pas en mètres" % code)
    return math.ceil(d(code, "q") * M_TO_PI - 1e-9)


# ---------------------------------------------------------------- METHODE.md : taux, facteur, lots
with open(P("METHODE.md"), encoding="utf-8") as f:
    METH = f.read()
for jeton in ("85 $/h", "1,244", "lot éclairage 11 893 $", "lot distribution 6 200 $", "NF 225 A 42 cct 3 400 $",
              "JCI 7 400 $ + 2 détecteurs 600 $", "GPR 1 250 $"):
    if jeton not in METH:
        erreur("METHODE.md ne contient plus %r" % jeton)
TAUX = 85.0
FACTEUR = 1.244
LOTS = {
    "eclairage": 11893.0,
    "distribution": 6200.0,
    "alarme_JCI": 7400.0,
    "detecteurs_fumee": 600.0,
    "GPR_scan_dalle": 1250.0,
}
DIST_PANNEAU = 3400.0                       # NF 225 A 42 cct
DIST_DISJ = LOTS["distribution"] - DIST_PANNEAU  # EDB 1P 95 $ x20 + EJB 3P 300 $ x3 = 2 800 $
if abs(DIST_DISJ - (95 * 20 + 300 * 3)) > 1e-9:
    erreur("ventilation du lot distribution incohérente")

# ---------------------------------------------------------------- lignes
LIGNES = []


def cat_ligne(code, desc, q, libelle, art, src, p0=False, u=None):
    """Ligne au prix/temps Dupuis (dernier). p0=True : matériel au prix de lot (0 $ en ligne)."""
    prix, temps, um = dupuis(libelle)
    src_txt = "catalogue Dupuis « %s » (%s $ / %s h, dernier)" % (libelle, prix, temps)
    if p0:
        src_txt += " — 0 $ en ligne, fourniture au prix de lot"
        prix = 0.0
    LIGNES.append({"code": code, "desc": desc, "q": q, "u": u or ("pi" if um == "P" else "u"),
                   "p": prix, "h": temps, "art": art, "src": src_txt + " ; qté : " + src})


def jug_ligne(code, desc, q, u, h, art, src, p=0.0):
    LIGNES.append({"code": code, "desc": desc, "q": q, "u": u, "p": p, "h": h, "art": art, "src": src})


def lot_ligne(code, desc, montant, art, src):
    LIGNES.append({"code": code, "desc": desc, "q": 1, "u": "lot", "p": montant, "h": 0.0, "art": art, "src": src})


R = "relevé"
J = "jugement (Estimation_2)"

# ===== A — Démantèlement / démolition (5.19) et démontage-relocalisation (3.22)
cat_ligne("A-01a", "Enlever luminaires encastrés existants (E.A.E)", rel("ENCASTRE", "ENLEVER"),
          "ENLEVER FIXTURE ENCAST", "5.19", R + " ENCASTRE ENLEVER (remplace A-01/A-02)")
cat_ligne("A-01b", "Enlever luminaires 1x4 existants (E.A.E)", rel("LED 1X4", "ENLEVER"),
          "ENLEVER FIXTURE 1X4", "5.19", R + " LED 1X4 ENLEVER")
cat_ligne("A-01c", "Enlever luminaires 2x2 existants (E.A.E)", rel("LED 2X2", "ENLEVER"),
          "ENLEVER FIXTURE 2X2", "5.19", R + " LED 2X2 ENLEVER")
cat_ligne("A-03a", "Enlever prises existantes (doubles + DDFT, E.A.E)",
          rel("PRISE DOUBLE 15A", "ENLEVER") + rel("PRISE DDFT", "ENLEVER"),
          "ENLEVER PRISE", "5.19", R + " PRISE DOUBLE 15A + PRISE DDFT ENLEVER (remplace A-03/A-04)")
cat_ligne("A-03b", "Enlever sorties informatiques existantes (E.A.E)", rel("SORTIE DATA", "ENLEVER"),
          "ENLEVER PRISE", "5.19", R + " SORTIE DATA ENLEVER (pas de libellé « enlever data » : temps prise)")
cat_ligne("A-03c", "Enlever interrupteur existant (E.A.E)", rel("INT SIMPLE", "ENLEVER"),
          "ENLEVER INTERRUPTEUR", "5.19", R + " INT SIMPLE ENLEVER")
n_eq = (rel("MONUMENT PLANCHER", "ENLEVER") + rel("BOITE JONCTION", "ENLEVER") + rel("BOITIER C", "ENLEVER")
        + rel("DETECTEUR MOUVEMENT", "ENLEVER") + rel("DETECTEUR FUMEE", "ENLEVER")
        + rel("HP ENCASTRE PLAFOND", "ENLEVER"))
cat_ligne("A-03d", "Enlever équipements divers (monuments 3, boîtes de jonction 2, boîtiers C 2, "
          "détecteur mouvement 1, détecteur fumée 1, haut-parleur 1)", n_eq,
          "EQUIPEMENT A ENLEVER", "5.19", R + " portées ENLEVER des familles listées")
cat_ligne("A-05a", "Relocaliser prises existantes (E.A.R)",
          rel("PRISE DOUBLE 15A", "RELOCALISER") + rel("PRISE DDFT", "RELOCALISER"),
          "RELO PRISE", "3.22", R + " PRISE DOUBLE 15A + PRISE DDFT RELOCALISER (remplace A-05)")
cat_ligne("A-05b", "Relocaliser sorties informatiques existantes (E.A.R)", rel("SORTIE DATA", "RELOCALISER"),
          "RELO PRISE", "3.22", R + " SORTIE DATA RELOCALISER")
cat_ligne("A-06", "Relocaliser luminaires encastrés existants (E.A.R)", rel("ENCASTRE", "RELOCALISER"),
          "RELO ENCASTRE", "3.22", R + " ENCASTRE RELOCALISER")
jug_ligne("A-07", d("A-07", "desc"), d("A-07", "q"), d("A-07", "u"), d("A-07", "h"), "3.22",
          J + " A-07 — aucun libellé Dupuis équivalent")
jug_ligne("A-08", d("A-08", "desc"), d("A-08", "q"), d("A-08", "u"), d("A-08", "h"), "3.22",
          J + " A-08 — retrait temporaire + réinstallation, aucun libellé Dupuis équivalent")
jug_ligne("A-09", d("A-09", "desc"), d("A-09", "q"), d("A-09", "u"), d("A-09", "h"), "3.22",
          J + " A-09 — retrait temporaire + réinstallation, aucun libellé Dupuis équivalent")
n_toit = rel("MOTEUR MONOPHASE", "RELOCALISER")
if n_toit != d("A-10", "q"):
    sys.stderr.write("AVIS : relevé toit %d vs Détail A-10 %s\n" % (n_toit, d("A-10", "q")))
cat_ligne("A-10a", "Toit : débrancher équipements de ventilation (E.A.R)", n_toit,
          "DEBRANCHEMENT APPAREILAGE MECANIQUE", "3.22", R + " MOTEUR MONOPHASE RELOCALISER (E-103)")
cat_ligne("A-10b", "Toit : enlever et entreposer sectionneurs", n_toit,
          "SECTIONNEUR ENL", "3.22", R + " MOTEUR MONOPHASE RELOCALISER (E-103)")
cat_ligne("A-10c", "Toit : rebrancher équipements de ventilation", n_toit,
          "MOTEUR", "3.22", R + " MOTEUR MONOPHASE RELOCALISER (E-103)")
jug_ligne("A-11", d("A-11", "desc"), d("A-11", "q"), d("A-11", "u"), d("A-11", "h"), "5.19",
          J + " A-11 — aucun libellé Dupuis équivalent")

# ===== B — Distribution
cat_ligne("B-01", "Panneau PP-3 neuf 225 A 347/600 V 42 cct NF — remplacement de l'existant", 1,
          "PANNEAU A CHANGER", "3.24", "Détail B-01 (1 u)", p0=True)
cat_ligne("B-02", "Disjoncteurs 1P (transferts PP-3 + P2-61/63/64 + P3-17/19)", 20,
          "DISJ15A1P", "3.25", "lot distribution confirmé (EDB 1P x20); Détail B-02/B-06/B-07/B-08 = 19", p0=True)
cat_ligne("B-03", "Disjoncteurs 3P (aérotherme, UTA-1, VE-2N)", 3,
          "DISJ 100A 3P 120/240V 22 K", "3.25",
          "lot distribution confirmé (EJB 3P x3) = Détail B-03/B-04/B-05", p0=True)
jug_ligne("B-09", d("B-09", "desc"), d("B-09", "q"), d("B-09", "u"), d("B-09", "h"), "3.24",
          J + " B-09 — aucun libellé Dupuis équivalent")
lot_ligne("LOT-02a", "Lot distribution — panneau NF 225 A 42 cct", DIST_PANNEAU, "3.24",
          "METHODE.md lot distribution 6 200 $ (part panneau 3 400 $), confirmé par Francis 2026-10-20")
lot_ligne("LOT-02b", "Lot distribution — disjoncteurs EDB 1P x20 + EJB 3P x3", DIST_DISJ, "3.25",
          "METHODE.md lot distribution 6 200 $ (part disjoncteurs 2 800 $), confirmé par Francis 2026-10-20")

# ===== C — Raccordements mécaniques
cat_ligne("C-01", "UTA-01 York 575 V/3Ø au toit — raccordement", 1, "UNITER UC AU TOIT", "3.23",
          R + " MOTEUR TRIPHASE INSTALLER (E-106) / Détail C-01")
cat_ligne("C-02", "VE-02N Plymovent 575 V/3Ø c/a VFD au toit — raccordement", 1, "MOTEUR", "3.23",
          R + " MOTEUR TRIPHASE INSTALLER (E-106) / Détail C-02")
cat_ligne("C-03", "Chaudières CH-01/CH-02 600 V/3Ø — réutiliser tiroirs CCM-1, raccord final", d("C-03", "q"),
          "CHAUDIERE", "3.26", R + " MOTEUR TRIPHASE INSTALLER (E-105) = Détail C-03")
cat_ligne("C-04a", "Humidificateurs HUM-1/2/3 120 V/1Ø — raccordement", rel("MOTEUR MONOPHASE", "INSTALLER"),
          "HUMIDIFICATEUR", "3.23", R + " MOTEUR MONOPHASE INSTALLER (E-105) = Détail C-04")
cat_ligne("C-04b", "Sectionneur sans fusible 30 A — humidificateurs", rel("MOTEUR MONOPHASE", "INSTALLER"),
          "INTERR 30A 2P 600V CEMA1", "3.23", R + " MOTEUR MONOPHASE INSTALLER / Détail C-04")
cat_ligne("C-05", "Conduit EMT 1\" + 4 #8 (UTA-01) — conducteurs C-06 inclus dans l'assemblage", pieds("C-05"),
          "CONDUIT 1 03C08", "3.23", "Détail C-05 %s m converti en pi" % d("C-05", "q"))
pi_ve = math.ceil(35 * M_TO_PI - 1e-9)
pi_hum = pieds("C-07") - pi_ve
cat_ligne("C-07a", "Conduit EMT 3/4\" + 4 #12 (VE-02N) — conducteurs C-08 inclus", pi_ve,
          "CONDUIT 3/4 04C12", "3.23", "Détail C-07 (35 m VE-02N) converti en pi")
cat_ligne("C-07b", "Conduit EMT 3/4\" + 3 #12 (HUM x3) — conducteurs C-08 inclus", pi_hum,
          "CONDUIT 3/4 03C12", "3.23", "Détail C-07 (3 x 15 m HUM) converti en pi")
cat_ligne("R-06", "Thermostats neufs", rel("THERMOSTAT", "INSTALLER"), "THERMOSTAT", "3.23",
          R + " THERMOSTAT INSTALLER (E-104)")

# ===== Services — dispositifs (quantités relevé E-104/E-105, remplacent D-01..D-06, E-01, E-03..E-05)
n_ddft = rel("PRISE DDFT", "INSTALLER")
n_gfi = d("D-07", "q")
cat_ligne("R-01", "Prises doubles 15 A 120 V neuves", rel("PRISE DOUBLE 15A", "INSTALLER"),
          "PRISE DUPLEX 15A 120V", "3.28", R + " PRISE DOUBLE 15A INSTALLER")
cat_ligne("R-02", "Prises doubles 15 A 120 V M.A.L.T. isolée (postes de travail)", n_ddft - n_gfi,
          "PRISE DUPLEX 15A 120V", "3.29",
          R + " PRISE DDFT INSTALLER (22) moins 1 DDFT fontaine — symbole barre pleine = M.A.L.T. isolée (R-007)")
cat_ligne("D-07", "Prise DDFT 15 A fontaine (P2-61)", n_gfi, "PRISE DUPLEX 20A 125V GFI", "3.30",
          "Détail D-07 (1 u), pris dans PRISE DDFT INSTALLER du relevé")
cat_ligne("R-03", "Prises doubles 15 A USB Leviton Decora", rel("PRISE USB", "INSTALLER"),
          "PRISE USB 125V 15 AMP", "3.27", R + " PRISE USB INSTALLER")
n_data = rel("SORTIE DATA", "INSTALLER")
n_data_mur = d("D-08", "q")
cat_ligne("D-08", "Sorties informatiques murales/colonne c/a conduit 3/4\"", n_data_mur,
          "TEL/DATA 3/4", "3.33", "Détail D-08 (10 u), pris dans SORTIE DATA INSTALLER du relevé")
cat_ligne("R-04", "Sorties informatiques (postes, tables, mobilier)", n_data - n_data_mur,
          "TEL/DATA", "3.33", R + " SORTIE DATA INSTALLER (65) moins D-08")
cat_ligne("R-05", "Lecteurs de carte neufs", rel("LECTEUR CARTE", "INSTALLER"), "LECTEUR DE CARTE", "3.33",
          R + " LECTEUR CARTE INSTALLER")
cat_ligne("E-02", "Monuments de plancher (poke-through) prise 20 A + 2 RJ45", rel("MONUMENT PLANCHER", "INSTALLER"),
          "MONUMENT PLANCHER", "3.34", R + " MONUMENT PLANCHER INSTALLER (E-105; Détail E-02 + E-03 = 4)")
cat_ligne("D-09", "Boîtes de jonction + raccords (barrière anti-vol, bornes de prêts)",
          rel("BOITE JONCTION", "INSTALLER"), "RACCORD DIRECT", "3.31",
          R + " BOITE JONCTION INSTALLER (E-104) = Détail D-09 (2) + D-10 (1)")
cat_ligne("D-11", "Bouton d'appel RDC — boîte (conduit vide compté en D-14)", d("D-11", "q"),
          "DECLENCHEUR APPEL", "3.31", "Détail D-11")
cat_ligne("E-06", "Boutons d'appel (2) + bouton de panique (1) étage — boîtes (conduit vide en E-09)",
          d("E-06", "q"), "DECLENCHEUR APPEL", "3.31", "Détail E-06")

# ===== Percements, conduits, filage
cat_ligne("D-12", "Percement de dalle carotté RDC (1 par course)", d("D-12", "q"), "PERCEMENT PLANCHER", "3.31",
          "Détail D-12")
cat_ligne("E-07", "Percement de dalle carotté étage", d("E-07", "q"), "PERCEMENT PLANCHER", "3.31", "Détail E-07")
cat_ligne("D-12b", "Coupe-feu homologué aux percements (1 tube par percement)", d("D-12", "q") + d("E-07", "q"),
          "COUPE FEU TUBE 10.1 FL O", "3.31", "Détail D-12 + E-07 (devis CAN/ULC-S115)")
cat_ligne("D-13", "Conduit 3/4\" + filage #12 sous dalle RDC (16 courses x 30 m) — D-15 inclus", pieds("D-13"),
          "CONDUIT 3/4 03C12", "3.31", "Détail D-13 %s m converti en pi" % d("D-13", "q"))
cat_ligne("E-08", "Conduit 3/4\" + filage #12 sous dalle étage (7 courses x 25 m) — E-10 inclus", pieds("E-08"),
          "CONDUIT 3/4 03C12", "3.31", "Détail E-08 %s m converti en pi" % d("E-08", "q"))
cat_ligne("D-14", "Conduit vide 3/4\" data/bouton RDC + corde de tirage (11 courses)", pieds("D-14"),
          "CONDUIT VIDE 3/4", "3.33", "Détail D-14 %s m converti en pi" % d("D-14", "q"))
cat_ligne("E-09", "Conduit vide 3/4\" data/boutons étage (7 courses)", pieds("E-09"),
          "CONDUIT VIDE 3/4", "3.33", "Détail E-09 %s m converti en pi" % d("E-09", "q"))
pq, tq, _ = dupuis("QUINCAILLERIE")
jug_ligne("D-16", d("D-16", "desc") + " (forfait quincaillerie Dupuis, couvre aussi E-11/F-10)", 1, "lot",
          d("D-16", "h"), "3.31",
          "prix : catalogue Dupuis « QUINCAILLERIE » (%s $, dernier) ; heures : %s D-16" % (pq, J), p=pq)
jug_ligne("E-11", d("E-11", "desc"), d("E-11", "q"), d("E-11", "u"), d("E-11", "h"), "3.31",
          J + " E-11 — matériel dans le forfait D-16")
ps, ts, _ = dupuis("SCAN DALLE BETON")
if abs(ps - LOTS["GPR_scan_dalle"]) > 1e-9:
    sys.stderr.write("AVIS : prix SCAN DALLE BETON %s ≠ lot GPR %s\n" % (ps, LOTS["GPR_scan_dalle"]))
jug_ligne("D-17", "Scan de dalle béton (GPR) avant percements — lot GPR", 1, "lot", ts, "3.32",
          "prix : lot GPR 1 250 $ confirmé (METHODE.md) ; heures : catalogue Dupuis « SCAN DALLE BETON » (%s h)" % ts,
          p=LOTS["GPR_scan_dalle"])

# ===== F — Éclairage (lot fournisseur, temps Dupuis)
cat_ligne("F-01", "Luminaire A1 Cooper 2x4 encastré 347 V", rel("LUMINAIRE A1", "INSTALLER"),
          "FIXTURE TYPE A1", "3.35", R + " LUMINAIRE A1 INSTALLER", p0=True)
cat_ligne("F-02", "Luminaire C1 Axis SL1 module 4'", rel("LUMINAIRE C1", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1 INSTALLER", p0=True)
cat_ligne("F-03", "Luminaire C1_1 Axis SL1 module 1'", rel("LUMINAIRE C1_1", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1_1 INSTALLER", p0=True)
cat_ligne("F-04", "Luminaire C1_3 Axis SL1 module 3'", rel("LUMINAIRE C1_3", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1_3 INSTALLER", p0=True)
jug_ligne("F-05", d("F-05", "desc"), d("F-05", "q"), d("F-05", "u"), d("F-05", "h"), "3.35",
          J + " F-05 — trousses de suspension dans le lot éclairage")
cat_ligne("F-06", "Luminaire E GreenCreative cylindre suspendu", rel("LUMINAIRE E", "INSTALLER"),
          "FIXTURE TYPE E", "3.35", R + " LUMINAIRE E INSTALLER (Détail F-06 = 4, lot = 3)", p0=True)
cat_ligne("F-07", "Enseigne de sortie Z6 Stanpro RMSOWH-UDC", rel("ENSEIGNE SORTIE Z6", "INSTALLER"),
          "ENSEIGNE SORTIE Z6", "3.35", R + " ENSEIGNE SORTIE Z6 INSTALLER", p0=True)
cat_ligne("R-07", "Interrupteurs neufs (éclairage E-203)", rel("INT SIMPLE", "INSTALLER"),
          "INTERRUPTEUR 15A 120V", "3.35", R + " INT SIMPLE INSTALLER")
cat_ligne("F-08", "Conduit 3/4\" + filage #12 éclairage (35 points x 8 m) — F-09 inclus", pieds("F-08"),
          "CONDUIT 3/4 03C12", "3.35", "Détail F-08 %s m converti en pi" % d("F-08", "q"))
cat_ligne("F-08b", "Alimentations de luminaires (flexible boîte → appareil)", 35,
          "ALIMENTATION ECLAIRAGE", "3.35", "Détail F-08 (35 points d'éclairage)")
jug_ligne("F-10", d("F-10", "desc"), d("F-10", "q"), d("F-10", "u"), d("F-10", "h"), "3.35",
          J + " F-10 — boîtes dans le forfait D-16")
cat_ligne("F-10b", "Tige filetée 3/8\" — suspension des rangées à la dalle (devis 16.5)", 200,
          "TIGE FILETE 3/8PO", "3.35", "quantité de la feuille Matériel Lumen (VISTF3/8, 200 pi)")
lot_ligne("LOT-01", "Lot éclairage (A1, Axis SL1 4'/3'/1', trousses, E, Z6)", LOTS["eclairage"], "3.35",
          "METHODE.md lot éclairage 11 893 $, confirmé par Francis 2026-10-20")

# ===== G — Alarme incendie
cat_ligne("G-01", "Détecteurs de fumée adressables Simplex (fournis JCI) — installation",
          rel("DETECTEUR FUMEE", "INSTALLER"), "DETECTEUR A FUMEE", "3.36",
          R + " DETECTEUR FUMEE INSTALLER = Détail G-01", p0=True)
cat_ligne("G-08", "Relocaliser détecteur de fumée existant", rel("DETECTEUR FUMEE", "RELOCALISER"),
          "RELO DETECTEUR A FUME", "3.36", R + " DETECTEUR FUMEE RELOCALISER")
cat_ligne("G-02", "Conduit EMT 3/4\" boucle AI", pieds("G-02"), "CONDUIT VIDE 3/4", "3.36",
          "Détail G-02 %s m converti en pi" % d("G-02", "q"))
cat_ligne("G-03", "Câble AI FT4 2c #16 (boucle adressable)", pieds("G-03"), "COND ALARME 3/16 FT4", "3.36",
          "Détail G-03 %s m converti en pi" % d("G-03", "q"))
lot_ligne("LOT-03", "Lot 2 détecteurs de fumée adressables (JCI)", LOTS["detecteurs_fumee"], "3.36",
          "METHODE.md « 2 détecteurs 600 $ », confirmé par Francis 2026-10-20")
cat_ligne("G-04", "Relocalisation panneau annonciateur Simplex", rel("PANNEAU ANNONCIATEUR", "INSTALLER"),
          "RELO PANN ANNOCIATEUR", "3.37", R + " PANNEAU ANNONCIATEUR (nouvel emplacement E-203) = Détail G-04")
jug_ligne("G-05", d("G-05", "desc"), d("G-05", "q"), d("G-05", "u"), d("G-05", "h"), "3.38", J + " G-05")
jug_ligne("G-06", d("G-06", "desc"), d("G-06", "q"), d("G-06", "u"), d("G-06", "h"), "3.38", J + " G-06")
lot_ligne("G-07", d("G-07", "desc"), LOTS["alarme_JCI"], "3.38",
          "METHODE.md lot JCI 7 400 $, confirmé par Francis 2026-10-20")

# ===== H — Frais de chantier
cat_ligne("H-01", "Mobilisation / démobilisation, plateforme élévatrice (nacelle)", 1, "NACELLE", "6.1",
          "Détail H-01 (1 lot)")
for c in ("H-02", "H-03", "H-04"):
    jug_ligne(c, d(c, "desc"), d(c, "q"), d(c, "u"), d(c, "h"), "6.1", J + " " + c)

# ---------------------------------------------------------------- totaux
codes = [l["code"] for l in LIGNES]
if len(codes) != len(set(codes)):
    erreur("codes de ligne en double")
for l in LIGNES:
    if not (isinstance(l["q"], (int, float)) and l["q"] > 0):
        erreur("quantité invalide ligne %s" % l["code"])
    l["p"] = round(float(l["p"]), 4)
    l["h"] = round(float(l["h"]), 4)

materiel = round(sum(l["q"] * l["p"] for l in LIGNES), 2)
heures_relevees = round(sum(l["q"] * l["h"] for l in LIGNES), 2)
heures_chargees = int(math.ceil(heures_relevees / 10.0 - 1e-9) * 10)
main_doeuvre = round(heures_chargees * TAUX, 2)
total_acceo = round(materiel + main_doeuvre, 2)
prix_vente = int(math.floor(total_acceo * FACTEUR / 1000.0 + 0.5) * 1000)

# ventilation par article : matériel + heures relevées; total au prorata des heures chargées
ratio = heures_chargees / heures_relevees
par_article = {}
for l in LIGNES:
    a = par_article.setdefault(l["art"], {"materiel": 0.0, "heures": 0.0, "total": 0.0})
    a["materiel"] += l["q"] * l["p"]
    a["heures"] += l["q"] * l["h"]
for art, a in par_article.items():
    a["total"] = round(a["materiel"] + a["heures"] * ratio * TAUX, 2)
    a["materiel"] = round(a["materiel"], 2)
    a["heures"] = round(a["heures"], 2)


def cle_art(a):
    return tuple(int(x) for x in a.split("."))


par_article = {k: par_article[k] for k in sorted(par_article, key=cle_art)}
lots_total = round(sum(LOTS.values()), 2)
lots_lignes = round(sum(l["p"] for l in LIGNES if l["code"].startswith("LOT-") or l["code"] in ("G-07", "D-17")), 2)

# ---------------------------------------------------------------- contrôles
assert abs(materiel + heures_chargees * TAUX - total_acceo) < 0.005, "cohérence total ACCEO"
assert prix_vente == int(round(total_acceo * FACTEUR / 1000.0)) * 1000, "arrondi prix de vente"
assert heures_chargees % 10 == 0 and heures_chargees >= heures_relevees, "heures chargées"
assert abs(lots_total - lots_lignes) < 0.005, "lots ≠ lignes de lot"
assert abs(sum(a["materiel"] for a in par_article.values()) - materiel) < 0.05, "ventilation matériel"
assert abs(sum(a["heures"] for a in par_article.values()) - heures_relevees) < 0.05, "ventilation heures"
assert abs(sum(a["total"] for a in par_article.values()) - total_acceo) < 0.5, "ventilation total"

sortie = {
    "taux": TAUX,
    "materiel": materiel,
    "heures_relevees": heures_relevees,
    "heures_chargees": heures_chargees,
    "main_doeuvre": main_doeuvre,
    "total_acceo": total_acceo,
    "facteur": FACTEUR,
    "prix_vente": prix_vente,
    "lignes": [{k: l[k] for k in ("code", "desc", "q", "u", "p", "h", "art", "src")} for l in LIGNES],
    "par_article": par_article,
    "lots": LOTS,
}
with open(P("chiffrage.json"), "w", encoding="utf-8") as f:
    json.dump(sortie, f, ensure_ascii=False, indent=2)

part_mo = main_doeuvre / total_acceo * 100
print("lignes=%d materiel=%.2f heures=%.2f chargees=%d MO=%.2f total=%.2f prix=%d | part MO %.1f %% | %.0f h/100k$"
      % (len(LIGNES), materiel, heures_relevees, heures_chargees, main_doeuvre, total_acceo, prix_vente,
         part_mo, heures_chargees / prix_vente * 100000))
