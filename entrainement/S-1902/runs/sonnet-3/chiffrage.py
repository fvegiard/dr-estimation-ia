#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » à l'aveugle — S-1902 Bibliothèque Georgette-Lepage (Brossard).

Déterministe : lit uniquement le kit (dossier du script), écrit chiffrage.json.
Dépendances : stdlib + openpyxl.

Sources lues :
  - releve.xlsx (feuille « Par libellé » + « Marques » pour contrôle) : quantités sur plan
  - rapport-releve.md : répartition par portée (CONSERVER / INSTALLER / ENLEVER / RELOCALISER)
  - entree/reference-Estimation_2.xlsx : feuille « Détail » (quantités hors plan, heures de jugement),
    feuille « Matériel Lumen » (prix réguliers pour les familles absentes du catalogue)
  - catalogue-prix-temps-dupuis.csv : colonnes prix_a_utiliser / temps_a_utiliser (dernier Dupuis)
  - METHODE.md : taux, facteur, lots confirmés (valeurs reprises ci-dessous, vérifiées à la lecture)
"""
import csv
import json
import math
import os
import re
import sys

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
M_VERS_PI = 3.28084

# --- Paramètres METHODE.md (vérifiés dans verifier_methode()) -------------------------------------
TAUX = 85.0          # $/h dossiers août-octobre 2026
FACTEUR = 1.244      # S-1844
LOTS = {             # prix de lot confirmés par Francis le 2026-10-20 (METHODE.md §2)
    "eclairage": 11893.0,
    "distribution": 6200.0,
    "jci": 7400.0,
    "detecteurs_jci": 600.0,
    "gpr": 1250.0,
}


def chemin(*p):
    return os.path.join(ICI, *p)


def verifier_methode():
    txt = open(chemin("METHODE.md"), encoding="utf-8").read()
    ok = (
        "85 $/h" in txt
        and "1,244" in txt
        and "11 893 $" in txt
        and "6 200 $" in txt
        and "7 400 $" in txt
        and "600 $" in txt
        and "1 250 $" in txt
    )
    if not ok:
        raise SystemExit("METHODE.md ne contient plus les paramètres attendus (taux/facteur/lots)")


# --- Lecture du kit -------------------------------------------------------------------------------
def lire_catalogue():
    cat = {}
    with open(chemin("catalogue-prix-temps-dupuis.csv"), encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            cat.setdefault(r["description"], []).append(r)
    return cat


def lire_releve():
    wb = openpyxl.load_workbook(chemin("releve.xlsx"), data_only=True)
    qte = {}
    for i, r in enumerate(wb["Par libellé"].iter_rows(values_only=True)):
        if i == 0:
            continue
        if r[0]:
            qte[r[0]] = int(r[2])
    # contrôle croisé avec la feuille Marques
    cnt = {}
    for i, r in enumerate(wb["Marques"].iter_rows(values_only=True)):
        if i == 0 or not r[0] or not r[2]:
            continue
        if r[0] == "Feuille (cartouche)":  # en-tête secondaire éventuel
            continue
        cnt[r[2]] = cnt.get(r[2], 0) + 1
    for k, v in qte.items():
        if cnt.get(k) != v:
            raise SystemExit("Incohérence releve.xlsx: %s Par libellé=%s Marques=%s" % (k, v, cnt.get(k)))
    return qte


def lire_portees(qte):
    """Retourne {famille: {portée: n}} depuis rapport-releve.md, contrôlé contre releve.xlsx."""
    por = {}
    for ln in open(chemin("rapport-releve.md"), encoding="utf-8"):
        m = re.match(r"\|\s*([A-Z0-9_ ]+?)\s*\|\s*(\d+)\s*\|\s*(.+?)\s*\|\s*$", ln)
        if not m:
            continue
        fam, tot, txt = m.group(1), int(m.group(2)), m.group(3)
        d = {}
        for part in txt.split(","):
            mm = re.match(r"\s*([A-Z]+)\s+(\d+)\s*$", part)
            if not mm:
                raise SystemExit("portée illisible: %r" % part)
            d[mm.group(1)] = int(mm.group(2))
        if sum(d.values()) != tot:
            raise SystemExit("rapport-releve.md: %s portées != total" % fam)
        por[fam] = d
    for fam, n in qte.items():
        if fam not in por:
            raise SystemExit("famille absente du rapport: " + fam)
        if sum(por[fam].values()) != n:
            raise SystemExit("famille %s: rapport %s != releve.xlsx %s" % (fam, sum(por[fam].values()), n))
    return por


def lire_detail():
    wb = openpyxl.load_workbook(chemin("entree", "reference-Estimation_2.xlsx"), data_only=True)
    det = {}
    for i, r in enumerate(wb["Détail"].iter_rows(values_only=True), 1):
        if i < 7 or not r[0]:
            continue
        code = str(r[0])
        if re.match(r"^[A-H]-\d\d$", code):
            det[code] = {"desc": r[1], "q": r[3], "u": r[4], "hpu": r[7]}
    return det, wb


def lire_lumen(wb):
    lum = {}
    for r in wb["Matériel Lumen"].iter_rows(values_only=True):
        if r[0] and r[4] is not None and isinstance(r[4], (int, float)):
            lum[str(r[0])] = {"desc": r[1], "p": float(r[4]), "u": r[3]}
    return lum


# --- Construction des lignes ------------------------------------------------------------------------
class Chiffrage:
    def __init__(self):
        verifier_methode()
        self.cat = lire_catalogue()
        self.qte = lire_releve()
        self.por = lire_portees(self.qte)
        self.det, wb_est = lire_detail()
        self.lum = lire_lumen(wb_est)
        self.lignes = []

    # quantité relevée par famille et portée
    def rq(self, fam, *portees):
        return sum(self.por[fam].get(p, 0) for p in portees)

    def dq(self, code):
        return self.det[code]["q"]

    def dm_vers_pi(self, code):
        d = self.det[code]
        if d["u"] != "m":
            raise SystemExit("unité inattendue pour " + code)
        return int(round(d["q"] * M_VERS_PI))

    def cat_ligne(self, label):
        rows = self.cat.get(label)
        if not rows or len(rows) != 1:
            raise SystemExit("libellé catalogue introuvable ou ambigu: %r" % label)
        r = rows[0]
        return float(r["prix_a_utiliser"]), float(r["temps_a_utiliser"]), r

    def add_cat(self, code, desc, q, u, label, art, p_force=None, note=""):
        p, h, r = self.cat_ligne(label)
        src = "catalogue Dupuis «%s» (PRIX À UTILISER %s $, TEMPS À UTILISER %s h, dernier %s)" % (
            label, r["prix_a_utiliser"], r["temps_a_utiliser"], r["date_dernier"])
        if p_force is not None:
            src += " — prix forcé à %s $ : %s" % (p_force, note)
            p = p_force
        elif note:
            src += " — " + note
        self.lignes.append({"code": code, "desc": desc, "q": q, "u": u, "p": p, "h": h, "art": art, "src": src})

    def add_jug(self, code, desc, q, u, h, art, why):
        self.lignes.append({"code": code, "desc": desc, "q": q, "u": u, "p": 0.0, "h": h, "art": art,
                            "src": "jugement (Estimation_2, feuille Détail %s : %s h/%s) — %s" % (code, h, u, why)})

    def add_lot(self, code, desc, montant, art, nom):
        self.lignes.append({"code": code, "desc": desc, "q": 1, "u": "lot", "p": montant, "h": 0.0, "art": art,
                            "src": "lot fournisseur confirmé METHODE.md (%s, prix confirmés par Francis 2026-10-20)" % nom})

    def construire(self):
        R = self.rq
        D = self.det

        # ---------- A. Démontage / relocalisation (3.22 et 5.19) ----------
        q = R("ENCASTRE", "ENLEVER") + R("LED 2X2", "ENLEVER") + R("LED 1X4", "ENLEVER")
        self.add_cat("A-01", "Démantèlement luminaires existants (E.A.E) RDC + 2e — ENCASTRE %d + LED 2X2 %d + LED 1X4 %d (relevé)" % (
            R("ENCASTRE", "ENLEVER"), R("LED 2X2", "ENLEVER"), R("LED 1X4", "ENLEVER")),
            q, "u", "EQUIPEMENT A ENLEVER", "3.22",
            note="temps 0,5 h = luminaire enlevé (METHODE §3); A-02 fusionné dans A-01 (relevé sans split par étage)")
        q = R("PRISE DDFT", "ENLEVER") + R("PRISE DOUBLE 15A", "ENLEVER") + R("SORTIE DATA", "ENLEVER")
        self.add_cat("A-03", "Démantèlement services (prises DDFT %d, prises doubles %d, sorties data %d) E.A.E (relevé); A-04 fusionné" % (
            R("PRISE DDFT", "ENLEVER"), R("PRISE DOUBLE 15A", "ENLEVER"), R("SORTIE DATA", "ENLEVER")),
            q, "u", "ENLEVER PRISE", "3.22", note="prise enlevée 0,35 h (METHODE §3)")
        self.add_cat("N-01", "Interrupteur simple enlevé (E.A.E, relevé)", R("INT SIMPLE", "ENLEVER"), "u",
                     "ENLEVER INTERRUPTEUR", "3.22")
        self.add_cat("N-02", "Monument de plancher enlevé (E.A.E, relevé)", R("MONUMENT PLANCHER", "ENLEVER"), "u",
                     "EQUIPEMENT A ENLEVER", "3.22")
        q = (R("BOITE JONCTION", "ENLEVER") + R("BOITIER C", "ENLEVER") + R("DETECTEUR FUMEE", "ENLEVER")
             + R("DETECTEUR MOUVEMENT", "ENLEVER") + R("HP ENCASTRE PLAFOND", "ENLEVER"))
        self.add_cat("N-03", "Autres appareils enlevés (boîtes de jonction, boîtier C, détecteurs fumée/mouvement, haut-parleur) E.A.E (relevé)",
                     q, "u", "EQUIPEMENT A ENLEVER", "3.22")

        q = R("PRISE DOUBLE 15A", "RELOCALISER") + R("PRISE DDFT", "RELOCALISER") + R("SORTIE DATA", "RELOCALISER")
        self.add_cat("A-05", "Relocalisation équipements services (E.A.R) : prises doubles %d, DDFT %d, data %d (relevé)" % (
            R("PRISE DOUBLE 15A", "RELOCALISER"), R("PRISE DDFT", "RELOCALISER"), R("SORTIE DATA", "RELOCALISER")),
            q, "u", "RELO PRISE", "5.19", note="relocalisation de prise 1,5 h (METHODE §3)")
        self.add_cat("A-06", "Relocalisation luminaires encastrés (E.A.R) (relevé)", R("ENCASTRE", "RELOCALISER"), "u",
                     "RELO ENCASTRE", "5.19")
        d = D["A-07"]
        self.add_jug("A-07", "Relocalisation interrupteurs maîtres du hall (note spécifique 1 E-201)", d["q"], d["u"], d["hpu"], "5.19",
                     "aucune famille catalogue équivalente à un lot d'interrupteurs maîtres, quantité inconnue")
        self.add_cat("A-08", "FabLab 210 : retrait temporaire + réinstallation luminaires linéaires existants (E-202 note 1)",
                     self.dq("A-08"), "u", "RELO ENCASTRE", "5.19",
                     note="quantité Détail (hors relevé); retrait+réinstallation assimilés à une relocalisation")
        self.add_cat("A-09", "FabLab 210 : retrait temporaire + réinstallation encastrés existants (E-202 note 1)",
                     self.dq("A-09"), "u", "RELO ENCASTRE", "5.19", note="quantité Détail (hors relevé)")
        n_mono_r = R("MOTEUR MONOPHASE", "RELOCALISER")
        if n_mono_r != self.dq("A-10"):
            raise SystemExit("A-10: relevé %s != Détail %s" % (n_mono_r, self.dq("A-10")))
        self.add_cat("A-10", "Toit : débrancher équipements de ventilation (E.A.R, E-103)", n_mono_r, "u",
                     "DEBRANCHEMENT APPAREILAGE MECANIQUE", "5.19")
        self.add_cat("N-04", "Toit : rebrancher équipements de ventilation relocalisés (E.A.R, E-103)", n_mono_r, "u",
                     "MOTEUR", "5.19", note="moteur 0,75 h (METHODE §3) pour le rebranchement")
        d = D["A-11"]
        self.add_jug("A-11", d["desc"], d["q"], d["u"], d["hpu"], "5.19", "traçage/continuité de service, pas de famille catalogue")

        # ---------- B. Distribution (3.24 / 3.25) ----------
        self.add_cat("B-01", "Panneau PP-3 neuf 225 A 347/600 V 42 cct NF Schneider — remplacement (E-006/E-502)", 1, "u",
                     "PANNEAU 100A 347/600V 3F", "3.24",
                     p_force=0.0, note="panneau 8 h (METHODE §3); «PANNEAU 225A 347/600V 3» a temps 0 au catalogue; prix dans lot distribution")
        n1p = D["B-02"]["q"]
        self.add_cat("B-02", D["B-02"]["desc"], n1p, "u", "DISJ15A1P", "3.25", note="prix dans lot distribution (EDB 1P)")
        for code in ("B-03", "B-04", "B-05"):
            self.add_cat(code, D[code]["desc"], D[code]["q"], "u", "DISJ 100A 3P 120/240V 22 K", "3.25",
                         p_force=0.0, note="disjoncteur 3P 1 h (METHODE §3); prix dans lot distribution (EJB 3P)")
        self.add_cat("B-06", D["B-06"]["desc"], D["B-06"]["q"], "u", "DISJ15A1P", "3.25", note="prix dans lot distribution")
        self.add_cat("B-07", D["B-07"]["desc"], D["B-07"]["q"], "u", "DISJ15A1P", "3.25", note="prix dans lot distribution")
        self.add_cat("B-08", D["B-08"]["desc"], D["B-08"]["q"], "u", "DISJ 20A 1P 120V 10KA", "3.25",
                     p_force=0.0, note="prix dans lot distribution")
        d = D["B-09"]
        self.add_jug("B-09", d["desc"], d["q"], d["u"], d["hpu"], "3.24", "mise à jour des fiches de panneaux, pas de famille catalogue")
        self.add_lot("LOT-DIST-PP3", "Lot distribution : panneau NF 225 A 42 cct", 3400.0, "3.24", "lot distribution")
        self.add_lot("LOT-DIST-DISJ", "Lot distribution : disjoncteurs EDB 1P x20 (95 $) + EJB 3P x3 (300 $)", 2800.0, "3.25", "lot distribution")

        # ---------- C. Raccordements mécaniques (3.23 / 3.26) ----------
        n_tri_i = R("MOTEUR TRIPHASE", "INSTALLER")
        n_mono_i = R("MOTEUR MONOPHASE", "INSTALLER")
        if n_tri_i != self.dq("C-01") + self.dq("C-02") + self.dq("C-03") or n_mono_i != self.dq("C-04"):
            raise SystemExit("C: moteurs relevé != Détail")
        self.add_cat("C-01", "UTA-01 York 575 V/3Ø au toit — raccordement (relevé: moteur triphasé)", self.dq("C-01"), "u",
                     "MOTEUR UAF-1", "3.23", note="UTA 1 h (METHODE §3)")
        self.add_cat("C-02", "VE-02N Plymovent 1,5 HP 575 V/3Ø au toit — raccordement (relevé: moteur triphasé)", self.dq("C-02"), "u",
                     "MOTEUR", "3.23", note="moteur 0,75 h (METHODE §3)")
        self.add_cat("C-03", "Chaudières CH-01/CH-02 600 V/3Ø — tiroirs CCM-1 F15A, câbles existants (relevé: moteur triphasé)", self.dq("C-03"), "u",
                     "MOTEUR", "3.26", note="moteur 0,75 h (METHODE §3)")
        self.add_cat("C-04", "Humidificateurs HUM-1/2/3 Carel 120 V (relevé: moteur monophasé)", n_mono_i, "u", "MOTEUR", "3.23",
                     note="moteur 0,75 h (METHODE §3)")
        self.add_cat("N-06", "Sectionneur 30 A 1Ø HUM-1/2/3 (C-04)", n_mono_i, "u", "SECTIONNEUR POR ENTRETIEN HUBHBL13R22", "3.23",
                     note="Hubbell, seul sectionneur 30 A avec prix au catalogue")
        self.add_cat("C-05", "Conduit EMT 1\" + 3 x #8 RW90 (alim. UTA-01) — %s m (Détail)" % self.dq("C-05"),
                     self.dm_vers_pi("C-05"), "pi", "CONDUIT 1 03C08", "3.23",
                     note="assemblage conduit+3 conducteurs; C-06 réduit au conducteur MALT; longueur = hypothèse Détail")
        self.add_cat("C-06", "Conducteur MALT #8 RW90 (UTA-01), 35 m", self.dm_vers_pi("C-05"), "pi", "8AWG 7 BRINS RW90 600V ", "3.23",
                     note="4e conducteur (3+MALT)")
        m_ve = 35
        m_hum = int(self.dq("C-07")) - m_ve
        self.add_cat("C-07", "Conduit EMT 3/4\" + 4 x #12 (VE-02N), 35 m", int(round(m_ve * M_VERS_PI)), "pi", "CONDUIT 3/4 04C12", "3.23",
                     note="longueur = hypothèse Détail; C-08 fusionné (fils inclus)")
        self.add_cat("N-08", "Conduit EMT 3/4\" + 3 x #12 (HUM-1/2/3), 3 x 15 m", int(round(m_hum * M_VERS_PI)), "pi", "CONDUIT 3/4 03C12", "3.23",
                     note="longueur = hypothèse Détail")

        # ---------- D/E. Services (3.27 à 3.34, 3.31, 3.33) ----------
        self.add_cat("D-06", "Prises doubles 15 A-120 V neuves (relevé INSTALLER)", R("PRISE DOUBLE 15A", "INSTALLER"), "u",
                     "PRISE DUPLEX 15A 120V", "3.28", note="prise duplex assemblage 64,81 $ (METHODE §1)")
        self.add_cat("D-07", "Prises doubles DDFT/GFI neuves (relevé INSTALLER; MALT isolées non distinguées)", R("PRISE DDFT", "INSTALLER"), "u",
                     "PRISE DUPLEX 20A 125V GFI", "3.30", note="DDFT 65,38 $ / 1,3 h (METHODE §1,§3)")
        self.add_cat("D-02", "Prises doubles USB Leviton Decora neuves (relevé INSTALLER)", R("PRISE USB", "INSTALLER"), "u",
                     "PRISE USB 125V 15 AMP", "3.27", note="USB 77,45 $ / 1,125 h (METHODE §1,§3)")
        n_data = R("SORTIE DATA", "INSTALLER")
        n_data_cond = self.dq("D-08")
        self.add_cat("D-08", "Sortie data murale/colonne : boîte + conduit 3/4\" vers entreplafond (D-08 Détail)", n_data_cond, "u",
                     "TEL/DATA 3/4", "3.33", note="tel/data 3/4 22,32 $ (METHODE §1)")
        self.add_cat("N-09", "Sorties data intégrées au mobilier (plaque seulement) — relevé INSTALLER %d moins D-08" % n_data,
                     n_data - n_data_cond, "u", "TEL/DATA", "3.33")
        self.add_cat("E-02", "Monument de plancher Hubbell S1R6PTFIT (relevé INSTALLER)", R("MONUMENT PLANCHER", "INSTALLER"), "u",
                     "MONUMENT PLANCHER", "3.34", note="monument 250 $ / 1 h (METHODE §1,§3)")
        self.add_cat("N-10", "Boîtes de jonction neuves (relevé INSTALLER)", R("BOITE JONCTION", "INSTALLER"), "u",
                     "BOITE CARRE 2-1/8PO KO 1", "3.31")
        d = D["D-09"]
        self.add_jug("D-09", d["desc"], d["q"], d["u"], d["hpu"], "3.33", "raccord panneau de contrôle barrière anti-vol, système fourni par d'autres")
        self.add_cat("N-11", "Lecteurs de carte neufs (relevé INSTALLER)", R("LECTEUR CARTE", "INSTALLER"), "u",
                     "LECTEUR DE CARTE", "3.33", note="lecteur de carte 15 $ (METHODE §1)")
        self.add_cat("D-11", D["D-11"]["desc"], D["D-11"]["q"], "u", "DECLENCHEUR APPEL", "3.33")
        self.add_cat("E-06", D["E-06"]["desc"], D["E-06"]["q"], "u", "DECLENCHEUR APPEL", "3.33")
        self.add_cat("D-12", "Percement de dalle carotté — conduits sous dalle RDC (note 5)", D["D-12"]["q"], "u",
                     "PERCEMENT PLANCHER", "3.31", note="percement 150 $ / 1 h (METHODE §1,§3)")
        self.add_cat("E-07", "Percement de dalle carotté — conduits sous dalle étage (note 3)", D["E-07"]["q"], "u",
                     "PERCEMENT PLANCHER", "3.31", note="percement 150 $ / 1 h")
        self.add_cat("N-12", "Coupe-feu CAN/ULC-S115 aux percements (1 par percement)", D["D-12"]["q"] + D["E-07"]["q"], "u",
                     "COUPE FEU TUBE 10.1 FL O", "3.31", note="0,5 h + percement 1 h = 1,5 h du Détail")
        self.add_cat("D-13", "Conduit EMT 1\" sous dalle RDC (16 courses), %s m" % D["D-13"]["q"], self.dm_vers_pi("D-13"), "pi",
                     "COND TME 1PO", "3.31", note="longueur = hypothèse Détail (30 m/course)")
        self.add_cat("D-14", "Conduit vide 3/4\" data sous dalle RDC (11 courses), %s m" % D["D-14"]["q"], self.dm_vers_pi("D-14"), "pi",
                     "CONDUIT VIDE 3/4", "3.31", note="conduit vide 3/4 1,41 $/pi, 0,056 h/pi (METHODE §1,§3); hypothèse Détail")
        self.add_cat("D-15", "Conducteurs #12 RW90 (5 cond./course x 16 courses x 30 m), %s m" % D["D-15"]["q"], self.dm_vers_pi("D-15"), "pi",
                     "12AWG SOLIDE RW90 600V", "3.31", note="hypothèse Détail")
        self.add_cat("D-16", "Boîtes de tirage 4\" carrées sous dalle RDC (2 par course)", 2 * D["D-12"]["q"], "u",
                     "BOITE CARREES 4 PO", "3.31", note="Lumen IBE72171K : 2 par course")
        self.add_cat("E-08", "Conduit EMT 1\" sous dalle étage (7 courses), %s m" % D["E-08"]["q"], self.dm_vers_pi("E-08"), "pi",
                     "COND TME 1PO", "3.31", note="longueur = hypothèse Détail (25 m/course)")
        self.add_cat("E-09", "Conduit vide 3/4\" data sous dalle étage (7 courses), %s m" % D["E-09"]["q"], self.dm_vers_pi("E-09"), "pi",
                     "CONDUIT VIDE 3/4", "3.31", note="hypothèse Détail")
        self.add_cat("E-10", "Conducteurs #12 RW90 étage (5 cond. x 7 x 25 m), %s m" % D["E-10"]["q"], self.dm_vers_pi("E-10"), "pi",
                     "12AWG SOLIDE RW90 600V", "3.31", note="hypothèse Détail")
        self.add_cat("E-11", "Boîtes de tirage 4\" carrées sous dalle étage (2 par course)", 2 * D["E-07"]["q"], "u",
                     "BOITE CARREES 4 PO", "3.31")
        # famille absente du catalogue : corde de tirage -> Lumen + jugement
        lum = self.lum["GRE430"]
        self.lignes.append({"code": "N-14", "desc": "Corde de tirage Poly Line 430 (6500 pi) — conduits vides", "q": 1, "u": "ch",
                            "p": lum["p"], "h": 0.5, "art": "3.31",
                            "src": "famille absente du catalogue Dupuis : prix Lumen GRE430 %s $ (prix régulier 2026-10-07) + temps par jugement 0,5 h" % lum["p"]})

        # ---------- F. Éclairage (3.35) ----------
        f = {"A1": ("LUMINAIRE A1", "FIXTURE TYPE A1"), "C1": ("LUMINAIRE C1", "FIXTURE TYPE C1"),
             "C1_1": ("LUMINAIRE C1_1", "FIXTURE TYPE C1"), "C1_3": ("LUMINAIRE C1_3", "FIXTURE TYPE C1"),
             "E": ("LUMINAIRE E", "FIXTURE TYPE E")}
        codes = {"A1": "F-01", "C1": "F-02", "C1_1": "F-03", "C1_3": "F-04", "E": "F-06"}
        for t in ("A1", "C1", "C1_1", "C1_3"):
            self.add_cat(codes[t], D[codes[t]]["desc"], R(f[t][0], "INSTALLER"), "u", f[t][1], "3.35",
                         p_force=0.0, note="luminaire 0,75 h (METHODE §3); prix dans lot éclairage")
        d = D["F-05"]
        self.add_jug("F-05", d["desc"], d["q"], d["u"], d["hpu"], "3.35", "suspension de rangées Axis (trousses dans le lot éclairage), pas de famille catalogue")
        self.add_cat("F-06", "Luminaire E — GreenCreative PXCYL-8 cylindre suspendu (relevé : 3 étiquettes; Détail : 4)",
                     R("LUMINAIRE E", "INSTALLER"), "u", "FIXTURE TYPE E", "3.35", p_force=0.0,
                     note="luminaire 0,75 h; prix dans lot éclairage (E x3)")
        self.add_cat("F-07", D["F-07"]["desc"], R("ENSEIGNE SORTIE Z6", "INSTALLER"), "u", "ENSEIGNE SORTIE Z6", "3.35",
                     note="enseigne 1 h (METHODE §3); prix dans lot éclairage (Z6 x3)")
        self.add_cat("F-08", "Conduit EMT 3/4\" + 3 x #12 RW90 raccordement luminaires, %s m (35 points x 8 m)" % D["F-08"]["q"], self.dm_vers_pi("F-08"), "pi",
                     "CONDUIT 3/4 03C12", "3.35", note="conduit 3/4 03C12 4,24 $/pi, 0,113 h/pi (METHODE §1,§3); F-09 fusionné (fils inclus); hypothèse Détail")
        self.add_cat("F-10", "Boîtes octogonales de jonction luminaires (1 par point, 35 points)", int(round(D["F-08"]["q"] / 8.0)), "u",
                     "BOITE OCTOG 1-1/2PO KO 1", "3.35", note="Lumen IBE54151L x35")
        self.add_lot("LOT-ECL", "Lot éclairage : A1 x11, Axis SL1 4'/3'/1', trousses, E x3, Z6 x3", LOTS["eclairage"], "3.35", "lot éclairage")

        # ---------- I. Interrupteurs, thermostats ----------
        self.add_cat("N-16", "Interrupteurs simples neufs (relevé INSTALLER)", R("INT SIMPLE", "INSTALLER"), "u",
                     "INTERRUPTEUR 15A 120V", "3.35", note="interrupteur 75,47 $ (METHODE §1)")
        self.add_cat("N-13", "Thermostats neufs (relevé INSTALLER)", R("THERMOSTAT", "INSTALLER"), "u", "THERMOSTAT", "3.23",
                     note="thermostat 35 $ (METHODE §1)")

        # ---------- G. Alarme incendie (3.36 à 3.38) ----------
        self.add_cat("G-01", "Détecteurs de fumée adressables Simplex neufs (relevé INSTALLER)", R("DETECTEUR FUMEE", "INSTALLER"), "u",
                     "DETECTEUR FUMEE", "3.36", note="détecteur 0,5 h (METHODE §3); prix dans lot JCI (2 détecteurs 600 $)")
        self.add_cat("N-05", "Détecteur de fumée relocalisé (relevé RELOCALISER)", R("DETECTEUR FUMEE", "RELOCALISER"), "u",
                     "RELO DETECTEUR A FUME", "3.36")
        self.add_cat("G-02", "Conduit vide 3/4\" boucle AI, %s m" % D["G-02"]["q"], self.dm_vers_pi("G-02"), "pi",
                     "CONDUIT VIDE 3/4", "3.36", note="hypothèse Détail")
        self.add_cat("G-03", "Câble AI FT4 2c #16 (boucle adressable), %s m" % D["G-03"]["q"], self.dm_vers_pi("G-03"), "pi",
                     "COND ALARME 3/16 FT4", "3.36", note="hypothèse Détail")
        self.add_lot("LOT-DET", "Lot JCI : 2 détecteurs de fumée", LOTS["detecteurs_jci"], "3.36", "détecteurs JCI")
        self.add_cat("G-04", "Relocalisation panneau annonciateur Simplex (relevé PANNEAU ANNONCIATEUR)", R("PANNEAU ANNONCIATEUR", "INSTALLER"), "u",
                     "RELO PANN ANNOCIATEUR", "3.37", note="annonciateur relocalisé 8 h (METHODE §3)")
        d = D["G-05"]
        self.add_jug("G-05", d["desc"], d["q"], d["u"], d["hpu"], "3.37", "allocation, dispositifs AI existants en zones touchées")
        d = D["G-06"]
        self.add_jug("G-06", d["desc"], d["q"], d["u"], d["hpu"], "3.38", "assistance à JCI, pas de famille catalogue")
        self.add_lot("G-07", "Sous-traitance JCI/Simplex : bypass, programmation 4007ES, vérification S537, certificat", LOTS["jci"], "3.38", "JCI")

        # ---------- Scan dalle (3.32) ----------
        self.add_cat("N-15", "Scan dalle béton (GPR) avant percements", 1, "u", "SCAN DALLE BETON", "3.32", p_force=0.0,
                     note="prix 1 250 $ porté par le lot GPR")
        self.add_lot("LOT-GPR", "Lot GPR / scan dalle", LOTS["gpr"], "3.32", "GPR")

        # ---------- H. Frais de chantier (6.1) ----------
        self.add_cat("H-01", "Mobilisation / nacelle", 1, "lot", "NACELLE", "6.1", note="nacelle 1 250 $ / 16 h (METHODE §3)")
        for c, why in (("H-02", "coordination, dessins d'atelier, tel que construit"),
                       ("H-03", "essais, mise en service, étiquetage"),
                       ("H-04", "nettoyage, gestion des déchets")):
            d = D[c]
            self.add_jug(c, d["desc"], d["q"], d["u"], d["hpu"], "6.1", why)

    # --- Totaux --------------------------------------------------------------------------------------
    def totaux(self):
        mat = sum(l["q"] * l["p"] for l in self.lignes)
        hrs = sum(l["q"] * l["h"] for l in self.lignes)
        mat = round(mat, 2)
        hrs = round(hrs, 4)
        charg = int(math.ceil(round(hrs / 10.0, 9)) * 10)
        mo = round(charg * TAUX, 2)
        total = round(mat + mo, 2)
        pv = int(math.floor(total * FACTEUR / 1000.0 + 0.5) * 1000)
        # ventilation par article
        art = {}
        for l in self.lignes:
            a = art.setdefault(l["art"], {"materiel": 0.0, "heures": 0.0})
            a["materiel"] += l["q"] * l["p"]
            a["heures"] += l["q"] * l["h"]
        arts = sorted(art, key=lambda s: [int(x) for x in s.split(".")])
        par_art = {}
        cumul = 0.0
        ratio = charg / hrs if hrs else 0.0
        for k in arts:
            m = round(art[k]["materiel"], 2)
            h = round(art[k]["heures"], 4)
            t = round(m + h * ratio * TAUX, 2)  # heures chargées réparties au prorata des heures relevées
            par_art[k] = {"materiel": m, "heures": h, "total": t}
            cumul += t
        ecart = round(total - cumul, 2)
        if ecart:
            gros = max(par_art, key=lambda k: par_art[k]["total"])
            par_art[gros]["total"] = round(par_art[gros]["total"] + ecart, 2)
        return mat, hrs, charg, mo, total, pv, par_art

    def ecrire(self):
        mat, hrs, charg, mo, total, pv, par_art = self.totaux()
        lignes = []
        for l in self.lignes:
            lignes.append({"code": l["code"], "desc": l["desc"], "q": l["q"], "u": l["u"],
                           "p": round(l["p"], 4), "h": round(l["h"], 4), "art": l["art"], "src": l["src"]})
        out = {
            "taux": TAUX,
            "materiel": mat,
            "heures_relevees": hrs,
            "heures_chargees": charg,
            "main_doeuvre": mo,
            "total_acceo": total,
            "facteur": FACTEUR,
            "prix_vente": pv,
            "lignes": lignes,
            "par_article": par_art,
            "lots": LOTS,
        }
        # contrôles de cohérence
        assert abs(out["materiel"] + out["heures_chargees"] * out["taux"] - out["total_acceo"]) < 0.005
        assert out["prix_vente"] == int(math.floor(out["total_acceo"] * out["facteur"] / 1000.0 + 0.5) * 1000)
        assert out["heures_chargees"] % 10 == 0 and out["heures_chargees"] >= out["heures_relevees"]
        assert out["heures_chargees"] - out["heures_relevees"] < 10
        assert abs(sum(a["materiel"] for a in par_art.values()) - out["materiel"]) < 0.01
        assert abs(sum(a["heures"] for a in par_art.values()) - out["heures_relevees"]) < 0.01
        assert abs(sum(a["total"] for a in par_art.values()) - out["total_acceo"]) < 0.01
        lots_en_lignes = round(sum(l["p"] for l in self.lignes if l["u"] == "lot" and l["src"].startswith("lot fournisseur")), 2)
        assert abs(lots_en_lignes - sum(LOTS.values())) < 0.01
        codes = [l["code"] for l in lignes]
        assert len(codes) == len(set(codes)), "codes de ligne dupliqués"
        with open(chemin("chiffrage.json"), "w", encoding="utf-8") as fh:
            json.dump(out, fh, ensure_ascii=False, indent=2)
        return out


def main():
    c = Chiffrage()
    c.construire()
    out = c.ecrire()
    print("lignes          :", len(out["lignes"]))
    print("taux            :", out["taux"])
    print("materiel        :", out["materiel"])
    print("heures_relevees :", out["heures_relevees"])
    print("heures_chargees :", out["heures_chargees"])
    print("main_doeuvre    :", out["main_doeuvre"])
    print("total_acceo     :", out["total_acceo"])
    print("facteur         :", out["facteur"])
    print("prix_vente      :", out["prix_vente"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
