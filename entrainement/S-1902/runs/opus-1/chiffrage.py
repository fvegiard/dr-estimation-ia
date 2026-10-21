#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » à l'aveugle — dossier S-1902 (Bibliothèque Georgette-Lepage, Brossard).

Script autonome et déterministe (stdlib + openpyxl). Lit le kit du dossier de travail et écrit chiffrage.json.
Règles appliquées (CONSIGNES.md / METHODE.md) :
  - Matériel = Σ qté × PRIX À UTILISER Dupuis (dernière valeur, jamais la médiane) ; familles absentes → prix Lumen.
  - Luminaires, enseignes, détecteurs, panneaux, disjoncteurs du lot : 0 $ en ligne + lots fournisseur confirmés.
  - Heures = Σ qté × TEMPS À UTILISER Dupuis ; conduits : pas de colonne TEMPSEC au catalogue → temps/pi.
  - Heures chargées = heures relevées arrondies vers le haut au multiple de 10.
  - Total ACCEO = matériel + heures chargées × taux ; prix de vente = arrondi au millier de total ACCEO × facteur.
Dans `lignes`, p = prix unitaire ($/u) et h = temps unitaire (h/u) ; les totaux de ligne sont q×p et q×h.
"""
import csv
import json
import math
import os
import re
import sys

import openpyxl

ICI = os.path.dirname(os.path.abspath(__file__))
F_RELEVE = os.path.join(ICI, "releve.xlsx")
F_RAPPORT = os.path.join(ICI, "rapport-releve.md")
F_CATALOGUE = os.path.join(ICI, "catalogue-prix-temps-dupuis.csv")
F_ESTIM = os.path.join(ICI, "entree", "reference-Estimation_2.xlsx")
F_BORDEREAU = os.path.join(ICI, "entree", "reference-Bordereau_de_prix.xlsx")
F_SORTIE = os.path.join(ICI, "chiffrage.json")

# METHODE.md §2 points 2, 5, 7
TAUX = 85.0          # $/h — dossiers Dupuis août–octobre 2026 (S-1797, S-1835, S-1854, S-1881)
FACTEUR = 1.244      # S-1844 : sommaire 191 909,19 $ → prix remis 238 744 $
PI_PAR_M = 3.28084

# Lots fournisseur confirmés par Francis le 2026-10-20 (METHODE.md §2 point 2) et leur ventilation au bordereau
LOTS = {
    "Lot éclairage (A1, Axis SL1 4'/3'/1', trousses, E, Z6)": (11893.0, {"3.35": 11893.0}),
    "Lot distribution (NF 225 A 42 cct + EDB 1P ×20 + EJB 3P ×3)": (6200.0, {"3.24": 3400.0, "3.25": 2800.0}),
    "JCI / Simplex (bypass, programmation, vérification ULC-S537, certificat)": (7400.0, {"3.38": 7400.0}),
    "Détecteurs de fumée adressables ×2 (JCI)": (600.0, {"3.36": 600.0}),
    "GPR — scan de dalle": (1250.0, {"3.32": 1250.0}),
}


def fail(msg):
    print("ERREUR : " + msg, file=sys.stderr)
    sys.exit(1)


# ---------------------------------------------------------------- lecture du kit
def lire_catalogue():
    cat = {}
    with open(F_CATALOGUE, encoding="utf-8-sig", newline="") as fh:
        for r in csv.DictReader(fh, delimiter=";"):
            d = r["description"].strip()
            cat.setdefault(d, []).append((r["um"], float(r["prix_a_utiliser"] or 0), float(r["temps_a_utiliser"] or 0)))
    return cat


def lire_releve():
    wb = openpyxl.load_workbook(F_RELEVE, read_only=True, data_only=True)
    ws = wb["Par libellé"]
    tot = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            tot[str(row[0]).strip()] = int(row[2])
    # contrôle croisé avec la feuille Marques (une ligne = une occurrence)
    wm = wb["Marques"]
    cnt = {}
    for row in wm.iter_rows(min_row=2, values_only=True):
        if row[2]:
            cnt[row[2]] = cnt.get(row[2], 0) + 1
    if cnt != tot:
        fail("Par libellé ≠ Marques")
    wb.close()
    # portées par famille : tableau du rapport-releve.md
    portees = {}
    with open(F_RAPPORT, encoding="utf-8") as fh:
        for ligne in fh:
            m = re.match(r"^\|\s*([A-Z0-9_ ]+?)\s*\|\s*(\d+)\s*\|\s*([A-Z ,0-9]+?)\s*\|\s*$", ligne)
            if m and m.group(1) in tot:
                d = {}
                for part in m.group(3).split(","):
                    k, v = part.strip().rsplit(" ", 1)
                    d[k] = int(v)
                if sum(d.values()) != int(m.group(2)) or int(m.group(2)) != tot[m.group(1)]:
                    fail("portées incohérentes pour " + m.group(1))
                portees[m.group(1)] = d
    if set(portees) != set(tot):
        fail("familles sans portée : %s" % (set(tot) - set(portees)))
    return tot, portees


def lire_estimation():
    wb = openpyxl.load_workbook(F_ESTIM, read_only=True, data_only=True)
    det = {}
    for row in wb["Détail"].iter_rows(min_row=6, values_only=True):
        code = row[0]
        if isinstance(code, str) and re.match(r"^[A-H]-\d\d$", code):
            det[code] = {"desc": row[1], "q": row[3], "u": row[4], "h": row[7]}
    lumen = {}
    for row in wb["Matériel Lumen"].iter_rows(min_row=5, values_only=True):
        if row[0] and row[4] is not None and isinstance(row[4], (int, float)):
            lumen[str(row[0]).strip()] = {"desc": row[1], "q": row[2], "u": row[3], "p": float(row[4])}
    wb.close()
    return det, lumen


def lire_articles_bordereau():
    wb = openpyxl.load_workbook(F_BORDEREAU, read_only=True, data_only=True)
    arts = set()
    for row in wb.active.iter_rows(values_only=True):
        if row[0] and re.match(r"^\d+\.\d+$", str(row[0])):
            arts.add(str(row[0]))
    wb.close()
    return arts


CAT = lire_catalogue()
TOT, PORT = lire_releve()
DET, LUMEN = lire_estimation()
ARTICLES = lire_articles_bordereau()


def por(fam, portee):
    return PORT[fam].get(portee, 0)


def pieds(m):
    """mètres (hypothèse Estimation_2) → pieds, arrondi vers le haut à la section de 10 pi"""
    return int(math.ceil(m * PI_PAR_M / 10.0) * 10)


# ---------------------------------------------------------------- construction des lignes
LIGNES = []


def cat_ligne(code, desc, q, libelle, art, src_q, prix_zero=False):
    if libelle not in CAT:
        fail("libellé absent du catalogue : " + libelle)
    if len(CAT[libelle]) != 1:
        fail("libellé catalogue ambigu (plusieurs unités) : " + libelle)
    um, p, h = CAT[libelle][0]
    u = {"U": "u", "P": "pi", "M": "m"}[um]
    src = 'Dupuis « %s » (PRIX/TEMPS À UTILISER %.4g $ / %.4g h) — qté : %s' % (libelle, p, h, src_q)
    if prix_zero:
        src += " — prix 0 $ en ligne (fourni au lot)"
        p = 0.0
    LIGNES.append({"code": code, "desc": desc, "q": q, "u": u, "p": round(p, 4), "h": h, "art": art, "src": src})


def lumen_ligne(code, desc, q, code_lumen, h, art, src_q, src_h):
    if code_lumen not in LUMEN:
        fail("code Lumen absent : " + code_lumen)
    L = LUMEN[code_lumen]
    LIGNES.append({"code": code, "desc": desc, "q": q, "u": L["u"], "p": L["p"], "h": h, "art": art,
                   "src": "Lumen %s « %s » %.4g $/%s (famille absente du catalogue Dupuis) — temps : %s — qté : %s"
                          % (code_lumen, L["desc"], L["p"], L["u"], src_h, src_q)})


def jug_ligne(code, art, desc=None):
    D = DET[code]
    LIGNES.append({"code": code, "desc": desc or D["desc"], "q": D["q"], "u": D["u"], "p": 0.0, "h": float(D["h"]),
                   "art": art, "src": "jugement (Estimation_2 %s : %s %s × %s h, à valider)" % (code, D["q"], D["u"], D["h"])})


R = "relevé"
# ===== Prises / services (E-104, E-105 : INSTALLER)
cat_ligne("R-PU", "Prise double 15 A USB Leviton Decora (installer)", por("PRISE USB", "INSTALLER"),
          "PRISE USB 125V 15 AMP", "3.27", R + " PRISE USB INSTALLER")
cat_ligne("R-PC", "Prise double 15 A 120 V (installer)", por("PRISE DOUBLE 15A", "INSTALLER"),
          "PRISE DUPLEX 15A 120V", "3.28", R + " PRISE DOUBLE 15A INSTALLER")
n_ddft = por("PRISE DDFT", "INSTALLER")
q_gfi = int(DET["D-07"]["q"])  # 1 DDFT (fontaine P2-61) ; le reste = M.A.L.T. isolée (R-007, bordereau 3.29)
cat_ligne("R-PI", "Prise double 15 A 120 V M.A.L.T. isolée — postes (symbole plein, R-007)", n_ddft - q_gfi,
          "PRISE DUPLEX 15A 120V", "3.29", R + " PRISE DDFT INSTALLER %d − 1 DDFT fontaine (D-07)" % n_ddft)
cat_ligne("R-PG", "Prise DDFT 20 A — fontaine P2-61", q_gfi, "PRISE DUPLEX 20A 125V GFI", "3.30",
          R + " PRISE DDFT INSTALLER (part DDFT = Estimation_2 D-07)")
n_data = por("SORTIE DATA", "INSTALLER")
q_data_mur = int(DET["D-08"]["q"])
cat_ligne("R-SD3", "Sortie data murale/colonne : boîte + conduit 3/4 vers entreplafond", q_data_mur,
          "TEL/DATA 3/4", "3.33", "Estimation_2 D-08 (part murale des %d sorties INSTALLER du relevé)" % n_data)
cat_ligne("R-SD", "Sortie informatique (postes/mobilier, plaque + boîte)", n_data - q_data_mur,
          "TEL/DATA", "3.33", R + " SORTIE DATA INSTALLER %d − %d murales" % (n_data, q_data_mur))
cat_ligne("R-LC", "Lecteur de carte (installer)", por("LECTEUR CARTE", "INSTALLER"),
          "LECTEUR DE CARTE", "3.33", R + " LECTEUR CARTE INSTALLER")
cat_ligne("R-MP", "Monument plancher Hubbell S1R6PTFIT / poke-through", por("MONUMENT PLANCHER", "INSTALLER"),
          "MONUMENT PLANCHER", "3.34", R + " MONUMENT PLANCHER INSTALLER")
q_jb = por("BOITE JONCTION", "INSTALLER")
cat_ligne("R-JB", "Boîte de jonction (barrière anti-vol, bornes de prêt)", q_jb,
          "BOITE CARRE 2-1/8PO KO 1", "3.31", R + " BOITE JONCTION INSTALLER")
cat_ligne("R-JBc", "Couvercle de boîte de jonction", q_jb, "COUV CARRE 4PO VIERGE", "3.31", R + " BOITE JONCTION INSTALLER")
cat_ligne("D-09", "Barrière anti-vol : raccord au panneau de contrôle (PU-1 23)", int(DET["D-09"]["q"]),
          "PANNEAU CONTROL", "3.31", "Estimation_2 D-09 (boîtes comptées au relevé BOITE JONCTION)")
q_bout = int(DET["D-11"]["q"]) + int(DET["E-06"]["q"])
cat_ligne("D-11/E-06", "Boutons d'appel (3) + bouton de panique (1) : boîte + conduit vide 3/4", q_bout,
          "TEL/DATA 3/4", "3.31", "Estimation_2 D-11 + E-06")
q_perc = int(DET["D-12"]["q"]) + int(DET["E-07"]["q"])
cat_ligne("D-12/E-07", "Percement de dalle carotté (16 RDC + 7 étage)", q_perc, "PERCEMENT PLANCHER", "3.31",
          "Estimation_2 D-12 + E-07")
cat_ligne("D-12f", "Coupe-feu des percements (CAN/ULC-S115)", q_perc, "COUPE FEU TUBE 10.1 FL O", "3.31",
          "1 tube par percement (D-12 + E-07)")
cat_ligne("D-13", "Conduit EMT sous dalle RDC + 5 #12 (16 courses × 30 m)", pieds(DET["D-13"]["q"]),
          "CONDUIT 3/4 05C12", "3.31", "Estimation_2 D-13 %s m → pi (D-15 #12 inclus dans l'assemblage)" % DET["D-13"]["q"])
cat_ligne("E-08", "Conduit EMT sous dalle étage + 5 #12 (7 courses × 25 m)", pieds(DET["E-08"]["q"]),
          "CONDUIT 3/4 05C12", "3.31", "Estimation_2 E-08 %s m → pi (E-10 #12 inclus)" % DET["E-08"]["q"])
cat_ligne("D-14", "Conduit vide 3/4 data sous dalle RDC (11 courses)", pieds(DET["D-14"]["q"]),
          "CONDUIT VIDE 3/4", "3.31", "Estimation_2 D-14 %s m → pi" % DET["D-14"]["q"])
cat_ligne("E-09", "Conduit vide 3/4 data sous dalle étage (7 courses)", pieds(DET["E-09"]["q"]),
          "CONDUIT VIDE 3/4", "3.31", "Estimation_2 E-09 %s m → pi" % DET["E-09"]["q"])
cat_ligne("D-16/E-11", "Boîtes de tirage/jonction des courses sous dalle (23 courses × 2)", int(LUMEN["IBE72171K"]["q"]),
          "BOITE CARRE 2-1/8PO KO 1", "3.31", "Matériel Lumen IBE72171K « 23 courses sous dalle x 2 » (remplace D-16/E-11)")
cat_ligne("D-16c", "Couvercles des boîtes de tirage", int(LUMEN["IBE72C1"]["q"]), "COUV CARRE 4PO VIERGE", "3.31", "idem IBE72C1")
lumen_ligne("GRE430", "Corde de tirage — conduits vides (devis 35.10)", LUMEN["GRE430"]["q"], "GRE430", 0.0, "3.31",
            "Matériel Lumen", "0 h (posée avec le conduit vide)")
cat_ligne("D-SCAN", "Scan de dalle béton (GPR) avant percements", 1, "SCAN DALLE BETON", "3.32",
          "1 lot", prix_zero=True)

# ===== Éclairage, urgence, contrôle (E-203) — 0 $ en ligne + lot éclairage
cat_ligne("F-01", "Luminaire A1 — Cooper 24CGTS 2'x4' encastré 347 V", por("LUMINAIRE A1", "INSTALLER"),
          "FIXTURE TYPE A1", "3.35", R + " LUMINAIRE A1", prix_zero=True)
cat_ligne("F-02", "Luminaire C1 — Axis SL1 module 4'", por("LUMINAIRE C1", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1", prix_zero=True)
cat_ligne("F-03", "Luminaire C1_1 — Axis SL1 module 1'", por("LUMINAIRE C1_1", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1_1 (famille C1)", prix_zero=True)
cat_ligne("F-04", "Luminaire C1_3 — Axis SL1 module 3'", por("LUMINAIRE C1_3", "INSTALLER"),
          "FIXTURE TYPE C1", "3.35", R + " LUMINAIRE C1_3 (famille C1)", prix_zero=True)
cat_ligne("F-06", "Luminaire E — GreenCreative PXCYL cylindre suspendu", por("LUMINAIRE E", "INSTALLER"),
          "FIXTURE TYPE E", "3.35", R + " LUMINAIRE E (3 étiquettes ; Estimation_2 en compte 4)", prix_zero=True)
cat_ligne("F-07", "Enseigne de sortie Z6 — Stanpro RMSOWH-UDC", por("ENSEIGNE SORTIE Z6", "INSTALLER"),
          "ENSEIGNE SORTIE Z6", "3.35", R + " ENSEIGNE SORTIE Z6", prix_zero=True)
jug_ligne("F-05", "3.35")
cat_ligne("R-INT", "Interrupteur simple 15 A (installer)", por("INT SIMPLE", "INSTALLER"),
          "INTERRUPTEUR 15A 120V", "3.35", R + " INT SIMPLE INSTALLER")
jug_ligne("A-07", "3.35")
cat_ligne("F-08", "Conduit EMT 3/4 + 3 #12 — raccordement luminaires (8 m × 35 points)", pieds(DET["F-08"]["q"]),
          "CONDUIT 3/4 03C12", "3.35", "Estimation_2 F-08 %s m → pi (F-09 #12 inclus)" % DET["F-08"]["q"])
cat_ligne("F-10", "Boîtes de jonction éclairage", int(LUMEN["IBE54151L"]["q"]), "BOITE OCTOG 1-1/2PO KO 1", "3.35",
          "Matériel Lumen IBE54151L « JB éclairage » 35 (remplace F-10)")
cat_ligne("F-10t", "Tige filetée 3/8 — suspension à la dalle (devis 16.5)", int(LUMEN["VISTF3/8"]["q"]), "TIGE FILETE 3/8PO", "3.35",
          "Matériel Lumen VISTF3/8 200 pi")
lumen_ligne("VISDIA3/816BX50", "Ancrages drop-in 3/8-16 (paquet de 50)", LUMEN["VISDIA3/816BX50"]["q"], "VISDIA3/816BX50", 0.0, "3.35",
            "Matériel Lumen", "0 h (pose comptée dans F-05)")

# ===== Alarme incendie
cat_ligne("G-01", "Détecteur de fumée adressable Simplex (fourni JCI)", por("DETECTEUR FUMEE", "INSTALLER"),
          "DETECTEUR FUMEE", "3.36", R + " DETECTEUR FUMEE INSTALLER", prix_zero=True)
cat_ligne("R-DFR", "Détecteur de fumée existant à relocaliser", por("DETECTEUR FUMEE", "RELOCALISER"),
          "RELO DETECTEUR A FUME", "3.36", R + " DETECTEUR FUMEE RELOCALISER")
cat_ligne("G-02", "Conduit vide 3/4 boucle AI", pieds(DET["G-02"]["q"]), "CONDUIT VIDE 3/4", "3.36",
          "Estimation_2 G-02 %s m → pi" % DET["G-02"]["q"])
cat_ligne("G-03", "Câble AI FT4 #16 (boucle adressable)", pieds(DET["G-03"]["q"]), "COND ALARME 3/16 FT4", "3.36",
          "Estimation_2 G-03 %s m → pi" % DET["G-03"]["q"])
jug_ligne("G-05", "3.36")
cat_ligne("G-04", "Relocalisation panneau annonciateur Simplex", por("PANNEAU ANNONCIATEUR", "INSTALLER"),
          "RELO PANN ANNOCIATEUR", "3.37", R + " PANNEAU ANNONCIATEUR (nouvel emplacement) / E-007 note 1")
jug_ligne("G-06", "3.38")

# ===== Distribution — 0 $ en ligne + lot distribution
cat_ligne("B-01", "Panneau PP-3 NF 225 A 347/600 V 42 cct (fourni au lot)", int(DET["B-01"]["q"]),
          "PANNEAU 225A 120/208V 3F", "3.24", "Estimation_2 B-01 — temps panneau 8 h", prix_zero=True)
cat_ligne("B-01d", "Démantèlement de l'ancien PP-3", 1, "ENL PANN 42CCT 120/208V", "3.24", "Estimation_2 B-01 (incl. démantèlement)")
q_transf = int(DET["B-02"]["q"]) + int(DET["B-03"]["q"])
cat_ligne("B-01t", "Transfert des circuits existants dans PP-3", q_transf,
          "CIRCUITS DU PANNEAU PC VERS PANNEAU P-2", "3.24", "Estimation_2 B-02 + B-03")
cat_ligne("B-1P", "Disjoncteurs 1P (EDB) — lot distribution", 20, "DISJ15A1P", "3.25",
          "lot distribution EDB 1P ×20 (B-02 14 + B-06 1 + B-07 2 + B-08 2 + 1)", prix_zero=True)
cat_ligne("B-3P", "Disjoncteurs 3P (EJB) — aérotherme, UTA-1, VE-2N", 3, "DISJ 100A 3P 120/240V 22 K", "3.25",
          "lot distribution EJB 3P ×3 (B-03, B-04, B-05) — temps 3P 1 h", prix_zero=True)
jug_ligne("B-09", "3.25")

# ===== Mécanique
cat_ligne("C-01", "UTA-01 York 575 V/3Ø — raccordement (sectionneur d'usine)", int(DET["C-01"]["q"]),
          "MOTEUR UAF-1", "3.23", "Estimation_2 C-01 / relevé MOTEUR TRIPHASE INSTALLER (E-106)")
cat_ligne("C-02", "VE-02N Plymovent 575 V/3Ø — raccordement", int(DET["C-02"]["q"]), "MOTEUR", "3.23",
          "Estimation_2 C-02 / relevé MOTEUR TRIPHASE INSTALLER (E-106)")
cat_ligne("C-04", "Humidificateurs HUM-1/2/3 120 V — raccordement", int(DET["C-04"]["q"]), "HUMIDIFICATEUR", "3.23",
          "Estimation_2 C-04 / relevé MOTEUR MONOPHASE INSTALLER 3 (E-105)")
cat_ligne("C-04s", "Sectionneur 30 A 2P Hubbell HBL1372 (HUM)", int(DET["C-04"]["q"]), "INTERR 30A 2P 600V CEMA1", "3.23",
          "Estimation_2 C-04 / Lumen HUBHBL1372 (même prix 117,84 $)")
cat_ligne("C-05", "Conduit EMT 1 + 4 #8 RW90 — alimentation UTA-01", pieds(DET["C-05"]["q"]), "CONDUIT 1 04C08", "3.23",
          "Estimation_2 C-05 %s m → pi (C-06 #8 inclus)" % DET["C-05"]["q"])
cat_ligne("C-07a", "Conduit EMT 3/4 + 4 #12 — VE-02N", pieds(35), "CONDUIT 3/4 04C12", "3.23",
          "Estimation_2 C-07 (35 m VE-02N) → pi (C-08 #12 inclus)")
cat_ligne("C-07b", "Conduit EMT 3/4 + 3 #12 — HUM-1/2/3", pieds(DET["C-07"]["q"] - 35), "CONDUIT 3/4 03C12", "3.23",
          "Estimation_2 C-07 (3 × 15 m HUM) → pi (C-08 #12 inclus)")
cat_ligne("R-TH", "Thermostat (installer)", por("THERMOSTAT", "INSTALLER"), "THERMOSTAT", "3.23",
          R + " THERMOSTAT INSTALLER")
q_toit = por("MOTEUR MONOPHASE", "RELOCALISER")
cat_ligne("A-10d", "Toit : débrancher équipements de ventilation (E.A.R)", q_toit,
          "DEBRANCHEMENT APPAREILAGE MECANIQUE", "3.23", R + " MOTEUR MONOPHASE RELOCALISER (E-103) / A-10")
cat_ligne("A-10s", "Toit : sectionneurs déposés, entreposés, réinstallés", q_toit, "SECTIONNEUR 30A 240V 1P", "3.23",
          R + " MOTEUR MONOPHASE RELOCALISER / A-10")
cat_ligne("A-10r", "Toit : rebranchement des équipements", q_toit, "MOTEUR", "3.23",
          R + " MOTEUR MONOPHASE RELOCALISER / A-10")
cat_ligne("C-03", "Chaudières CH-01/02 600 V/3Ø — réutiliser tiroirs CCM-1, raccord final", int(DET["C-03"]["q"]),
          "MOTEUR", "3.26", "Estimation_2 C-03 / relevé MOTEUR TRIPHASE INSTALLER (E-105)")
lumen_ligne("C-03c", "Câble TECK90 3C #12 600 V (chaudières)", LUMEN["CAB12TECK3C600V"]["q"], "CAB12TECK3C600V", 0.03, "3.26",
            "Matériel Lumen 10 m", "jugement 0,03 h/m")
lumen_ligne("C-03k", "Connecteurs Star Teck 3/4", LUMEN["IBETEK75120"]["q"], "IBETEK75120", 0.25, "3.26",
            "Matériel Lumen 4", "jugement 0,25 h/connecteur")

# ===== Démontage / relocalisation (3.22) — retrait temporaire, relocalisations
cat_ligne("R-PR", "Prises existantes à relocaliser (E.A.R)",
          por("PRISE DOUBLE 15A", "RELOCALISER") + por("PRISE DDFT", "RELOCALISER"), "RELO PRISE", "3.22",
          R + " PRISE DOUBLE 15A + PRISE DDFT RELOCALISER")
cat_ligne("R-DR", "Sorties data existantes à relocaliser (E.A.R)", por("SORTIE DATA", "RELOCALISER"), "RELO PRISE", "3.22",
          R + " SORTIE DATA RELOCALISER (famille relocalisation de sortie)")
cat_ligne("R-LR", "Encastrés existants à relocaliser (E.A.R)", por("ENCASTRE", "RELOCALISER"), "RELO ENCASTRE", "3.22",
          R + " ENCASTRE RELOCALISER")
jug_ligne("A-08", "3.22")
jug_ligne("A-09", "3.22")

# ===== Démolition (5.19) — E.A.E
cat_ligne("R-XE", "Encastrés à enlever (E.A.E)", por("ENCASTRE", "ENLEVER"), "ENLEVER FIXTURE ENCAST", "5.19",
          R + " ENCASTRE ENLEVER")
cat_ligne("R-X14", "Luminaires 1'x4' à enlever", por("LED 1X4", "ENLEVER"), "ENLEVER FIXTURE 1X4", "5.19",
          R + " LED 1X4 ENLEVER")
cat_ligne("R-X22", "Luminaires 2'x2' à enlever", por("LED 2X2", "ENLEVER"), "ENLEVER FIXTURE 2X2", "5.19",
          R + " LED 2X2 ENLEVER")
cat_ligne("R-XP", "Prises à enlever", por("PRISE DOUBLE 15A", "ENLEVER") + por("PRISE DDFT", "ENLEVER"),
          "ENLEVER PRISE", "5.19", R + " PRISE DOUBLE 15A + PRISE DDFT ENLEVER")
cat_ligne("R-XD", "Sorties data à enlever", por("SORTIE DATA", "ENLEVER"), "ENLEVER PRISE", "5.19",
          R + " SORTIE DATA ENLEVER (famille la plus proche)")
cat_ligne("R-XI", "Interrupteur à enlever", por("INT SIMPLE", "ENLEVER"), "ENLEVER INTERRUPTEUR", "5.19",
          R + " INT SIMPLE ENLEVER")
q_xeq = (por("MONUMENT PLANCHER", "ENLEVER") + por("BOITE JONCTION", "ENLEVER") + por("BOITIER C", "ENLEVER")
         + por("DETECTEUR FUMEE", "ENLEVER") + por("DETECTEUR MOUVEMENT", "ENLEVER") + por("HP ENCASTRE PLAFOND", "ENLEVER"))
cat_ligne("R-XQ", "Monuments, JB, boîtiers C, détecteurs, haut-parleur à enlever", q_xeq, "EQUIPEMENT A ENLEVER", "5.19",
          R + " ENLEVER : MONUMENT 3 + BOITE JONCTION 2 + BOITIER C 2 + DET. FUMÉE 1 + DET. MOUV. 1 + HP 1")

# ===== Conditions de mise en œuvre (6.1)
cat_ligne("H-01", "Nacelle / plateforme roulante, mobilisation", int(DET["H-01"]["q"]), "NACELLE", "6.1",
          "Estimation_2 H-01")
jug_ligne("H-02", "6.1")
jug_ligne("H-03", "6.1")
jug_ligne("H-04", "6.1")
jug_ligne("A-11", "6.1")

# ---------------------------------------------------------------- totaux
for L in LIGNES:
    if L["art"] not in ARTICLES:
        fail("article absent du bordereau : " + L["art"])
    if not (isinstance(L["q"], (int, float)) and L["q"] > 0):
        fail("quantité nulle/invalide : " + L["code"])

codes = [L["code"] for L in LIGNES]
if len(codes) != len(set(codes)):
    fail("code de ligne en double")

mat_lignes = sum(L["q"] * L["p"] for L in LIGNES)
mat_lots = sum(v[0] for v in LOTS.values())
for nom, (mt, vent) in LOTS.items():
    if abs(sum(vent.values()) - mt) > 1e-6:
        fail("ventilation du lot incohérente : " + nom)
materiel = round(mat_lignes + mat_lots, 2)
heures_relevees = round(sum(L["q"] * L["h"] for L in LIGNES), 2)
heures_chargees = int(math.ceil(heures_relevees / 10.0) * 10)
main_doeuvre = round(heures_chargees * TAUX, 2)
total_acceo = round(materiel + main_doeuvre, 2)
prix_vente = int(round(total_acceo * FACTEUR / 1000.0) * 1000)

# ventilation par article : heures relevées ; MO au prorata des heures chargées (l'arrondi est réparti)
coef = heures_chargees / heures_relevees
par = {}
for L in LIGNES:
    a = par.setdefault(L["art"], {"materiel": 0.0, "heures": 0.0})
    a["materiel"] += L["q"] * L["p"]
    a["heures"] += L["q"] * L["h"]
for nom, (mt, vent) in LOTS.items():
    for art, m in vent.items():
        par.setdefault(art, {"materiel": 0.0, "heures": 0.0})["materiel"] += m


def cle_art(a):
    return tuple(int(x) for x in a.split("."))


par_article = {}
for art in sorted(par, key=cle_art):
    m = par[art]["materiel"]
    h = par[art]["heures"]
    par_article[art] = {"materiel": round(m, 2), "heures": round(h, 2), "total": round(m + h * coef * TAUX, 2)}

out = {
    "taux": TAUX,
    "materiel": materiel,
    "heures_relevees": heures_relevees,
    "heures_chargees": heures_chargees,
    "main_doeuvre": main_doeuvre,
    "total_acceo": total_acceo,
    "facteur": FACTEUR,
    "prix_vente": prix_vente,
    "lignes": LIGNES,
    "par_article": par_article,
    "lots": {nom: v[0] for nom, v in LOTS.items()},
}

# contrôles de cohérence
assert abs(out["materiel"] + out["heures_chargees"] * out["taux"] - out["total_acceo"]) < 0.01
assert out["prix_vente"] == int(round(out["total_acceo"] * out["facteur"] / 1000.0) * 1000)
assert out["heures_chargees"] >= out["heures_relevees"] > out["heures_chargees"] - 10
assert abs(sum(v["total"] for v in par_article.values()) - total_acceo) < 0.5
assert abs(sum(v["heures"] for v in par_article.values()) - heures_relevees) < 0.05

with open(F_SORTIE, "w", encoding="utf-8") as fh:
    json.dump(out, fh, ensure_ascii=False, indent=1)

print("lignes=%d materiel=%.2f heures_relevees=%.2f heures_chargees=%d MO=%.2f total_acceo=%.2f prix_vente=%d"
      % (len(LIGNES), materiel, heures_relevees, heures_chargees, main_doeuvre, total_acceo, prix_vente))
for art, v in par_article.items():
    print("  %-5s mat %10.2f  h %7.2f  total %11.2f" % (art, v["materiel"], v["heures"], v["total"]))
