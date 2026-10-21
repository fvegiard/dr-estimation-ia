# Contrôle + génération de lots/releveur-A3-2.json (releveur A3-2, sous-lot éclairage/AI et services sous dalle)
import json, math, pathlib
import openpyxl

BASE = pathlib.Path('/home/claude/entrainement/S-1902/swarm-2')
wb = openpyxl.load_workbook(BASE / 'entree/reference-Estimation_2.xlsx', data_only=True)
ws = wb['Détail']

def lire(code):
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == code]
    assert len(rows) == 1, (code, rows)
    r = rows[0]
    return {'ligne': r, 'ref': ws.cell(r, 3).value, 'qte': ws.cell(r, 4).value,
            'unite': ws.cell(r, 5).value, 'note': ws.cell(r, 12).value}

def pi(m):
    return math.floor(m * 3.2808 + 0.5)

def fr(x):
    return f"{x:.2f}".replace('.', ',')

def note_txt(d):
    return d['note'] if d['note'] is not None else 'vide'

def ligne_m(code, desc, art, suffixe=''):
    d = lire(code)
    assert d['unite'] == 'm', (code, d)
    m = d['qte']
    q = pi(m)
    src = (f"Estimation_2 Détail {code} Qté {m} m × 3,2808 = {fr(m * 3.2808)} → {q} pi ; "
           f"note L : {note_txt(d)} ; Réf. plan {d['ref']}{suffixe}")
    return {'code': code, 'desc': desc, 'q': q, 'u': 'PI', 'p': None, 'h': None, 'art': art, 'src': src}, m, q

def ligne_u(code, desc, art, q_attendu, u):
    d = lire(code)
    unite_attendue = 'u' if u == 'U' else 'lot'
    assert d['qte'] == q_attendu and d['unite'] == unite_attendue, (code, d)
    src = (f"Estimation_2 Détail {code} Qté {d['qte']} {d['unite']} ; "
           f"note L : {note_txt(d)} ; Réf. plan {d['ref']}")
    return {'code': code, 'desc': desc, 'q': q_attendu, 'u': u, 'p': None, 'h': None, 'art': art, 'src': src}

lignes = []
l, m_f08, q_f08 = ligne_m('F-08', 'Conduit EMT 3/4" + 3 #12 RW90 éclairage — assemblage 03C12', '3.35',
                          suffixe=' ; 3 cond. RW90 #12 : F-09 exclu (840 m)')
assert (m_f08, q_f08) == (280, 919)
lignes.append(l)

l, m_g02, q_g02 = ligne_m('G-02', 'Conduit EMT 3/4" boucle alarme incendie', '3.36')
assert (m_g02, q_g02) == (30, 98)
lignes.append(l)

l, m_g03, q_g03 = ligne_m('G-03', 'Câble FT4 2c #16 rouge boucle adressable (dans G-02)', '3.36',
                          suffixe=' ; câble AI, reprise séparée (exception, pas un fil RW90 d\'assemblage)')
assert (m_g03, q_g03) == (60, 197)
lignes.append(l)

lignes.append(ligne_u('D-11', "Bouton d'appel — boîte + conduit vide (système fourni par d'autres)", '3.31', 1, 'U'))
lignes.append(ligne_u('D-12', 'Percement de dalle carotté + coupe-feu (CAN/ULC-S115, Hilti CP-506) — conduits sous dalle RDC', '3.31', 16, 'U'))
lignes.append(ligne_u('D-16', 'Boîtes de tirage/jonction, connecteurs, supports, corde de tirage — conduits sous dalle RDC', '3.31', 1, 'LOT'))
lignes.append(ligne_u('E-06', "Bouton d'appel (2) + bouton de panique (1) — boîte + conduit vide", '3.31', 3, 'U'))
lignes.append(ligne_u('E-07', 'Percement de dalle carotté + coupe-feu — étage (7 courses)', '3.31', 7, 'U'))
lignes.append(ligne_u('E-11', 'Boîtes, connecteurs, supports — étage', '3.31', 1, 'LOT'))

# Exclusion F-09 : longueur lue et conversion refaite
d_f09 = lire('F-09')
assert d_f09['unite'] == 'm' and d_f09['qte'] == 840
pi_f09 = pi(d_f09['qte'])
assert pi_f09 == 2756
exclusions = [{'code': 'F-09',
               'raison': (f"incluse dans l'assemblage 03C12 de F-08 (3 cond. = 840 / 280) ; "
                          f"fil RW90 #12 éclairage 840 m = {pi_f09} pi (840 × 3,2808 = {fr(840 * 3.2808)} → {pi_f09} pi) ; "
                          f"à rétablir si le directeur B chiffre un conduit vide")}]

# Contrôles croisés (lecture des feuilles Lumen et alarme, aucune quantité modifiée)
ws_l = wb['Matériel Lumen']
lumen = {ws_l.cell(r, 1).value: ws_l.cell(r, 7).value for r in range(1, ws_l.max_row + 1) if ws_l.cell(r, 1).value}
assert '895 m' in lumen['CONEMT3/4'] and 80 + 330 + 175 + 280 + 30 == 895
assert 275 + 2400 + 875 + 840 == 4390
assert '23 courses' in lumen['IBE72171K'] and '23 percements' in lumen['TROCP25WB+10.1OZ']
assert '35 points' in lumen['THOABRH03830'] and 35 * 1.8 == 63
qte_lumen_12 = sum(ws_l.cell(r, 3).value for r in range(1, ws_l.max_row + 1)
                   if str(ws_l.cell(r, 1).value).startswith('CAB12/1RW90'))
assert qte_lumen_12 == 4390, qte_lumen_12
assert round(895 * 3.2808) == 2936  # base CONEMT3/4 en pieds (Lumen 2940 pi = arrondi à la boîte de 10 pi)

wb_a = openpyxl.load_workbook(BASE / 'entree/reference-Releve_Alarme_Incendie.xlsx', data_only=True)
wa = wb_a['Relevé']
alarme = {wa.cell(r, 1).value: (wa.cell(r, 2).value, wa.cell(r, 3).value, wa.cell(r, 4).value)
          for r in range(1, wa.max_row + 1) if wa.cell(r, 1).value}
assert alarme['AI-08'][1] == 60 and alarme['AI-08'][2] == 'm'
assert alarme['AI-09'][1] == 30 and alarme['AI-09'][2] == 'm'

wb_b = openpyxl.load_workbook(BASE / 'entree/reference-Bordereau_de_prix.xlsx', data_only=True)
wbd = wb_b['Bordereau de prix']
libelles = {str(wbd.cell(r, 1).value).strip(): wbd.cell(r, 2).value for r in range(1, wbd.max_row + 1) if wbd.cell(r, 1).value}
assert libelles['3.31'] == 'Percements, conduit filage et boîte de jonction'
assert libelles['3.35'] == "Appareil de contrôle + appareils d'éclairage"
assert libelles['3.36'] == 'Détecteur de fumée'
assert libelles['5.11'] == 'Percement du plancher'

notes = [
    "p et h sont null partout (prix et temps réservés à l'autre directeur) ; les lignes heures de Détail (estimé Léna AI) ne sont pas utilisées ; longueurs de Détail en mètres converties en pieds (× 3,2808, arrondi à l'entier).",
    "Contrôle croisé Lumen CONEMT3/4 : base « 895 m (C-07 80 + D-14 330 + E-09 175 + F-08 280 + G-02 30) » = 895 m, vérifié. Part de ce sous-lot : F-08 280 + G-02 30 = 310 m ; les 585 m restants (C-07, D-14, E-09) sont à A3-1. Lumen 2940 pi = arrondi à la boîte de 10 pi (895 m = 2936 pi).",
    "Contrôle croisé Lumen #12 : 275 + 2400 + 875 + 840 = 4390 m, vérifié ; Qté Lumen des bobines (1500 + 1200 + 900 + 790) = 4390 m. Part de ce sous-lot : F-09 840 m, exclu (incluse dans l'assemblage 03C12 de F-08).",
    "Hypothèse F-08 « Hyp. 8 m x 35 points » : 8 × 35 = 280 m, vérifié ; cohérent avec le Sommaire note 2 (8 m/point éclairage). Le flexible Lumen THOABRH03830 (63 m = 35 points × 1,8 m) confirme les 35 points comme base de calcul.",
    "Réserve R-003 (reserves.md) : le relevé automatique compte 44 luminaires typés sur E-203 (A1 11, C1 12, C1_1 6, C1_3 15) plus 3 étiquettes E sur bloc (estimation en compte 4). Ce comptage ne concorde pas avec les 35 points de l'hypothèse F-08 : à revérifier au plan. Quantité F-08 non modifiée, écart signalé.",
    "Alarme incendie (reference-Releve_Alarme_Incendie.xlsx, feuille Relevé) : AI-08 câble AI FT4 2c #16 = 60 m, confirme G-03 (60 m) ; AI-09 conduit EMT 3/4\" + boîtes = 30 m, confirme G-02 (30 m). Les deux lignes AI sont marquées « longueur estimée » : cohérence, pas mesure indépendante au plan.",
    "Bordereau (reference-Bordereau_de_prix.xlsx, feuille Bordereau de prix) : libellés exacts vérifiés : 3.31 « Percements, conduit filage et boîte de jonction » ; 3.35 « Appareil de contrôle + appareils d'éclairage » ; 3.36 « Détecteur de fumée » ; 5.11 « Percement du plancher ».",
    "Alternatives d'article à trancher : F-08 (conduit d'éclairage) pourrait relever de 3.31 « conduit filage » plutôt que 3.35. G-02 et G-03 (conduit et câble de boucle AI) sont portés en 3.36 « Détecteur de fumée » selon le brief : ce libellé ne décrit pas un conduit ni un câble ; 3.31 est à considérer.",
    "Percements : alternative 5.11 « Percement du plancher » notée pour D-12 (16) et E-07 (7), non retenue (3.31 conservé selon le brief). Cohérence : D-12 = 16 percements RDC = 16 courses de D-13 (« 16 courses ») ; E-07 = 7 percements étage = 7 courses de E-08 (« 7 courses ») ; total 23 courses = Lumen IBE72171K « Tirage : 23 courses sous dalle » et TROCP25WB « 23 percements + extras ». Aucune quantité changée.",
    "Lots D-16 et E-11 (Qté 1, unité lot, sans décompte unitaire dans Détail) : quantité et unité lues telles quelles ; le contenu Lumen (boîtes, connecteurs, corde GRE430 « conduits vides ») n'est pas décompté ici ; article 3.31 et unité LOT à valider.",
    "D-11 : Qté 1 u ; la note L « Aide au lecteur 116 » est une référence de page, pas une justification de quantité ; système fourni par d'autres, boîte + conduit vide seulement.",
    "Longueurs de conduit du sous-lot (F-08 280 m, G-02 30 m) sont des hypothèses du Sommaire note 2 (8 m/point éclairage) ; G-02 et G-03 n'ont que la mention « Hyp. » dans Détail. À valider sur plan.",
    "Exclusion F-09 réversible : si le directeur B chiffre un conduit vide, rétablir F-09 = 840 m = 2756 pi ; G-03 (câble FT4) est repris séparément, sa description « (dans G-02) » indique seulement le cheminement dans le conduit G-02, sans recalcul du conduit.",
    "Aucune réserve de reserves.md ne touche directement les lignes D-11 à E-11 ; R-005 (pictogrammes détecteurs de fumée) concerne les dispositifs AI fournis par JCI, hors de ce sous-lot.",
]

sortie = {
    'niveau': 'releveur',
    'nom': 'A3-2',
    'modele_demande': 'haiku',
    'effort_demande': 'high',
    'lignes': lignes,
    'notes': notes,
    'exclusions': exclusions,
    'lots': {},
}

chemin = BASE / 'lots/releveur-A3-2.json'
chemin.write_text(json.dumps(sortie, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

# ---- Contrôles (critères 3, 4, 5, 6) ----
relu = json.loads(chemin.read_text(encoding='utf-8'))
assert 'lignes' in relu and 'notes' in relu
codes_lignes = [x['code'] for x in relu['lignes']]
codes_excl = [x['code'] for x in relu['exclusions']]
assert len(codes_lignes) == len(set(codes_lignes)), 'doublon dans lignes'
assert not (set(codes_lignes) & set(codes_excl)), 'code à la fois repris et exclu'
requis = {'F-08', 'F-09', 'G-02', 'G-03', 'D-11', 'D-12', 'D-16', 'E-06', 'E-07', 'E-11'}
assert set(codes_lignes) | set(codes_excl) == requis and len(requis) == 10
assert len(codes_lignes) == 9 and codes_excl == ['F-09']
for x in relu['lignes']:
    assert x['p'] is None and x['h'] is None
    assert x['u'] in {'U', 'PI', 'LOT'}
    assert isinstance(x['art'], str)
assert relu['modele_demande'] == 'haiku'
# recalcul indépendant des conversions à partir des Qté relues
for x in relu['lignes']:
    if x['u'] == 'PI':
        qte_m = lire(x['code'])['qte']
        assert x['q'] == math.floor(qte_m * 3.2808 + 0.5), (x['code'], x['q'])
        assert f"{qte_m} m" in x['src'] and f"→ {x['q']} pi" in x['src']
    else:
        assert x['q'] == lire(x['code'])['qte']

print(f"{len(relu['lignes'])} lignes")
print(', '.join(f"{x['code']} = {x['q']}" for x in relu['lignes']))
print('exclusions:', codes_excl)
print('CONTROLE OK')
