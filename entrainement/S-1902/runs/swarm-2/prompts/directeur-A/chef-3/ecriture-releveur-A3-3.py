#!/usr/bin/env python3
# Releveur A3-3 : écrit lots/releveur-A3-3.json à partir des cellules réellement lues
# (Estimation_2 Détail, catalogue Dupuis, METHODE.md) puis le relit et le vérifie.
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

BASE = Path("/home/claude/entrainement/S-1902/swarm-2")
SORTIE = BASE / "lots" / "releveur-A3-3.json"
CLASSEUR = BASE / "entree" / "reference-Estimation_2.xlsx"
CATALOGUE = BASE / "catalogue-prix-temps-dupuis.csv"
METHODE = BASE / "METHODE.md"


def lire_detail(classeur_wb, code):
    ws = classeur_wb["Détail"]
    lignes = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == code]
    assert len(lignes) == 1, (code, lignes)
    r = lignes[0]
    return {
        "row": r,
        "desc": ws.cell(r, 2).value,
        "D": ws.cell(r, 4).value,
        "E": ws.cell(r, 5).value,
        "L": ws.cell(r, 12).value,
    }


def lire_catalogue(libelle):
    with open(CATALOGUE, encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f, delimiter=";"))
    hits = [(i + 2, r) for i, r in enumerate(rows) if r["description"] == libelle]
    assert len(hits) == 1, (libelle, len(hits))
    return hits[0]


def normaliser(texte):
    return re.sub(r"[\s  ]", "", texte)


def construire():
    wb = openpyxl.load_workbook(CLASSEUR, data_only=True)
    methode = normaliser(METHODE.read_text(encoding="utf-8"))
    for montant in ("11893", "6200", "7400", "600", "1250"):
        assert montant in methode, montant

    def src_detail(code):
        d = lire_detail(wb, code)
        return d, (
            f"Estimation_2 Détail {code} (ligne {d['row']}) : Qté (D) = {d['D']}, "
            f"Unité (E) = {d['E']}, Notes (L) = « {d['L']} »"
        )

    lignes = []
    for code, art in (("A-11", "6.1"), ("H-01", "6.1"), ("H-02", "6.1"),
                      ("H-03", "6.1"), ("H-04", "6.1")):
        d, src = src_detail(code)
        assert d["D"] == 1 and d["E"] == "lot", (code, d)
        lignes.append({
            "code": code,
            "desc": d["desc"],
            "q": 1,
            "u": "LOT",
            "p": None,
            "h": None,
            "art": art,
            "src": src,
        })

    n_nac, nac = lire_catalogue("NACELLE")
    assert nac["um"] == "U" and nac["prix_a_utiliser"] == "1250.0" and nac["temps_a_utiliser"] == "16.0", nac
    lignes.append({
        "code": "X-02",
        "desc": "Nacelle (location, 16 h)",
        "q": 1,
        "u": "U",
        "p": None,
        "h": None,
        "art": "6.1",
        "src": (
            "METHODE.md §2 point 3 : nacelle 1 250 $ / 16 h ; catalogue-prix-temps-dupuis.csv "
            f"ligne {n_nac} « NACELLE » : prix_a_utiliser 1250.0, temps_a_utiliser 16.0 "
            f"(dernier projet {nac['projet_dernier']}, {nac['date_dernier']})"
        ),
    })

    lots_fournisseur = [
        ("LOT-ECL", "Lot éclairage (A1 ×11, Axis SL1 4′ ×12, 3′ ×15, 1′ ×6, trousses, E ×3, Z6 ×3)", "3.35",
         "METHODE.md §2 point 2 : lot éclairage 11 893 $ (prix confirmé Francis 2026-10-20)"),
        ("LOT-DIST", "Lot distribution (NF 225 A 42 cct + EDB 1P ×20 + EJB 3P ×3)", "3.24",
         "METHODE.md §2 point 2 : lot distribution 6 200 $ (prix confirmé Francis 2026-10-20)"),
        ("LOT-JCI", "Sous-traitance JCI/Simplex", "3.38",
         "METHODE.md §2 point 2 : JCI 7 400 $ (prix confirmé Francis 2026-10-20)"),
        ("LOT-DET", "2 détecteurs de fumée (fournis JCI)", "3.36",
         "METHODE.md §2 point 2 : 2 détecteurs 600 $ (prix confirmé Francis 2026-10-20)"),
        ("LOT-GPR", "GPR (géoradar = scan de dalle avant carottage)", "3.32",
         "METHODE.md §2 point 2 : GPR 1 250 $ (prix confirmé Francis 2026-10-20)"),
    ]
    for code, desc, art, src in lots_fournisseur:
        lignes.append({
            "code": code,
            "desc": desc,
            "q": 1,
            "u": "LOT",
            "p": None,
            "h": None,
            "art": art,
            "src": src,
        })

    exclusions = []
    d_g07, src_g07 = src_detail("G-07")
    assert d_g07["D"] == 1 and d_g07["E"] == "lot", d_g07
    exclusions.append({
        "code": "G-07",
        "raison": "double compte : remplacée par LOT-JCI (" + src_g07 + ")",
    })
    for code, libelle in (
        ("OPT-1", "câblage data Cat6 jusqu'à la salle des serveurs (devis 36.1), contradiction avec « réutiliser câble data existant »"),
        ("OPT-2", "conception parasismique par firme spécialisée + certificat (devis 16.2-16.6)"),
        ("OPT-3", "prime travaux hors heures normales (devis 1.17)"),
        ("OPT-4", "postes de travail temporaires (A-102) + chute à livres temporaire"),
    ):
        d = lire_detail(wb, code)
        exclusions.append({
            "code": code,
            "raison": f"OPTIONS / À CLARIFIER (hors total) dans Détail : {libelle} (Détail ligne {d['row']})",
        })
    n_scan, scan = lire_catalogue("SCAN DALLE BETON")
    assert scan["prix_a_utiliser"] == "1250.0", scan
    exclusions.append({
        "code": "X-01",
        "raison": (
            "GPR (géoradar) = scan de dalle béton : un seul objet, porté par LOT-GPR (1 250 $, art 3.32) ; "
            f"catalogue-prix-temps-dupuis.csv ligne {n_scan} « SCAN DALLE BETON » : U, prix_a_utiliser 1250.0, "
            f"temps_a_utiliser {scan['temps_a_utiliser']}, aussi 1 250 $"
        ),
    })

    lots = {"LOT-ECL": 11893, "LOT-DIST": 6200, "LOT-JCI": 7400, "LOT-DET": 600, "LOT-GPR": 1250}

    # Contrôle arithmétique du lot éclairage (résultat dans notes, rien corrigé).
    ecl_postes = [11 * 185, 12 * 290, 15 * 230, 6 * 110, 1500, 3 * 475, 3 * 161.60]
    ecl_somme = round(sum(ecl_postes), 2)
    ecl_ecart = round(ecl_somme - lots["LOT-ECL"], 2)
    dist_somme = 3400 + 95 * 20 + 300 * 3
    assert dist_somme == lots["LOT-DIST"]

    notes = [
        "Lecture : chaque code Détail est retrouvé par sa cellule A (A-11 ligne 18, H-01 à H-04 lignes 101 à 104, "
        "G-07 ligne 97, titre OPT ligne 107, OPT-1 à OPT-4 lignes 108 à 111). Qté (D), Unité (E) et Notes (L) sont "
        "relues dans le classeur et recopiées dans src. Les heures de Détail (estimé Léna AI) ne sont pas utilisées : "
        "h = null partout.",
        "Frais généraux A-11, H-01, H-02, H-03, H-04 : la colonne L de Détail dit « Hypothèse » pour les cinq lignes "
        "(forfaits de 1 lot sans relevé sur plan) ; à confirmer par le directeur A.",
        "X-02 nacelle (16 h) chevauche possiblement H-01 « plateforme/échafaudage roulant » : les deux sont reprises "
        "ici, le directeur A tranchera. Les 16 h ne sont pas reprises en h (h = null) : elles relèvent du directeur B.",
        "X-02 nacelle : prix et temps du catalogue concordent avec METHODE.md §2 point 3 (1 250 $ / 16 h).",
        "X-01 SCAN DALLE exclu. Recherche de « GPR » dans CONSIGNES.md, METHODE.md, rapport-releve.md, reserves.md, "
        "ARCHITECTURE.md et IDEE-FRANCIS.md : le kit ne définit pas le sigle GPR et aucune occurrence ne lui donne un "
        "autre sens qu'un lot de 1 250 $ (METHODE.md §2 point 2 ; ARCHITECTURE.md « GPR 1 250 $ »). L'équivalence "
        "géoradar = scan de dalle (3.32) vient de la décision du directeur A ; le montant identique à SCAN DALLE BETON "
        "la soutient.",
        f"Catalogue : SCAN DALLE BETON (ligne {n_scan}) : U, prix_a_utiliser 1250.0, temps_a_utiliser "
        f"{scan['temps_a_utiliser']}, dernier projet {scan['projet_dernier']} ({scan['date_dernier']}). Variantes non "
        "retenues : SCAN DALLE 750 (ligne 361), SCAN BETON 750 (ligne 1811), SCANN DALLE 950 (ligne 1812). Le temps de "
        "2,0 h n'est pas repris ici (X-01 exclu, LOT-GPR sans heures) : à la charge du directeur B.",
        "METHODE.md §2 point 2 relu : les montants 11 893 $, 6 200 $, 7 400 $, 600 $ et 1 250 $ sont identiques à "
        "ceux du brief ; aucune divergence.",
        "Contrôle arithmétique lot éclairage (rien corrigé) : 11 × 185 = 2 035 ; 12 × 290 = 3 480 ; 15 × 230 = 3 450 ; "
        "6 × 110 = 660 ; trousses 1 500 ; 3 × 475 = 1 425 ; 3 × 161,60 = 484,80 ; somme = "
        f"{ecl_somme:,.2f} $ (écart avec le montant confirmé 11 893 $ : {ecl_ecart:+,.2f} $). Le montant confirmé fait foi.",
        "Lot distribution : 3 400 + 95 × 20 (1 900) + 300 × 3 (900) = 6 200 $, écart nul. Ventilation par article à faire "
        "par le directeur B : 3 400 $ panneau NF 225 A → 3.24 ; 1 900 $ EDB 1P et 900 $ EJB 3P (disjoncteurs) → 3.25. "
        "Le panneau et les disjoncteurs NF sont hors de la liste Matériel Lumen (« demandes de prix séparées »), donc "
        "pas de doublon avec Lumen.",
        "Lot JCI 7 400 $ et détecteurs 600 $ : total 8 000 $ ; le kit n'a pas d'autre total indépendant pour contrôler. "
        "Sommaire « Sous-traitance JCI/Simplex (AI) — prix à obtenir » remplacé par LOT-JCI. Le Sommaire « Matériel » "
        "(23 159,67 $) égale le seul total de Matériel Lumen : JCI n'y figure pas.",
        "G-07 (ligne 97, colonne L) : « PRIX À OBTENIR DE JCI — inscrire en matériel/s-t » ; LOT-JCI (7 400 $) tient lieu "
        "de prix, d'où l'exclusion de G-07 (double compte évité).",
        "Sommaire « Sous-traitance parasismique — prix à obtenir (si incluse) » correspond à OPT-2 : exclu, hors total, "
        "comme les autres OPT.",
        "LOT-DET (2 détecteurs neufs) : rapport-releve.md donne DETECTEUR FUMEE INSTALLER 2, concordant. reserves.md R-005 : "
        "3 détecteurs existants pointillés et neufs pleins, à confirmer au zoom.",
        "LOT-ECL : rapport-releve.md concorde pour A1 11, C1 12, C1_1 6, C1_3 15, E 3 et Z6 3 (INSTALLER). Divergence à "
        "signaler : ARCHITECTURE.md décrit le relevé d'éclairage d'urgence avec « E 4 » et « C-SUSP 20 rangées », alors "
        "que le lot dit « E ×3 » et ne nomme pas C-SUSP. reserves.md R-003 : LUMINAIRE E = 3 étiquettes sur bloc, "
        "« l'estimation en compte 4 ». Les « trousses 1 500 $ » sont un forfait sans quantité dans le kit. Rien n'est "
        "corrigé ici.",
        "LOT-ECL, Z6 : reserves.md R-005 signale 3 enseignes Z6 à tracé pointillé, « confirmer neuf » ; le lot compte "
        "Z6 ×3 à 161,60 $ (METHODE.md §2 point 2).",
        "Accessoires Matériel Lumen non repris (inclus dans les assemblages Dupuis ou relevant du directeur B ; aucune "
        "ligne A créée) : APPTWC100S (accouplement ABB CI5508 1 po) ; APPTWC75S (accouplement ABB CI5506 3/4 po) ; "
        "APPTW100S (connecteur ABB Iberville CI5408 1 po) ; APPTW75S (connecteur ABB Iberville CI5406 3/4 po) ; "
        "APPTWCL100 (courroie ABB Iberville 1 po) ; APPTWCL75 (courroie ABB Iberville CI1306 3/4 po) ; THO5333 "
        "(connecteur ABB CONSEAL isolé 3/4 po) ; IBE2210 (connecteur ABB Iberville BX 90° 3/8 po) ; GRE430 (corde de "
        "tirage Poly Line 430, conduits vides) ; IBE52151K (boîte carrée ABB BC52151-K 4 po) ; IBE52C13 (couvercle carré "
        "1/2 po surélevé) ; IBE72171K (boîte carrée 4-11/16 po) ; IBE72C1 (couvercle carré ABB BC72-C-1) ; IBE54151L "
        "(boîte octogonale 4 po) ; TROCP25WB+10.1OZ (coupe-feu 3M CP25WB+) ; VISTF3/8 (tige filetée 3/8 po) ; "
        "VISDIA3/816BX50 (ancrage drop-in 3/8-16) ; IDE30452 (marrettes Ideal 30-452) ; THOABRH03830 (conduit flexible "
        "alu 3/8 po) ; CONCONSEAL3/4 (conduit flexible étanche 3/4 po). Lumen exclut déjà luminaires, urgence, AI "
        "Simplex, panneau PP-3 et disjoncteurs NF.",
        "Contrôle des longueurs : aucune longueur dans ce sous-lot, donc aucune conversion m → pi (contrôle 6 sans objet).",
        "Autres codes Détail (A-01…A-10, B, C, D, E, F, G-01…G-06) non repris : ils relèvent des autres sous-lots.",
    ]

    return {
        "niveau": "releveur",
        "nom": "A3-3",
        "modele_demande": "haiku",
        "effort_demande": "high",
        "lignes": lignes,
        "notes": notes,
        "exclusions": exclusions,
        "lots": lots,
    }


def verifier(d):
    erreurs = []
    assert "lignes" in d and "notes" in d
    if d["nom"] != "A3-3" or d["modele_demande"] != "haiku":
        erreurs.append("en-tête")

    codes_lignes = [l["code"] for l in d["lignes"]]
    codes_excl = [e["code"] for e in d["exclusions"]]
    if len(set(codes_lignes)) != len(codes_lignes):
        erreurs.append("doublon dans lignes")
    if len(set(codes_excl)) != len(codes_excl):
        erreurs.append("doublon dans exclusions")
    if set(codes_lignes) & set(codes_excl):
        erreurs.append("code à la fois dans lignes et exclusions")

    attendu = {"A-11", "H-01", "H-02", "H-03", "H-04", "G-07", "OPT-1", "OPT-2", "OPT-3", "OPT-4",
               "LOT-ECL", "LOT-DIST", "LOT-JCI", "LOT-DET", "LOT-GPR", "X-01", "X-02"}
    if set(codes_lignes) | set(codes_excl) != attendu or len(attendu) != 17:
        erreurs.append("couverture des 17 codes")
    if len(codes_lignes) != 11 or len(codes_excl) != 6:
        erreurs.append("compte 11 repris / 6 exclus")

    for l in d["lignes"]:
        if l["p"] is not None or l["h"] is not None:
            erreurs.append(f"p/h non null : {l['code']}")
        if l["u"] not in ("U", "PI", "LOT"):
            erreurs.append(f"unité invalide : {l['code']}")
        if not isinstance(l["art"], str) or not l["src"]:
            erreurs.append(f"art/src : {l['code']}")
        if l["u"] == "PI":
            erreurs.append("conversion PI présente sans objet")

    lots_attendus = {"LOT-ECL": 11893, "LOT-DIST": 6200, "LOT-JCI": 7400, "LOT-DET": 600, "LOT-GPR": 1250}
    if d["lots"] != lots_attendus:
        erreurs.append("objet lots")
    for code in lots_attendus:
        l = next(x for x in d["lignes"] if x["code"] == code)
        if l["p"] is not None:
            erreurs.append(f"montant dans p : {code}")

    # Relecture indépendante du classeur : chaque src Détail doit correspondre aux cellules réellement lues.
    wb = openpyxl.load_workbook(CLASSEUR, data_only=True)
    for code in ("A-11", "H-01", "H-02", "H-03", "H-04", "G-07"):
        l_cell = lire_detail(wb, code)
        src = next((x["src"] for x in d["lignes"] if x.get("code") == code), None)
        if src is None:
            src = next((x["raison"] for x in d["exclusions"] if x.get("code") == code), None)
        if src is None:
            erreurs.append(f"src absent : {code}")
            continue
        if f"Qté (D) = {l_cell['D']}" not in src or f"Unité (E) = {l_cell['E']}" not in src:
            erreurs.append(f"src ne correspond pas à la cellule lue : {code}")
    for code in ("OPT-1", "OPT-2", "OPT-3", "OPT-4"):
        lire_detail(wb, code)  # présence et unicité du code en colonne A

    return erreurs


def main():
    d = construire()
    SORTIE.parent.mkdir(parents=True, exist_ok=True)
    with open(SORTIE, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)
        f.write("\n")

    relu = json.loads(SORTIE.read_text(encoding="utf-8"))
    erreurs = verifier(relu)
    print(f"lignes = {len(relu['lignes'])}, exclusions = {len(relu['exclusions'])}, notes = {len(relu['notes'])}")
    for l in relu["lignes"]:
        print(f"{l['code']} = {l['q']} {l['u']} (art {l['art']})")
    if erreurs:
        print("ECHEC :", erreurs)
        return 1
    print("VERIFICATION OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
