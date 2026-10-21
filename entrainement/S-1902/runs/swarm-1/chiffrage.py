#!/usr/bin/env python3
"""Chiffrage « méthode Dupuis » — S-1902 Bibliothèque Georgette-Lepage (à l'aveugle).

Déterministe, stdlib + openpyxl. Lit les lots (lots/lot-*.json) et le kit, revérifie chaque prix/temps
contre sa source, contrôle la couverture du relevé, puis écrit chiffrage.json.

Grammaire du champ `src` d'une ligne :
  Dupuis[LIBELLÉ]        -> p et h = colonnes prix_a_utiliser / temps_a_utiliser du catalogue (dernière valeur)
  Dupuis-temps[LIBELLÉ]  -> h = temps_a_utiliser ; p forcé à 0 (luminaires, enseignes, détecteurs, panneaux,
                            disjoncteurs : prix au lot fournisseur, règle 2)
  Lumen[CODE] / Lumen[CODE1+CODE2] -> p = somme des prix unitaires de la feuille « Matériel Lumen »
  jugement (...)         -> h repris de la feuille Détail (Estimation_2) quand un code Détail est cité
  lot confirmé (...)     -> montant de lot de METHODE.md
Dans chiffrage.json, p et h d'une ligne sont UNITAIRES ($/unité, h/unité).
"""
import glob
import json
import math
import re
import sys
from collections import OrderedDict
from pathlib import Path

import csv
import openpyxl

ICI = Path(__file__).resolve().parent
TAUX = 85.0          # $/h — METHODE.md §5 (dossiers Dupuis août–octobre 2026)
FACTEUR = 1.244      # METHODE.md §7
TOL = 1e-6

# Lots fournisseur confirmés (METHODE.md §2, prix confirmés par Francis le 2026-10-20)
LOTS_METHODE = OrderedDict([
    ("eclairage", 11893.0),
    ("distribution", 6200.0),
    ("alarme_incendie_JCI", 7400.0),
    ("detecteurs_AI", 600.0),
    ("GPR_scan_dalle", 1250.0),
])
# Lignes qui portent chaque lot (code -> nom de lot ; la distribution est ventilée sur deux lignes)
LIGNES_LOTS = {
    "LOT-ECL": "eclairage",
    "LOT-DIST-PAN": "distribution",
    "LOT-DIST-DISJ": "distribution",
    "G-07": "alarme_incendie_JCI",
    "LOT-AI-DET": "detecteurs_AI",
    "GPR": "GPR_scan_dalle",
}


def erreur(msg):
    print("ERREUR : " + msg, file=sys.stderr)
    sys.exit(1)


def lire_catalogue():
    chemin = ICI / "catalogue-prix-temps-dupuis.csv"
    cat = {}
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        for r in csv.DictReader(f, delimiter=";"):
            # un même libellé peut exister en deux unités (P et M) : marqué ambigu, refusé s'il est utilisé
            cat[r["description"]] = None if r["description"] in cat else r
    return cat


def lire_lumen_et_detail():
    wb = openpyxl.load_workbook(ICI / "entree" / "reference-Estimation_2.xlsx", data_only=True)
    lumen = {}
    for r in wb["Matériel Lumen"].iter_rows(values_only=True):
        if r[0] and isinstance(r[4], (int, float)):
            lumen[str(r[0]).strip()] = float(r[4])
    detail = {}
    for r in wb["Détail"].iter_rows(values_only=True):
        if r[0] and re.match(r"^[A-H]-\d\d$", str(r[0])):
            detail[str(r[0])] = r
    return lumen, detail


def heures_detail(row):
    """(h / unité, qté) de la feuille Détail : colonnes H (index 7) et D (index 3)."""
    if not isinstance(row[7], (int, float)) or not isinstance(row[3], (int, float)):
        erreur("ligne Détail illisible : %r" % (row,))
    return float(row[7]), float(row[3])


def lire_portees():
    """Tableau familles × portées de rapport-releve.md."""
    portees = {}
    for ligne in (ICI / "rapport-releve.md").read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\| ([A-Z0-9_ ]+?) \| (\d+) \| (.+) \|$", ligne)
        if m:
            fam = m.group(1).strip()
            for p, n in re.findall(r"([A-Z]+) (\d+)", m.group(3)):
                portees[(fam, p)] = int(n)
    return portees


def main():
    cat = lire_catalogue()
    lumen, detail = lire_lumen_et_detail()
    portees = lire_portees()
    if len({f for f, _ in portees}) != 31 or sum(portees.values()) != 645:
        erreur("tableau des portées de rapport-releve.md mal lu (%d familles, %d marques ; attendu 31 / 645)"
               % (len({f for f, _ in portees}), sum(portees.values())))

    fichiers = sorted(glob.glob(str(ICI / "lots" / "lot-*.json")))
    if len(fichiers) != 5:
        erreur("5 fichiers de lot attendus, trouvé %d" % len(fichiers))
    lignes, lots_meta = [], OrderedDict()
    for f in fichiers:
        d = json.loads(Path(f).read_text(encoding="utf-8"))
        lots_meta["lot-%d" % d["lot"]] = {"titre": d["titre"], "lignes": len(d.get("lignes", []))}
        for x in d.get("lignes", []):
            lignes.append({k: x[k] for k in ("code", "desc", "q", "u", "p", "h", "art", "src")})

    codes = [x["code"] for x in lignes]
    doublons = sorted({c for c in codes if codes.count(c) > 1})
    if doublons:
        erreur("codes en double : %s" % doublons)

    # 1. Revérification de chaque ligne contre sa source
    for x in lignes:
        src = x["src"]
        m = re.match(r"^Dupuis(-temps)?\[(.+?)\]", src)
        if m:
            if m.group(2) not in cat:
                erreur("%s : libellé absent du catalogue : %s" % (x["code"], m.group(2)))
            r = cat[m.group(2)]
            if r is None:
                erreur("%s : libellé ambigu (deux unités) dans le catalogue : %s" % (x["code"], m.group(2)))
            p = 0.0 if m.group(1) else float(r["prix_a_utiliser"])
            h = float(r["temps_a_utiliser"])
        elif src.startswith("Lumen["):
            codes_l = re.match(r"^Lumen\[(.+?)\]", src).group(1).split("+")
            if any(c not in lumen for c in codes_l):
                erreur("%s : code Lumen introuvable : %s" % (x["code"], codes_l))
            p = round(sum(lumen[c] for c in codes_l), 4)
            mj = re.search(r"jugement \(Estimation_2\) ([A-H]-\d\d)", src)
            h = heures_detail(detail[mj.group(1)])[0] if mj else 0.0
        elif src.startswith("jugement"):
            p = 0.0
            mj = re.search(r"jugement \(Estimation_2\) ([A-H]-\d\d)", src)
            if mj:
                h = heures_detail(detail[mj.group(1)])[0]
            else:  # analogie avec un temps Dupuis cité en clair « analogie Dupuis LIBELLÉ x h »
                ma = re.search(r"analogie Dupuis (.+?) [\d,]+ h", src)
                if not ma or ma.group(1) not in cat:
                    erreur("%s : jugement sans source vérifiable" % x["code"])
                h = float(cat[ma.group(1)]["temps_a_utiliser"])
        elif src.startswith("lot confirmé"):
            p, h = x["p"], 0.0
            nom = LIGNES_LOTS.get(x["code"])
            if nom is None:
                erreur("%s : ligne de lot inconnue" % x["code"])
        else:
            erreur("%s : source non reconnue : %s" % (x["code"], src))
        if abs(p - x["p"]) > TOL or abs(h - x["h"]) > TOL:
            erreur("%s : p/h (%s, %s) ≠ source (%s, %s)" % (x["code"], x["p"], x["h"], p, h))

    # 2. Lots : montants des lignes == lots confirmés
    lots = OrderedDict((k, 0.0) for k in LOTS_METHODE)
    for x in lignes:
        if x["code"] in LIGNES_LOTS:
            lots[LIGNES_LOTS[x["code"]]] += x["q"] * x["p"]
    for k, v in LOTS_METHODE.items():
        if abs(lots[k] - v) > TOL:
            erreur("lot %s : %s ≠ confirmé %s" % (k, lots[k], v))

    # 3. Couverture du relevé : chaque famille × portée non CONSERVER doit être citée par une ligne
    texte = " ".join(x["desc"] + " " + x["src"] for x in lignes)
    manquants = [f"{f} {p} {n}" for (f, p), n in sorted(portees.items())
                 if p != "CONSERVER" and f"{f} {p}" not in texte]
    if manquants:
        erreur("portées du relevé non chiffrées : %s" % manquants)

    # 4. Totaux
    materiel = round(sum(x["q"] * x["p"] for x in lignes), 2)
    heures_relevees = round(sum(x["q"] * x["h"] for x in lignes), 2)
    heures_chargees = int(math.ceil(heures_relevees / 10.0 - 1e-9) * 10)
    main_doeuvre = round(heures_chargees * TAUX, 2)
    total_acceo = round(materiel + main_doeuvre, 2)
    prix_vente = int(round(total_acceo * FACTEUR / 1000.0) * 1000)

    # 5. Ventilation par article (heures relevées ; MO au prorata des heures chargées pour que Σ total = total ACCEO)
    def cle(a):
        return [int(t) for t in a.split(".")]
    par = {}
    for x in lignes:
        a = par.setdefault(x["art"], {"materiel": 0.0, "heures": 0.0})
        a["materiel"] += x["q"] * x["p"]
        a["heures"] += x["q"] * x["h"]
    coef = heures_chargees / sum(a["heures"] for a in par.values())  # somme non arrondie
    par_article = OrderedDict()
    for art in sorted(par, key=cle):
        a = par[art]
        par_article[art] = {"materiel": round(a["materiel"], 2), "heures": round(a["heures"], 2),
                            "total": round(a["materiel"] + a["heures"] * coef * TAUX, 2)}

    # 6. Contrôles de cohérence
    if abs(materiel + heures_chargees * TAUX - total_acceo) > 0.005:
        erreur("incohérence materiel + heures_chargees × taux ≠ total_acceo")
    if prix_vente != int(round(total_acceo * FACTEUR, -3)):
        erreur("prix de vente mal arrondi")
    if abs(sum(v["total"] for v in par_article.values()) - total_acceo) > 0.01 * len(par_article):
        erreur("Σ par_article ≠ total_acceo")
    if abs(sum(v["materiel"] for v in par_article.values()) - materiel) > 0.05:
        erreur("Σ matériel par article ≠ matériel")

    sortie = OrderedDict([
        ("taux", TAUX), ("materiel", materiel), ("heures_relevees", heures_relevees),
        ("heures_chargees", heures_chargees), ("main_doeuvre", main_doeuvre), ("total_acceo", total_acceo),
        ("facteur", FACTEUR), ("prix_vente", prix_vente),
        ("lignes", [{k: (round(v, 4) if isinstance(v, float) else v) for k, v in x.items()} for x in lignes]),
        ("par_article", par_article), ("lots", lots),
    ])
    (ICI / "chiffrage.json").write_text(json.dumps(sortie, ensure_ascii=False, indent=1), encoding="utf-8")
    part_mo = main_doeuvre / total_acceo * 100
    print(f"lignes={len(lignes)} materiel={materiel:.2f} heures_relevees={heures_relevees:.2f} "
          f"heures_chargees={heures_chargees} main_doeuvre={main_doeuvre:.2f} total_acceo={total_acceo:.2f} "
          f"prix_vente={prix_vente} part_MO={part_mo:.1f}% h/100k$={heures_chargees / prix_vente * 1e5:.0f}")
    print("lots par équipe : " + ", ".join(f"{k}={v['lignes']}" for k, v in lots_meta.items()))


if __name__ == "__main__":
    main()
