"""Build separate schedule and diagram inventories; never silently aggregate scopes."""
import csv
import json
from collections import Counter
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
ROOT = Path(__file__).resolve().parents[1]
HEADERS=['Panneau','Circuits','Ampères','Pôles','Quantité','Désignation','Protection source','Feuille','Portée','Réserve','Repère']
data=json.loads((ROOT/'audit/E105-page09-records.json').read_text())
rows=[]
for rec in data['records']:
    rows.append([rec['parent'],rec['circuits'],rec['ampere'],rec['poles'],1,rec['label'],rec.get('protection','Non indiquée'),'E105-page09','PRINCIPAL' if rec['family']=='Principal' else 'DÉPART CÉDULE',rec.get('reserve',''),rec['id']])
with (ROOT/'S-1294-breakers-detailed.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(HEADERS);w.writerows(rows)
book=Workbook();readme=book.active;readme.title='Lire avant usage'
for line in ['S-1294 — Disjoncteurs par panneau','194 départs de cédule et 3 principaux, dont 10 repères avec réserve.',
             'Un départ multipolaire = un appareil. Les circuits constitutifs ne sont pas des disjoncteurs supplémentaires.',
             'LIBRE avec calibre inclus; ESPACE sans calibre exclu; ESPACE avec calibre conservé en réserve.',
             'Un circuit simple non fusionné = 1 pôle; 2P/3P reproduits tels qu’écrits. Principaux : pôles non indiqués.',
             'Schémas CM dans onglet séparé : ne pas additionner aux départs de cédule ni aux plans.',
             'Contradictions de tension, Icc, barres et principal : voir onglet Caractéristiques. Aucun arbitrage.',
             'Marque, série et caractéristiques non explicites restent à confirmer. Ce classeur ne remplace pas le relevé des appareils.']:
    readme.append([line])
detail=book.create_sheet('Départs et principaux');detail.append(HEADERS)
for row in rows:detail.append(row)
summary=book.create_sheet('Synthèse par panneau');summary.append(['Panneau','Portée','Ampères','Pôles','Quantité','Dont réserve'])
counts=Counter((r[0],r[8],r[2],str(r[3])) for r in rows)
for key,count in sorted(counts.items()):summary.append([*key,count,sum(bool(r[9]) for r in rows if (r[0],r[8],r[2],str(r[3]))==key)])
panels=book.create_sheet('Caractéristiques');panels.append(['Panneau','Tension cédule','Principal cédule','Barres cédule A','Icc cédule kA','Réserve / comparaison E102'])
metadata=[('PP1','347/600 V','600 A',600,14,'E102 Icc35kA contre 14kA en cédule.'),('PP1A','347/600 V','200 A',225,14,'E102 225A et Icc10kA; départ PP1 200A. Barres / principal à distinguer.'),('PS1','347/600 V','N/A',100,14,'E102 120/208V,225A,Icc10kA; incohérence de tension et capacité.'),('PS1A','120/208 V','100 A',100,14,'E102 Icc10kA contre 14kA.'),('PPU','347/600 V','N/A',225,14,'Pompe incendie 175A3P en cédule; schéma E102 pompe sur branche dédiée400A. À réconcilier.'),('PPU1','347/600 V','N/A',225,14,'E102100A,Icc25kA contre barres225A,Icc14kA.'),('PPU2','347/600 V','N/A',100,14,'E102 principal avec shunt trip; cédule N/A. À confirmer.'),('PSU1','120/208 V','N/A',225,14,'E102100A,Icc10kA contre barres225A,Icc14kA.'),('PSU1A','120/208 V','N/A; 100A adjacent',100,14,'Principal ambigu, non compté. E102 Icc10kA; départ PSU1 30A3P.')]
for row in metadata:panels.append(row)
diagrams=book.create_sheet('CM schémas non additionner');diagrams.append(HEADERS)
for key in ['E103','E104']:
    for rec in json.loads((ROOT/'audit'/f'{key}-records.json').read_text())['records']:
        if rec.get('ampere') is not None:
            diagrams.append([rec['parent'],rec['label'],rec['ampere'],'NON INDIQUÉ',1,rec['label'],'Non indiquée',key,'SCHÉMA — NE PAS ADDITIONNER',rec['reserve'],rec['id']])
for sheet in book:
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='24476A');cell.font=Font(color='FFFFFF',bold=True)
    for col in sheet.columns:
        width=min(75,max(14,max(len(str(c.value or '')) for c in col)+2))
        sheet.column_dimensions[get_column_letter(col[0].column)].width=width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    sheet.sheet_view.showGridLines=False
path=ROOT/'S-1294-breakers-by-panel.xlsx';book.save(path)
check=load_workbook(path,read_only=True,data_only=True)
assert check['Départs et principaux'].max_row==198
assert check['CM schémas non additionner'].max_row==87
seen=set()
for row in rows:
    if row[8]=='PRINCIPAL':continue
    circuits=row[1].split('-')
    assert len(circuits)==row[3]
    for circuit in circuits:
        key=(row[0],circuit)
        assert key not in seen, key
        seen.add(key)
print('Workbook verified: 197 schedule records; 86 independent diagram references; no circuit overlap.')
