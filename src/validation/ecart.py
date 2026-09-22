"""
Comparateur déterministe référence vs décompte IA — NOUVEAU module
(voir SPEC.md). Produit `ecart.md` : tableau ligne par ligne, écart total
≤ 5 %, écart par poste majeur ≤ 10 %, section « Items manquants » jamais
silencieuse, verdict PASS/FAIL **calculé**. Ce n'est PAS un texte
narratif : chaque chiffre est dérivé des deux fichiers d'entrée, et le
verdict FINAL reste à Francis (rappelé en en-tête du livrable).

CLI :
  python -m src.validation.ecart --reference ref.csv --ia compteurs.csv --sortie ecart.md

Entrées :
  --reference  Estimation de référence, CSV (séparateur , ou ; détecté)
               ou XLSX. Colonnes : poste, description, quantité
               (variantes d'accents/anglais acceptées : poste|famille,
               description|libelle|label, quantite|quantité|qty|count).
  --ia         Décompte IA par libellé :
               - CSV (label|libelle|description + quantite|count|qty) ;
               - JSON liste [{libelle|name|label, quantite|count}, ...]
                 ou dict {libelle: quantite} ;
               - JSON « compteurs » exporté par `src.releve verifier
                 --export-counters` ({"counters": [{name, elements}]} :
                 quantité = nombre d'éléments du compteur).

Appariement : par description normalisée (minuscules, accents retirés,
espaces compactées — `src.qpl.charte.normaliser`), jamais par position de
ligne. Une description de la référence sans contrepartie IA est un
**item manquant** : toujours listé, jamais silencieux. Une description IA
sans référence est un **item en trop** : listé à part.

Verdict calculé (indicatif) :
  FAIL si |écart total| > 5 %, ou si un poste majeur (luminaires,
  distribution, filage, conduits — familles de `src/qpl/charte.py`)
  dépasse 10 %, ou s'il existe au moins un item manquant ; sinon PASS.
Code de sortie : 0 = PASS, 1 = FAIL, 2 = erreur d'entrée.
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from src.qpl.charte import POSTES_MAJEURS, famille_pour_libelle, normaliser

SEUIL_TOTAL_PCT = 5.0
SEUIL_POSTE_MAJEUR_PCT = 10.0

_COLS_POSTE = ("poste", "famille", "lot")
_COLS_DESC = ("description", "libelle", "libellé", "label", "item")
_COLS_QTE = ("quantite", "quantité", "qty", "count", "qte", "qté")


class EcartEntreeError(ValueError):
    """Entrée illisible ou incomplète — refusée, jamais corrigée en silence."""


@dataclass
class Ligne:
    poste: str
    description: str
    reference: float | None   # None = absent de la référence (item en trop)
    ia: float | None          # None = absent du décompte IA (item manquant)

    @property
    def famille(self) -> str:
        """Famille de la ligne : le poste déclaré s'il est connu, sinon la
        classification du libellé par la charte."""
        poste_norm = normaliser(self.poste) if self.poste else ""
        familles_connues = set(POSTES_MAJEURS) | {
            "commandes", "secours", "prises", "alarme", "telecom", "mecanique", "autre",
        }
        if poste_norm in familles_connues:
            return poste_norm
        return famille_pour_libelle(self.description)


@dataclass
class Resultat:
    lignes: list[Ligne] = field(default_factory=list)
    manquants: list[Ligne] = field(default_factory=list)
    en_trop: list[Ligne] = field(default_factory=list)
    total_reference: float = 0.0
    total_ia: float = 0.0
    ecart_total_pct: float = 0.0
    par_poste: dict[str, dict] = field(default_factory=dict)
    verdict: str = "FAIL"
    causes: list[str] = field(default_factory=list)


# --------------------------------------------------------------------- #
# Lecture des entrées
# --------------------------------------------------------------------- #

def _parse_quantite(valeur: object) -> float:
    """Quantité numérique ; accepte « 1 234 », « 12,5 » et « 12.5 ».
    Toute autre valeur est refusée (jamais remplacée par 0 en silence)."""
    if valeur is None:
        raise EcartEntreeError("quantité vide")
    txt = str(valeur).strip().replace(" ", "").replace(" ", "").replace(",", ".")
    if not txt:
        raise EcartEntreeError("quantité vide")
    try:
        return float(txt)
    except ValueError as exc:
        raise EcartEntreeError(f"quantité non numérique : {valeur!r}") from exc


def _colonne(fieldnames: list[str], candidats: tuple[str, ...]) -> str | None:
    norm = {normaliser(f): f for f in fieldnames if f}
    for c in candidats:
        if normaliser(c) in norm:
            return norm[normaliser(c)]
    return None


def _lire_csv(path: Path) -> tuple[list[str], list[dict]]:
    echantillon = path.read_text(encoding="utf-8-sig")[:4096]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=",;\t")
    except csv.Error:
        dialecte = csv.excel
    with open(path, encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh, dialect=dialecte)
        if reader.fieldnames is None:
            raise EcartEntreeError(f"{path} : CSV sans en-tête")
        return list(reader.fieldnames), list(reader)


def lire_reference(path: Path) -> dict[str, Ligne]:
    """Référence CSV ou XLSX -> {description normalisée: Ligne}."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        fieldnames, rows = _lire_xlsx(path)
    else:
        fieldnames, rows = _lire_csv(path)
    c_poste = _colonne(fieldnames, _COLS_POSTE)
    c_desc = _colonne(fieldnames, _COLS_DESC)
    c_qte = _colonne(fieldnames, _COLS_QTE)
    if not c_desc or not c_qte:
        raise EcartEntreeError(
            f"{path} : colonnes obligatoires introuvables "
            f"(description parmi {_COLS_DESC}, quantité parmi {_COLS_QTE})"
        )
    out: dict[str, Ligne] = {}
    for i, row in enumerate(rows, start=2):
        desc = (row.get(c_desc) or "").strip()
        if not desc:
            continue
        cle = normaliser(desc)
        if cle in out:
            raise EcartEntreeError(f"{path} ligne {i} : description en double « {desc} »")
        out[cle] = Ligne(
            poste=(row.get(c_poste) or "").strip() if c_poste else "",
            description=desc,
            reference=_parse_quantite(row.get(c_qte)),
            ia=None,
        )
    if not out:
        raise EcartEntreeError(f"{path} : aucune ligne exploitable")
    return out


def _lire_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        entete = next(rows_iter)
    except StopIteration:
        raise EcartEntreeError(f"{path} : classeur vide") from None
    fieldnames = [str(c).strip() if c is not None else "" for c in entete]
    rows = []
    for valeurs in rows_iter:
        rows.append(
            {fieldnames[i]: ("" if v is None else v) for i, v in enumerate(valeurs) if i < len(fieldnames)}
        )
    wb.close()
    return fieldnames, rows


def lire_ia(path: Path) -> dict[str, float]:
    """Décompte IA -> {description normalisée: quantité}."""
    path = Path(path)
    if path.suffix.lower() == ".json":
        return _lire_ia_json(path)
    fieldnames, rows = _lire_csv(path)
    c_desc = _colonne(fieldnames, _COLS_DESC + ("name", "compteur"))
    c_qte = _colonne(fieldnames, _COLS_QTE)
    if not c_desc or not c_qte:
        raise EcartEntreeError(
            f"{path} : colonnes introuvables (libellé parmi {_COLS_DESC}, "
            f"quantité parmi {_COLS_QTE})"
        )
    out: dict[str, float] = {}
    for i, row in enumerate(rows, start=2):
        desc = (row.get(c_desc) or "").strip()
        if not desc:
            continue
        cle = normaliser(desc)
        if cle in out:
            raise EcartEntreeError(f"{path} ligne {i} : libellé en double « {desc} »")
        out[cle] = _parse_quantite(row.get(c_qte))
    if not out:
        raise EcartEntreeError(f"{path} : aucun compteur exploitable")
    return out


def _lire_ia_json(path: Path) -> dict[str, float]:
    data = json.loads(path.read_text(encoding="utf-8"))
    out: dict[str, float] = {}

    def ajoute(desc: str, qte: float) -> None:
        cle = normaliser(desc)
        if not cle:
            return
        if cle in out:
            raise EcartEntreeError(f"{path} : libellé en double « {desc} »")
        out[cle] = qte

    if isinstance(data, dict) and "counters" in data:
        # Export `src.releve verifier --export-counters` : quantité =
        # nombre d'éléments (marques) du compteur, toutes feuilles.
        for compteur in data["counters"]:
            ajoute(str(compteur.get("name", "")), float(len(compteur.get("elements", []))))
    elif isinstance(data, dict):
        for desc, qte in data.items():
            ajoute(str(desc), _parse_quantite(qte))
    elif isinstance(data, list):
        for item in data:
            desc = item.get("libelle") or item.get("label") or item.get("name") or item.get("description")
            qte = item.get("quantite") or item.get("count") or item.get("qty")
            if desc is None or qte is None:
                raise EcartEntreeError(f"{path} : item JSON sans libellé/quantité : {item!r}")
            ajoute(str(desc), _parse_quantite(qte))
    else:
        raise EcartEntreeError(f"{path} : JSON d'un type non pris en charge")
    if not out:
        raise EcartEntreeError(f"{path} : aucun compteur exploitable")
    return out


# --------------------------------------------------------------------- #
# Comparaison
# --------------------------------------------------------------------- #

def comparer(reference: dict[str, Ligne], ia: dict[str, float]) -> Resultat:
    res = Resultat()
    cles = sorted(set(reference) | set(ia))
    for cle in cles:
        if cle in reference:
            ligne = reference[cle]
            ligne.ia = ia.get(cle)
        else:
            ligne = Ligne(poste="", description=cle, reference=None, ia=ia[cle])
        res.lignes.append(ligne)
        if ligne.ia is None:
            res.manquants.append(ligne)
        if ligne.reference is None:
            res.en_trop.append(ligne)
        res.total_reference += ligne.reference or 0.0
        res.total_ia += ligne.ia or 0.0
        p = res.par_poste.setdefault(ligne.famille, {"reference": 0.0, "ia": 0.0})
        p["reference"] += ligne.reference or 0.0
        p["ia"] += ligne.ia or 0.0

    res.ecart_total_pct = _pct(res.total_ia, res.total_reference)
    for poste, p in res.par_poste.items():
        p["ecart_pct"] = _pct(p["ia"], p["reference"])

    # Verdict calculé
    ok = True
    if abs(res.ecart_total_pct) > SEUIL_TOTAL_PCT:
        ok = False
        res.causes.append(
            f"écart total {res.ecart_total_pct:+.1f} % au-delà du seuil ±{SEUIL_TOTAL_PCT:.0f} %"
        )
    for poste in POSTES_MAJEURS:
        p = res.par_poste.get(poste)
        if p is None:
            continue
        if abs(p["ecart_pct"]) > SEUIL_POSTE_MAJEUR_PCT:
            ok = False
            res.causes.append(
                f"poste majeur « {poste} » {p['ecart_pct']:+.1f} % au-delà du seuil "
                f"±{SEUIL_POSTE_MAJEUR_PCT:.0f} %"
            )
    if res.manquants:
        ok = False
        res.causes.append(f"{len(res.manquants)} item(s) de la référence sans contrepartie IA")
    res.verdict = "PASS" if ok else "FAIL"
    return res


def _pct(partie: float, total: float) -> float:
    if total == 0:
        return 0.0 if partie == 0 else float("inf")
    return (partie - total) / total * 100.0


def _fmt_qte(valeur: float | None) -> str:
    if valeur is None:
        return "—"
    return f"{valeur:g}"


def _fmt_pct(valeur: float) -> str:
    if valeur == float("inf"):
        return "+∞"
    return f"{valeur:+.1f} %"


# --------------------------------------------------------------------- #
# Livrable ecart.md
# --------------------------------------------------------------------- #

def ecrire_ecart_md(res: Resultat, reference_path: Path, ia_path: Path, sortie: Path) -> None:
    L: list[str] = []
    L.append("# Écart — relevé IA vs estimation de référence")
    L.append("")
    L.append("> **Verdict FINAL réservé à Francis** — le verdict ci-dessous est ")
    L.append("> **calculé** (indicatif) : l'IA mesure, elle ne décide pas.")
    L.append("")
    L.append(f"- Date : {date.today().isoformat()}")
    L.append(f"- Référence : `{reference_path}`")
    L.append(f"- Décompte IA : `{ia_path}`")
    L.append(f"- Seuils : écart total ≤ {SEUIL_TOTAL_PCT:.0f} %, poste majeur "
             f"(luminaires, distribution, filage, conduits) ≤ {SEUIL_POSTE_MAJEUR_PCT:.0f} %")
    L.append("")
    L.append(f"## Verdict calculé : **{res.verdict}**")
    L.append("")
    if res.causes:
        for cause in res.causes:
            L.append(f"- {cause}")
    else:
        L.append("- tous les critères calculés sont respectés")
    L.append("")
    L.append("## Synthèse")
    L.append("")
    L.append("| Indicateur | Seuil | Mesuré | Statut |")
    L.append("|---|---:|---:|---|")
    statut_total = "OK" if abs(res.ecart_total_pct) <= SEUIL_TOTAL_PCT else "DÉPASSE"
    L.append(f"| Écart total | ±{SEUIL_TOTAL_PCT:.0f} % | {_fmt_pct(res.ecart_total_pct)} | {statut_total} |")
    for poste in POSTES_MAJEURS:
        p = res.par_poste.get(poste)
        if p is None:
            L.append(f"| Poste majeur : {poste} | ±{SEUIL_POSTE_MAJEUR_PCT:.0f} % | — | sans ligne |")
            continue
        statut = "OK" if abs(p["ecart_pct"]) <= SEUIL_POSTE_MAJEUR_PCT else "DÉPASSE"
        L.append(f"| Poste majeur : {poste} | ±{SEUIL_POSTE_MAJEUR_PCT:.0f} % | {_fmt_pct(p['ecart_pct'])} | {statut} |")
    statut_manquants = "OK" if not res.manquants else "DÉPASSE"
    L.append(f"| Items manquants signalés | 0 | {len(res.manquants)} | {statut_manquants} |")
    L.append("")
    L.append("## Détail ligne par ligne")
    L.append("")
    L.append("| Poste | Description | Référence | IA | Écart | Écart % |")
    L.append("|---|---|---:|---:|---:|---:|")
    for ligne in res.lignes:
        ref = ligne.reference or 0.0
        ia = ligne.ia or 0.0
        ecart = ia - ref
        pct = _pct(ia, ref) if ligne.reference is not None else float("inf")
        L.append(
            f"| {ligne.famille} | {ligne.description} | {_fmt_qte(ligne.reference)} "
            f"| {_fmt_qte(ligne.ia)} | {ecart:+g} | {_fmt_pct(pct)} |"
        )
    L.append(
        f"| **Total** | | **{_fmt_qte(res.total_reference)}** | **{_fmt_qte(res.total_ia)}** "
        f"| **{res.total_ia - res.total_reference:+g}** | **{_fmt_pct(res.ecart_total_pct)}** |"
    )
    L.append("")
    L.append("## Items manquants")
    L.append("")
    L.append("Items présents dans l'estimation de référence sans contrepartie "
             "dans le décompte IA — **jamais silencieux** :")
    L.append("")
    if res.manquants:
        L.append("| Poste | Description | Quantité de référence |")
        L.append("|---|---|---:|")
        for ligne in res.manquants:
            L.append(f"| {ligne.famille} | {ligne.description} | {_fmt_qte(ligne.reference)} |")
    else:
        L.append("Aucun : chaque item de la référence a une contrepartie IA.")
    L.append("")
    L.append("## Items sans référence")
    L.append("")
    L.append("Items comptés par l'IA mais absents de l'estimation de référence "
             "(à arbitrer par Francis) :")
    L.append("")
    if res.en_trop:
        L.append("| Description | Quantité IA |")
        L.append("|---|---:|")
        for ligne in res.en_trop:
            L.append(f"| {ligne.description} | {_fmt_qte(ligne.ia)} |")
    else:
        L.append("Aucun : chaque compteur IA correspond à un item de la référence.")
    L.append("")
    sortie = Path(sortie)
    sortie.parent.mkdir(parents=True, exist_ok=True)
    sortie.write_text("\n".join(L), encoding="utf-8")


# --------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------- #

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="python -m src.validation.ecart",
        description="Comparateur déterministe référence vs décompte IA -> ecart.md",
    )
    parser.add_argument("--reference", type=Path, required=True,
                        help="Estimation de référence (CSV ou XLSX : poste, description, quantité)")
    parser.add_argument("--ia", type=Path, required=True,
                        help="Décompte IA (CSV ou JSON de compteurs par libellé)")
    parser.add_argument("--sortie", type=Path, required=True, help="ecart.md à écrire")
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        reference = lire_reference(args.reference)
        ia = lire_ia(args.ia)
    except (EcartEntreeError, OSError, json.JSONDecodeError) as exc:
        print(f"ERREUR : {exc}", file=sys.stderr)
        return 2
    res = comparer(reference, ia)
    ecrire_ecart_md(res, args.reference, args.ia, args.sortie)
    print(
        f"OK : {args.sortie} — verdict calculé {res.verdict} "
        f"(écart total {res.ecart_total_pct:+.1f} %, {len(res.manquants)} manquant(s), "
        f"{len(res.en_trop)} sans référence)"
    )
    return 0 if res.verdict == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
