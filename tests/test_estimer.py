"""src.estimer — synthetic plans only (no client data).

A synthetic plan PDF is drawn with pymupdf: walls, "receptacle" circles and
"fixture" squares with circuit tags, a title block with a sheet number and a
scale. Two such dossiers train a small model; the third is estimated end to end
and its outputs are checked (JSON readable by ecart, xlsx, .qpl verified).
"""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import numpy as np
import pymupdf
import pytest
from openpyxl import load_workbook

from src.estimer import conduits as K
from src.estimer import evaluate as EV
from src.estimer import export as E
from src.estimer import features as F
from src.estimer import model as M
from src.estimer import pages as P
from src.estimer import train as T
from src.estimer.__main__ import estimate, parse_pages
from src.estimer.families import INDETERMINE, LabelFamilyMap
from src.estimer.gold import DossierGold, GoldLine, GoldMark, SheetInfo, paper_to_feet
from src.releve import verifier
from src.validation import ecart

W_PT, H_PT = 1224.0, 792.0


def draw_plan(path: Path, seed: int, n_pages: int = 1, coloured_marks: bool = False) -> list[list[tuple]]:
    """Returns, per page, [(x_pt, y_pt, family)] of the drawn symbols."""
    rng = np.random.default_rng(seed)
    doc = pymupdf.open()
    truth = []
    for p in range(n_pages):
        page = doc.new_page(width=W_PT, height=H_PT)
        for x in range(60, 1000, 180):                           # walls
            page.draw_line((x, 40), (x, 740), width=1.5)
        for y in range(60, 740, 160):
            page.draw_line((40, y), (1000, y), width=1.5)
        syms = []
        xs = np.arange(100, 960, 45)
        ys = np.arange(90, 720, 45)
        cells = [(x, y) for x in xs for y in ys]
        pick = rng.choice(len(cells), 70, replace=False)
        for k, i in enumerate(pick):
            x, y = cells[i]
            x += rng.uniform(-6, 6); y += rng.uniform(-6, 6)
            if k % 2 == 0:
                page.draw_circle((x, y), 5, width=1.2)
                page.draw_line((x - 2, y - 3), (x - 2, y + 3), width=1)
                page.draw_line((x + 2, y - 3), (x + 2, y + 3), width=1)
                fam = "dispositif"
            else:
                page.draw_rect(pymupdf.Rect(x - 5, y - 5, x + 5, y + 5), width=1.2, fill=(0, 0, 0))
                fam = "luminaire"
            page.insert_text((x + 7, y + 10), "C,3", fontsize=5)
            syms.append((x, y, fam))
            if coloured_marks and k % 5 == 0:
                page.draw_circle((x + 30, y + 25), 4, color=(1, 0, 0), fill=(1, 0, 0))
        page.insert_text((1080, 700), "E101", fontsize=24)
        page.insert_text((1060, 740), 'ECHELLE 1/8" = 1\'-0"', fontsize=9)
        truth.append(syms)
    doc.save(path)
    return truth


def synthetic_gold(pdf: Path, truth: list[list[tuple]], name: str) -> DossierGold:
    k = P.WORK_WIDTH / W_PT
    sheets = [SheetInfo(f"S{i}", f"E10{i}", i, "plan", 96.0, W_PT, P.WORK_WIDTH, H_PT * k) for i in range(len(truth))]
    g = DossierGold(name, pdf, sheets)
    for i, syms in enumerate(truth):
        for x, y, fam in syms:
            g.marks.append(GoldMark(i, x * k, y * k, "PRISE" if fam == "dispositif" else "FIXTURE TYPE A", fam))
        g.lines.append(GoldLine(i, "COND 3/4", 4000.0, 100.0))
    g.styles = {"PRISE": (0, 22, -16776961), "FIXTURE TYPE A": (1, 26, -65536)}
    return g


@pytest.fixture(scope="module")
def trained(tmp_path_factory):
    d = tmp_path_factory.mktemp("estimer")
    golds = []
    for s in (1, 2):
        pdf = d / f"train{s}.pdf"
        golds.append(synthetic_gold(pdf, draw_plan(pdf, s), f"SYN-{s}"))
    model = T.train(golds, calibrate=False, log=lambda m: None)
    model.threshold, model.nms_radius = 0.5, 12.0
    test_pdf = d / "test.pdf"
    truth = draw_plan(test_pdf, 7, n_pages=2, coloured_marks=True)
    return d, model, test_pdf, truth, golds


# ---------------------------------------------------------------- unit
def test_label_family_map_learns_only_clear_majorities():
    m = LabelFamilyMap()
    for _ in range(4):
        m.add_vote("KS", "dispositif")
    m.add_vote("KS", "luminaire")
    m.add_vote("ZZ", "luminaire")
    assert m.family("KS") == "dispositif"          # 4/5 votes
    assert m.family("ZZ") == INDETERMINE            # a single vote is not enough
    assert m.family("PRISE") == "dispositif"        # keyword categoriser wins first
    assert LabelFamilyMap.from_json(m.to_json()).family("ks") == "dispositif"


def test_parse_scale_and_units():
    assert K.parse_scale('ÉCHELLE 1/8" = 1\'-0"') == 96.0
    assert K.parse_scale("ECHELLE 1:100") == 100.0
    assert K.parse_scale('1/4"=1\'-0" et 1/8"=1\'-0"') is None      # ambiguous: two scales
    assert K.parse_scale("aucune") is None
    # 72 pt of paper at 1/8" = 1'-0" is 8 ft; page 2997 px for 1224 pt of paper
    assert K.px_to_ft(72 * P.WORK_WIDTH / W_PT, P.WORK_WIDTH, W_PT, 96.0) == pytest.approx(8.0)
    assert paper_to_feet(1.0, 0.125, 1) == pytest.approx(8.0)             # Plan Expert imperial
    assert paper_to_feet(1.0, 96.0, 0) == pytest.approx(8.0)              # Plan Expert metric 1:96


def test_tree_length_and_ratio():
    xy = np.array([(0, 0), (10, 0), (10, 10), (0, 10)], float)
    length, edges = K.tree_length_px(xy)
    assert length == pytest.approx(30.0) and len(edges) == 3
    assert K.learn_ratio([(20, 10), (30, 10), (0, 5)]) == pytest.approx(2.5)
    assert not K.is_conduit_line("Distance 3") and K.is_conduit_line("COND 3/4")


def test_overlay_mask_and_candidates_skip_coloured_markup():
    rgb = np.full((200, 200, 3), 255, np.uint8)
    rgb[90:110, 90:110] = (230, 40, 40)          # coloured mark
    rgb[20:24, 20:60] = 0                         # black ink
    ov = P.overlay_mask(rgb)
    assert ov[100, 100] and not ov[22, 30]
    gray = np.full((200, 200), 255, np.uint8)
    gray[20:24, 20:60] = 0
    gray[95:105, 60:140] = 0                      # ink crossing the mark
    xs, ys = F.candidate_grid(gray, ov, 4)
    pts = set(zip(xs.tolist(), ys.tolist()))
    assert any(abs(x - 40) < 8 and abs(y - 22) < 8 for x, y in pts)
    assert not any(abs(x - 100) < 4 and abs(y - 100) < 4 for x, y in pts)


def test_dense_features_shape():
    gray = np.full((300, 400), 255, np.uint8)
    gray[100:140, 100:104] = 0
    X = F.DenseFeatures(gray).at(np.array([0, 102, 399]), np.array([0, 120, 299]))
    assert X.shape == (3, F.N_FEATURES) and np.isfinite(X).all()
    assert X[1].sum() > X[0].sum()


def test_match_is_one_to_one():
    pred = np.array([(0, 0), (1, 0), (50, 50)], float)
    gold = np.array([(0, 0), (100, 100)], float)
    assert M.match(pred, gold, 5) == 1
    assert M.match(pred[:0], gold, 5) == 0


def test_parse_pages():
    assert parse_pages("1-3,7") == [0, 1, 2, 6]
    assert parse_pages(None) is None


# ---------------------------------------------------------------- end to end
def test_model_learns_symbols_and_families(trained):
    d, model, test_pdf, truth, golds = trained
    assert set(model.classes) >= {"dispositif", "luminaire", "none"}
    assert model.family_style["dispositif"][1] == 0 and model.family_style["luminaire"][1] == 1
    assert model.conduit_ratio is not None and model.conduit_ratio > 0
    doc = pymupdf.open(test_pdf)
    page = P.load_page(doc, 0)
    det = model.detect(page)
    k = P.WORK_WIDTH / W_PT
    gxy = np.array([(x * k, y * k) for x, y, _ in truth[0]])
    readable = M.readable_centre(page.overlay)[gxy[:, 1].astype(int), gxy[:, 0].astype(int)] <= 0
    tp = M.match(np.array([(q.x, q.y) for q in det]).reshape(-1, 2), gxy[readable], 15)
    assert tp >= 0.6 * readable.sum()                 # finds most readable symbols
    assert len(det) <= 1.5 * len(gxy)                 # without flooding the page
    fams = Counter(q.family for q in det)
    assert fams["dispositif"] > 0 and fams["luminaire"] > 0


def test_estimate_end_to_end_outputs(trained, tmp_path):
    d, model, test_pdf, truth, golds = trained
    out = tmp_path / "out"
    data = estimate(test_pdf, out, model, log=lambda m: None)
    assert (out / "estimate.json").is_file() and (out / "releve.xlsx").is_file() and (out / "sheets.csv").is_file()
    assert len(data["sheets"]) == 2
    s0 = data["sheets"][0]
    assert s0["sheet"] == "E101"                                  # title-block number from the text layer
    assert s0["scale_ratio"] == 96.0 and s0["conduit_estimate_ft"] is not None
    # ecart-ready: src.validation.ecart reads the JSON directly
    ia = ecart.lire_ia(out / "estimate.json")
    assert sum(ia.values()) == sum(data["totals"].values()) > 0
    # xlsx: Relevé quantities sum to the detections
    wb = load_workbook(out / "releve.xlsx")
    q = sum(r[2] for r in wb["Relevé"].iter_rows(min_row=2, values_only=True))
    assert q == sum(data["totals"].values())
    assert "Conduits" in wb.sheetnames
    # Plan Expert project: one plan per page, element count == detections
    qpl = next((out / "planexpert").glob("*.qpl"))
    counts = verifier.compute_counts(verifier._parse(qpl))
    assert counts["n_plans"] == 2
    assert counts["n_counter_elements"] == sum(data["totals"].values())
    assert len(list((out / "planexpert").glob("*.png"))) == 2
    # the coloured marks drawn on the test plan are reported, not counted
    assert data["uncertainties"]["unreadable_markup_zones"] > 0


def test_sheets_csv_overrides(trained, tmp_path):
    d, model, test_pdf, truth, golds = trained
    csv_path = tmp_path / "sheets.csv"
    csv_path.write_text("page,name,scale_ratio,paper_width_pt,skip\n1,A-1,48,1224,\n2,,,,1\n", encoding="utf-8")
    data = estimate(test_pdf, tmp_path / "o", model, csv_path, qpl=False, log=lambda m: None)
    assert [s["sheet"] for s in data["sheets"]] == ["A-1"]
    assert data["sheets"][0]["scale_ratio"] == 48.0


def test_evaluation_scoring_on_synthetic_fold(trained, tmp_path):
    d, model, test_pdf, truth, golds = trained
    g = synthetic_gold(test_pdf, truth, "SYN-T")
    from src.estimer import pipeline
    sheets, _ = pipeline.run(test_pdf, model, keep_gray=False, log=lambda m: None)
    res = EV.score_position_dossier(g, sheets, {})
    e = res["detection"]["all_pages"]["R25"]
    assert e["gold"] == sum(len(t) for t in truth) and 0 < e["tp"] <= e["predicted"]
    assert res["occlusion"]["visible"]["gold"] + res["occlusion"]["occluded"]["gold"] == e["gold"]
    assert set(res["families"]) >= {"dispositif", "luminaire"}
    pool = EV.pooled([dict(res, dossier="SYN-T", gold_type="positions")])
    assert pool["detection"]["all_pages"]["R25"]["tp"] == e["tp"]
    EV.write_report(tmp_path / "REPORT.md", [dict(res, dossier="SYN-T", gold_type="positions", trained_on=["SYN-1"],
                                                  threshold=0.5, notes=[])], pool,
                    {"generated": "test", "commit": "test"})
    assert "SYN-T" in (tmp_path / "REPORT.md").read_text(encoding="utf-8")


def test_count_gold_scoring():
    g = DossierGold("CNT", Path("x.pdf"), [], has_positions=False)
    g.counts = [("E101", "PRISE", INDETERMINE, 10), ("E101", "FIXTURE TYPE A", INDETERMINE, 4)]
    sr = E.SheetResult(0, "E101", 100, 100, "rendered", None, "none", None, 0.0)
    sr.detections = [M.Detection(0, 1, 1, "dispositif", 0.9, 0.9, False)] * 8
    res = EV.score_count_dossier(g, [sr], LabelFamilyMap())
    assert res["families"]["dispositif"] == {"gold": 10, "predicted": 8, "count_error": -0.2}
    assert res["families"]["luminaire"]["predicted"] == 0
    assert res["totals"]["count_error"] == pytest.approx((8 - 14) / 14)
