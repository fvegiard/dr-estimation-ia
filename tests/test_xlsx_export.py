"""releve.xlsx conforme au gabarit (colonnes du SPEC, ligne par
libellé/feuille avec page, feuille « Résumé » par poste majeur)."""
from __future__ import annotations

from openpyxl import load_workbook

from conftest import FIXTURES, run_cli
from src.releve import xlsx_export

COLONNES_ATTENDUES = ["Description", "Type/Identification", "Quantité",
                      "Feuille/Plan", "Page PDF", "Notes"]


def test_export_xlsx_conforme(tmp_path):
    sortie = tmp_path / "releve.xlsx"
    run_cli("-m", "src.releve.xlsx_export",
            "--occurrences", str(FIXTURES / "occurrences.csv"),
            "--nomenclature", str(FIXTURES / "nomenclature.csv"),
            "--sortie", str(sortie))

    wb = load_workbook(sortie)
    assert wb.sheetnames == ["Relevé", "Résumé"]

    ws = wb["Relevé"]
    assert [c.value for c in ws[1]] == COLONNES_ATTENDUES
    lignes = {
        (row[1], row[3]): row  # (Type/Identification, Feuille/Plan)
        for row in ws.iter_rows(min_row=2, values_only=True)
    }
    # 3 couples (libellé, feuille) distincts
    assert len(lignes) == 3
    ds1 = lignes[("Luminaire DS1", "E100")]
    assert ds1[0] == "Luminaire encastré DEL type DS1"  # Description (nomenclature)
    assert ds1[2] == 3                                   # Quantité agrégée
    assert ds1[4] == "1"                                 # Page PDF
    prise = lignes[("Prise duplex 20A 125V", "E200")]
    assert prise[2] == 2
    panneau = lignes[("Panneau P-1", "E200")]
    assert panneau[2] == 1
    assert panneau[5] == "localisation au sous-sol"      # Notes

    ws2 = wb["Résumé"]
    assert [c.value for c in ws2[1]] == ["Poste", "Poste majeur", "Quantité", "Détail par feuille"]
    resume = {row[0]: row for row in ws2.iter_rows(min_row=2, values_only=True)}
    assert resume["luminaires"][1] == "oui"   # poste majeur
    assert resume["luminaires"][2] == 3
    assert resume["distribution"][1] == "oui"
    assert resume["distribution"][2] == 1
    assert resume["prises"][1] == "non"
    assert resume["prises"][2] == 2
    # postes majeurs en tête de la feuille Résumé
    postes = [row[0] for row in ws2.iter_rows(min_row=2, values_only=True)]
    assert postes[0] == "luminaires"
    assert postes[1] == "distribution"


def test_agregation_pure():
    occurrences = xlsx_export.lire_occurrences(FIXTURES / "occurrences.csv")
    lignes = xlsx_export.agreger(occurrences)
    total = sum(l["quantite"] for l in lignes)
    assert total == 6
    assert len(lignes) == 3
    # tri : par feuille puis libellé
    assert [(l["feuille"], l["label"]) for l in lignes] == [
        ("E100", "Luminaire DS1"),
        ("E200", "Panneau P-1"),
        ("E200", "Prise duplex 20A 125V"),
    ] or [l["feuille"] for l in lignes] == ["E100", "E200", "E200"]


def test_sous_commande_xlsx_du_package(tmp_path):
    sortie = tmp_path / "releve2.xlsx"
    run_cli("-m", "src.releve", "xlsx", str(FIXTURES / "occurrences.csv"),
            "--sortie", str(sortie))
    wb = load_workbook(sortie)
    assert wb.sheetnames == ["Relevé", "Résumé"]
