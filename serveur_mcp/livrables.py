"""Bordereau et main-d'œuvre issus des quantités contrôlées du serveur."""
from __future__ import annotations

import csv
import json
import math
from pathlib import Path

from openpyxl import load_workbook

from src.releve.xlsx_export import ecrire_xlsx

TAUX_MO = 124
RESERVE_PORTEE = (
    "Calcul sur les quantités agrégées, sans distinction INSTALLER / CONSERVER / "
    "ENLEVER / RELOCALISER. Ventiler les travaux et les heures par portée avant soumission."
)


def exporter_bom(travail: Path, destination: Path, nom: str, quantites: dict) -> dict:
    """Exporte les quantités multipliées; les heures exigent une source explicite.

    Colonnes facultatives de nomenclature : heures_unitaires et source_heures.
    Une valeur absente ou non sourcée reste non chiffrable, jamais zéro heure.
    """
    with (travail / "nomenclature.csv").open(encoding="utf-8-sig", newline="") as fichier:
        nomenclature = {r["label"]: r for r in csv.DictReader(fichier)}
    with (travail / "feuilles.csv").open(encoding="utf-8-sig", newline="") as fichier:
        pages = {r["feuille"]: r.get("page", "") for r in csv.DictReader(fichier)}
    lignes = []
    for feuille, articles in sorted(quantites["par_feuille"].items()):
        for label, compte in sorted(articles.items()):
            lignes.append({
                "label": label, "feuille": feuille, "quantite": compte["total"],
                "page_pdf": pages.get(feuille, ""),
                "notes": f'{compte["represente"]} repère(s) représenté(s); quantité après multiplicateurs.',
            })
    ecrire_xlsx(lignes, {label: r.get("description", "") for label, r in nomenclature.items()},
                destination / f"{nom}-BOM.xlsx")
    colonnes = ["label", "code", "famille", "description", "represente", "total"]
    with (destination / f"{nom}-BOM-QUANTITES.csv").open("w", encoding="utf-8", newline="") as fichier:
        writer = csv.DictWriter(fichier, fieldnames=colonnes)
        writer.writeheader()
        writer.writerows({
            **ligne, "description": nomenclature.get(ligne["label"], {}).get("description", ""),
        } for ligne in quantites["lignes"])

    postes = []
    for compte in quantites["lignes"]:
        label = compte["label"]
        article = nomenclature.get(label, {})
        valeur = (article.get("heures_unitaires") or "").strip()
        source = (article.get("source_heures") or "").strip()
        heures_unitaires = None
        if valeur:
            try:
                heures_unitaires = float(valeur.replace(",", "."))
            except ValueError as erreur:
                raise ValueError(f"{label} : heures_unitaires illisibles : {valeur!r}") from erreur
            if not math.isfinite(heures_unitaires) or heures_unitaires < 0:
                raise ValueError(f"{label} : heures_unitaires doit être fini et positif ou nul")
        chiffre = heures_unitaires is not None and bool(source)
        heures = (round(compte["total"] * heures_unitaires, 4)
                  if heures_unitaires is not None and source else None)
        postes.append({
            **compte, "heures_unitaires": heures_unitaires if chiffre else None,
            "source_heures": source, "heures": heures,
            "cout_mo": round(heures * TAUX_MO, 2) if heures is not None else None,
            "reserve": "" if chiffre else "Heures unitaires sourcées manquantes : poste non chiffrable.",
        })
    complet = all(p["cout_mo"] is not None for p in postes)
    sous_total = round(sum(p["cout_mo"] for p in postes if p["cout_mo"] is not None), 2)
    estimation = {
        "taux_horaire": TAUX_MO, "devise": "CAD", "complet": complet,
        "reserve_portee": RESERVE_PORTEE,
        "cout_mo_total": sous_total if complet else None,
        "sous_total_postes_chiffres": sous_total, "postes": postes,
    }
    (destination / f"{nom}-ESTIMATION-MO.json").write_text(
        json.dumps(estimation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    bilan = f"{sous_total:.2f} $ CAD" if complet else "non chiffrable en entier (heures sourcées manquantes)"
    rapport = [
        "# Quantités et estimation de main-d'œuvre", "",
        f"Taux de référence : {TAUX_MO} $ CAD/h.",
        f"Repères représentés : {quantites['represente']}; quantité à soumissionner : {quantites['total']}.",
        f"Coût MO total : {bilan}.", "",
        f"Réserve de portée : {RESERVE_PORTEE}", "",
        "| Article | Quantité | Heures unitaires | Heures totales | Coût MO CAD | Source / réserve |",
        "|---|---:|---:|---:|---:|---|",
    ]
    for poste in postes:
        def texte(valeur):
            return "" if valeur is None else str(valeur).replace("|", "/").replace("\n", " ")
        rapport.append("| " + " | ".join(texte(v) for v in (
            poste["label"], poste["total"], poste["heures_unitaires"], poste["heures"],
            poste["cout_mo"], poste["reserve"] or poste["source_heures"],
        )) + " |")
    (destination / f"{nom}-ESTIMATION-MO.md").write_text("\n".join(rapport) + "\n", encoding="utf-8")

    classeur = load_workbook(destination / f"{nom}-BOM.xlsx")
    feuille = classeur.create_sheet("Main-d'œuvre")
    feuille.append(["Article", "Quantité", "Heures unitaires", "Heures totales", "Taux CAD/h",
                    "Coût MO CAD", "Source des heures", "Réserve"])
    for poste in postes:
        feuille.append([poste["label"], poste["total"], poste["heures_unitaires"], poste["heures"],
                        TAUX_MO, poste["cout_mo"], poste["source_heures"], poste["reserve"]])
    feuille.freeze_panes = "A2"
    feuille.append([])
    feuille.append(["Réserve de portée", RESERVE_PORTEE])
    classeur.save(destination / f"{nom}-BOM.xlsx")
    classeur.close()
    return estimation
