# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl", "pydantic>=2"]
# ///
# Exécution : .venv/Scripts/python.exe outils/controler_bom_campagne.py
"""Contrôle croisé des BOM, classeurs, coûts MO et fichiers QPL de la campagne."""
from __future__ import annotations

import csv
import json
from collections import Counter
from contextlib import closing
from datetime import datetime, timezone
from pathlib import Path
from typing import Final
from xml.etree import ElementTree
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field, TypeAdapter, ValidationError

RACINE: Final = Path(__file__).resolve().parents[1]
BASE: Final = RACINE / "travail-mcp/campagne-25-plans"


class Modele(BaseModel):
    model_config = ConfigDict(frozen=True)


class Detail(Modele):
    feuille: str
    repere: str
    code: str
    qte: float


class Quantite(Modele):
    label: str
    code: str
    represente: float
    total: float


class PosteMO(Quantite):
    heures_unitaires: float | None
    source_heures: str
    heures: float | None
    cout_mo: float | None


class EstimationMO(Modele):
    taux_horaire: float
    complet: bool
    cout_mo_total: float | None
    sous_total_postes_chiffres: float
    postes: list[PosteMO]


class Production(Modele):
    produit: bool = False
    reperes: int | None = None
    raison: str | None = None


class Texte(Modele):
    type: str
    text: str = ""


class Reponse(Modele):
    content: list[Texte] = Field(default_factory=list)
    structuredContent: Production | None = None


class Preuve(Modele):
    reponse: Reponse


class Cas(Modele):
    id: str


class Manifeste(Modele):
    cas: list[Cas]


class Controle(Modele):
    id: str
    statut: str
    erreurs: list[str]
    reperes: int | None = None
    lignes_bom: int | None = None
    quantite_totale: float | None = None
    taux_mo: float | None = None
    postes_mo_non_chiffres: int | None = None
    racine_qpl: str | None = None


def controler_produit(base: Path, cas: Cas, production: Production) -> Controle:
    """Relit les artefacts d'un produit et conserve chaque écart détecté."""
    erreurs: list[str] = []

    def verifier(condition: bool, description: str) -> None:
        if not condition:
            erreurs.append(description)

    sortie = base / cas.id / "sortie"
    png = sortie / "png"
    with (png / f"{cas.id}-BOM.csv").open(encoding="utf-8-sig", newline="") as fichier:
        details = TypeAdapter(list[Detail]).validate_python(list(csv.DictReader(fichier)))
    with (png / f"{cas.id}-BOM-QUANTITES.csv").open(encoding="utf-8-sig", newline="") as fichier:
        quantites = TypeAdapter(list[Quantite]).validate_python(list(csv.DictReader(fichier)))
    mo = EstimationMO.model_validate_json(
        (png / f"{cas.id}-ESTIMATION-MO.json").read_text(encoding="utf-8"))
    racine_qpl = ElementTree.parse(sortie / f"{cas.id}.qpl").getroot().tag
    verifier(racine_qpl == "QuoterPlanSession", "Racine XML QPL différente de QuoterPlanSession")
    verifier(len(details) == production.reperes, "Nombre de lignes BOM différent des repères MCP")
    verifier(len({(r.feuille, r.repere) for r in details}) == len(details), "Repères BOM en doublon")
    verifier(all(r.feuille and r.repere for r in details), "Feuille ou repère BOM vide")
    verifier(len({q.label for q in quantites}) == len(quantites), "Libellés agrégés en doublon")
    verifier(sum(q.represente for q in quantites) == len(details), "Total représenté différent du BOM")
    comptes_detail = Counter(r.code for r in details)
    comptes_agreges: Counter[str] = Counter()
    for quantite in quantites:
        comptes_agreges[quantite.code] += quantite.represente
    verifier(comptes_detail == comptes_agreges, "Comptes par code différents entre BOM et quantités")
    attendues = {q.label: q.total for q in quantites}
    verifier(mo.taux_horaire == 124, "Taux JSON MO différent de 124 CAD/h")
    verifier({p.label: p.total for p in mo.postes} == attendues, "Quantités JSON MO différentes du CSV")
    verifier(len({p.label for p in mo.postes}) == len(mo.postes), "Postes JSON MO en doublon")
    with closing(load_workbook(png / f"{cas.id}-BOM.xlsx", read_only=True, data_only=True)) as classeur:
        releve: Counter[str] = Counter()
        for ligne in classeur["Relevé"].iter_rows(min_row=2, values_only=True):
            if ligne[1]:
                releve[str(ligne[1])] += float(ligne[2])
        verifier(dict(releve) == attendues, "Quantités XLSX Relevé différentes du CSV agrégé")
        resume = sum(float(ligne[2]) for ligne in classeur["Résumé"].iter_rows(
            min_row=2, values_only=True) if ligne[0])
        verifier(resume == sum(attendues.values()), "Total XLSX Résumé différent du CSV agrégé")
        lignes_mo = [ligne for ligne in classeur["Main-d'œuvre"].iter_rows(
            min_row=2, values_only=True) if ligne[0] and ligne[0] != "Réserve de portée"]
        excel_mo = {str(ligne[0]): ligne for ligne in lignes_mo}
        verifier(len(excel_mo) == len(lignes_mo), "Articles XLSX MO en doublon")
        verifier(set(excel_mo) == set(attendues), "Articles XLSX MO différents du CSV agrégé")
        for poste in mo.postes:
            ligne = excel_mo.get(poste.label)
            if ligne is not None:
                verifier(ligne[1:6] == (poste.total, poste.heures_unitaires, poste.heures, 124, poste.cout_mo),
                         f"{poste.label} : quantité, heures, taux ou coût XLSX différents du JSON")
                verifier((ligne[6] or "") == poste.source_heures,
                         f"{poste.label} : source d'heures XLSX différente du JSON")
            if poste.heures_unitaires is None or not poste.source_heures:
                verifier(poste.heures_unitaires is None and poste.heures is None and poste.cout_mo is None,
                         f"{poste.label} : heures ou coûts inventés sans heures sourcées")
            else:
                heures = round(poste.total * poste.heures_unitaires, 4)
                verifier(poste.heures == heures and poste.cout_mo == round(heures * 124, 2),
                         f"{poste.label} : calcul MO incorrect")
    complet = all(p.cout_mo is not None for p in mo.postes)
    sous_total = round(sum(p.cout_mo for p in mo.postes if p.cout_mo is not None), 2)
    verifier(mo.complet == complet, "Statut MO complet incohérent")
    verifier(mo.sous_total_postes_chiffres == sous_total, "Sous-total MO incohérent")
    verifier(mo.cout_mo_total == (sous_total if complet else None), "Total MO incohérent")
    return Controle(id=cas.id, statut="ECHEC" if erreurs else "CONFORME",
                    erreurs=erreurs, reperes=production.reperes, lignes_bom=len(details),
                    quantite_totale=sum(attendues.values()), taux_mo=mo.taux_horaire,
                    postes_mo_non_chiffres=sum(p.cout_mo is None for p in mo.postes),
                    racine_qpl=racine_qpl)


def main() -> int:
    """Conserve tous les verdicts; les cas non produits sont explicitement hors contrôle BOM."""
    manifeste = Manifeste.model_validate_json((BASE / "manifest.json").read_text(encoding="utf-8"))
    controles: list[Controle] = []
    for cas in manifeste.cas:
        chemin = BASE / "preuves" / f"{cas.id}.json"
        if not chemin.exists():
            controles.append(Controle(id=cas.id, statut="EN_ATTENTE", erreurs=[]))
            continue
        try:
            preuve = Preuve.model_validate_json(chemin.read_text(encoding="utf-8"))
            production = preuve.reponse.structuredContent or Production.model_validate_json(
                "\n".join(c.text for c in preuve.reponse.content if c.type == "text"))
            controle = (controler_produit(BASE, cas, production) if production.produit else
                        Controle(id=cas.id, statut="NON_PRODUIT", erreurs=[]))
        except (OSError, ValidationError, KeyError, ValueError, TypeError, ElementTree.ParseError, BadZipFile) as erreur:
            controle = Controle(id=cas.id, statut="ECHEC", erreurs=[f"{type(erreur).__name__}: {erreur}"])
        controles.append(controle)
    bilan = dict(Counter(c.statut for c in controles))
    cible = RACINE / "docs/assurance-empirique-25-plans/controle-bom.json"
    cible.parent.mkdir(parents=True, exist_ok=True)
    cible.write_text(json.dumps({
        "genere_le": datetime.now(timezone.utc).isoformat(),
        "portee": "Concordance des artefacts produits; ne valide ni exhaustivité du relevé ni acceptation Plan Expert.",
        "bilan": bilan, "cas": [c.model_dump() for c in controles],
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(bilan, ensure_ascii=False))
    return 1 if any(c.erreurs for c in controles) else 0


if __name__ == "__main__":
    raise SystemExit(main())
