"""Consolidation locale selon méthode openpyxl existante. Aucun total inter-portées."""
import csv,json,hashlib,datetime,collections,re,zipfile,math
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
D=Path(__file__).resolve().parents[1];W=D.parents[3]
keys=['E200','E202','E203','E204','E205','E206','E300','E301']
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def csvout(p,head,rows):
 with p.open('w',encoding='utf-8-sig',newline='') as f:w=csv.writer(f);w.writerow(head);w.writerows(rows)
downloads=json.loads((W/'evidence/source-downloads.json').read_text('utf-8'));sourcebyname={x['name']:x for x in downloads}
scans=[];seen={}
for p in sorted((D/'sources').glob('doc*.png'),key=lambda p:('(' in p.name,p.name)):
 h=sha(p);canonical=seen.setdefault(h,p.name)
 n=p.name
 if '82318 - 1' in n or '83916' in n:cat='Tableau éclairage urgence';sem='Même tableau que doc20241218082318 - 1 / doc20241218083916; pas un second inventaire.'
 elif '82318 - 2' in n or '84050' in n:cat='Alarme incendie, tableau et notes manuscrites';sem='Même contenu visuel que doc20241218082318 - 2 / doc20241218084050.'
 elif '82318 - 3' in n or '84335' in n:cat='Cédule luminaires avec quantités manuscrites';sem='Même tableau que doc20241218082318 - 3 / doc20241218084335; quantités humaines non utilisées.'
 elif '82318 - 4' in n or '84436' in n:cat='Feuille humaine génératrice / transferts';sem='Même portée que doc20241218082318 - 4 / doc20241218084436; pas une implantation supplémentaire.'
 elif '82318 - 5' in n:cat='Distribution manuscrite, feuille 1/3';sem='Relevé humain à comparer après gel, non utilisé comme source de quantités.'
 elif '82318 - 6' in n:cat='Distribution manuscrite, feuille 2/3';sem='Relevé humain à comparer après gel, non utilisé comme source de quantités.'
 else:cat='Distribution manuscrite, feuille 3/3, recharge';sem='Relevé humain à comparer après gel, non utilisé comme source de quantités.'
 scans.append([n,cat,h,canonical,'DOUBLON BINAIRE' if canonical!=n else 'FICHIER UNIQUE',sem,'Classé visuellement; pas de validation des quantités humaines'])
csvout(D/'S-1294-classement-scans.csv',['Fichier','Classe','SHA-256','Exemplaire canonique','Statut','Doublon de contenu / portée','Utilisation'],scans)
reserves=[
 ['R01','PL206 / PL208','E104, page 8','Repères PL-206 et PL-208 présents au niveau 02 et à nouveau au niveau 01.','Ne pas renommer PL106/108 ni fusionner automatiquement. Clarification de conception requise.','conflit-E104-PL206-208.jpg'],
 ['R02','PS1','E102 p6 / E105 p9','E102 : 120/208 V, 225 A, Icc 10 kA. E105 : 347/600 V, barres 100 A, Icc 14 kA, principal N/A.','Tension et caractéristiques non arbitrées; ne pas appliquer 347/600 V aux prises PS1 par défaut.','conflit-E102-PS1.jpg; conflit-E105-PS1.jpg'],
 ['R03','PP1','E102 p6 / E105 p9','Icc 35 kA sur E102, 14 kA sur E105.','Conserver les deux exigences; validation concepteur avant sélection du matériel.','conflit-E102-PP1.jpg; conflit-E105-PP1.jpg'],
 ['R04','Plans / schémas / types','E200, E202–E206, E102–E105, E400','Plusieurs représentations des mêmes équipements.','Aucun total immeuble par somme de ces feuilles. Les types ne sont pas multipliés par un nombre de logements.','Plans entiers et colonnes Portée/Parent'],
 ['R05','Batterie 2 phares','E100 / E200','La légende du symbole indique 320 W et la remarque 144 W.','Capacité réservée; un ensemble comprend les deux phares.','E100-batteries-complet.jpg'],
 ['R06','Massif 6 × 100 mm','E300 / E301','550 × 405 sur E300, 625 × 455 sur le tableau E301 pour 6 conduits de 100 mm.','Dimensions non arbitrées; le 550 × 405 correspond à la colonne 75 mm du tableau.','E300-canalisation.jpg; E301-detail-1.jpg'],
 ['R07','Symboles hors emprise','E200','18 symboles C isolés hors emprise représentée.','Localisation ou doublon possible à clarifier; ne pas les ajouter aux 64 C implantés.','E200-float-grid.jpg; E200-float2-grid.jpg'],
 ['R08','Symboles ambigus','E200 / E202 / E204 / E206','Carrés diagonaux des commerces; deux cercles près MRA; commande S/T; carré P avec triangle plein.','Identification réservée. Ni suppression ni conversion arbitraire en appareil connu.','Vignettes liées aux ID'],
 ['R09','Éclairage logements','E108 / E203–E206','Référence JUNO contient 1000LM, colonne de flux indique 2000 lm.','Référence et divergence conservées; choix de couleur par architecte/propriétaire.','E108-catalog.jpg'],
 ['R10','Fonction CO','E100 / E204','KIDDE P12040CA pour L. CO conditionnel : E100 BRK SC7010B, note E204 BRK SC9120.','Applicabilité garage/gaz à confirmer, sans ajout de détecteur; références non arbitrées.','E204-notes.jpg'],
 ['R11','Mesures','Toutes feuilles','Les câbles et conduits ne sont pas métrés; les longueurs E300 sont les mentions approximatives du dessin.','Pas de conversion automatique en quantités d’achat.','Notes intégrées'],
 ['R12','Protocole','Reprise du 2026-10-02','Aucun QPL ouvert; exposition antérieure aux scans annotés déclarée.','Aucune revendication de protocole intégralement aveugle. Comparaison estimateur non réalisée.','PROTOCOL-EXPOSURE.md'],
 ['R13','E201','Page réservée','E201 exclue de la lecture, de la modification et du paquet de reprise.','La livraison S-1294 comprend 22 pages; E201 reste sous responsabilité du writer réservé.','Périmètre explicite du mandat'],
 ['R14','État de revue','Huit pages nouvelles','PASS bornés : E203 v2, E204 v3, E206 v2, E300/E301 v3. E200/E202 : D1–D6 corrigés pour revue delta. E205 : lettre f corrigée pour revue delta.','Revue bornée; aucun résultat automatique assimilé à une certification métier.','STATUS-S-1294.md'],
 ['R15','Instance retirée E202-0269','E202 source page 15','La prétendue prise était sur le thermostat 0285; la prise locale réelle est déjà 0270.','Suppression explicite de 0269, sans remplacement ni renumérotation. 19 → 18 prises doubles, 357 → 356 lignes E202. Aucune conclusion d’exhaustivité globale.','delta-E200-E202/D2-instance-retiree.json; E202-D2-local-grid.jpg']]
csvout(D/'S-1294-reserves-reprise.csv',['Réserve','Objet','Sources','Constat','Conséquence','Preuve locale'],reserves)
master=[];data={};HEAD=None
for k in keys:
 d=json.loads((D/'audit'/f'{k}-records.json').read_text('utf-8'));data[k]=d
 p=D/f'S-1294-{k}-{d["type"]}.csv';rows=list(csv.reader(p.open(encoding='utf-8-sig',newline='')))
 if HEAD is None:HEAD=rows[0]
 assert rows[0]==HEAD
 master += [[k,*r] for r in rows[1:]]
csvout(D/'S-1294-reprise-detail.csv',['Feuille',*HEAD],master)
book=load_workbook(D/'S-1294-breakers-by-panel.xlsx');original={s.title:list(s.values) for s in book}
def table(name,head,rows):
 s=book.create_sheet(name);s.append(head)
 for r in rows:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.sheet_view.showGridLines=False
 for c in s[1]:c.fill=PatternFill('solid',fgColor='24476A');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[1].height=32
 for row in s.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if c.row%2==0:c.fill=PatternFill('solid',fgColor='EFF4F8')
  s.row_dimensions[row[0].row].height=44 if len(head)>8 else 38
 for i,col in enumerate(s.columns,1):
  width=min(72,max(14,max(len(str(c.value or '')) for c in col)+2));s.column_dimensions[get_column_letter(i)].width=width
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
 return s
accueil=book.create_sheet('Accueil reprise',0)
intro=[['S-1294 — classeur de reprise local'],['22 pages : 14 conservées + 8 nouvelles. E201 exclue.'],['Version du '+datetime.datetime.now(datetime.timezone.utc).isoformat()],['Départs de cédule hérités : 197 lignes, sans revalidation complète des cédules.'],['86 références de schémas CM séparées; ne pas ajouter aux départs de cédule.'],['Relevés huit feuilles : '+str(len(master))+' repères, comprenant symboles et références documentaires.'],['P-… : onglet de chaque panneau; CM-… : schémas seulement.'],['Les types PL restent distincts; aucun multiplicateur de logements.'],['Réserves PS1, PP1 et PL206/208 conservées; voir Réserves reprise.'],['Le calibre 30 A des sectionneurs E202 ne devient pas une quantité de disjoncteurs.'],['Les champs non transcrits restent vides; aucune valeur n’est déduite d’un numéro de circuit.'],['Les contrôles CSV/JSON/classeur et SHA-256 ne certifient pas le métier.'],['Références de produits : prescriptions source, pas sélection d’achat.'],['Aucun QPL ouvert; exposition aux scans humains déclarée, pas de revendication d’aveugle intégral.']]
for r in intro:accueil.append(r)
accueil.column_dimensions['A'].width=125
for row in accueil:
 for c in row:c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=12)
 accueil.row_dimensions[row[0].row].height=32
accueil['A1'].font=Font(name='Calibri',size=20,bold=True,color='24476A');accueil.sheet_view.showGridLines=False
numeric={'Qté','X pixels source','Y pixels source','Pastille','Ampères','Pôles'}
expected=[]
for row in master:
 nr=[row[0]]
 for h,v in zip(HEAD,row[1:]):
  if v and h in numeric:
   try:v=float(v);v=int(v) if v.is_integer() else v
   except ValueError:pass
  nr.append(v if v!='' else None)
 expected.append(nr)
inv=table('Relevés huit feuilles',['Feuille',*HEAD],expected)
resrows=[[r[0],r[1],r[2],r[8],r[7]] for r in master if r[7]]
table('Réserves par repère',['Feuille','ID','Famille','Parent','Réserve'],resrows)
table('Réserves reprise',['Réserve','Objet','Sources','Constat','Conséquence','Preuve locale'],reserves)
table('Classement scans',['Fichier','Classe','SHA-256','Exemplaire canonique','Statut','Doublon / portée','Usage'],scans)
source_rows=[]
for p in sorted((D/'sources').glob('*.png')):
 assert not p.name.endswith(' - 14.png')
 meta=sourcebyname[p.name];assert sha(p)==meta['sha256'];source_rows.append([p.name,sha(p),p.stat().st_size,meta['remote']])
table('Sources',['Fichier source','SHA-256','Octets','Chemin SharePoint'],source_rows)
details=original['Départs et principaux'];caracs=original['Caractéristiques'];diagram=original['CM schémas non additionner']
for panel in sorted(set(r[0] for r in details[1:])):
 s=table('P-'+panel,list(details[0]),[list(r) for r in details[1:] if r[0]==panel])
 s.append([]);s.append(['CARACTÉRISTIQUES ET RÉSERVES — deux sources conservées'])
 s.append(list(caracs[0]));s.append(next(list(r) for r in caracs[1:] if r[0]==panel))
 for reserve in reserves:
  if reserve[1]==panel:
   s.append(['RÉSERVE '+reserve[0],reserve[3],reserve[4]])
 refs=[r for r in master if r[8]==panel or (r[2]=='Panneau' and r[3]==panel)]
 s.append([]);s.append(['RENVOIS D’IMPLANTATION — NON ADDITIONNABLES AUX DÉPARTS'])
 s.append(['Feuille','ID','Famille','Désignation','Qté symboles','Portée','Modèle','Réserve','Parent'])
 for r in refs:s.append(r[:9])
 for row in s.iter_rows():
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 s.auto_filter.ref=f'A1:K{1+sum(r[0]==panel for r in details[1:])}'
for cm in sorted(set(r[0] for r in diagram[1:])):table(cm,list(diagram[0]),[list(r) for r in diagram[1:] if r[0]==cm])
families=sorted(set(r['family'] for k in keys[2:6] for r in data[k]['records']))
types=[]
for k in keys[2:6]:
 for p in sorted(set(r['parent'] for r in data[k]['records'])):
  types.append((k,p))
type_rows=[[f,*[sum(r['quantity'] for r in data[k]['records'] if r['parent']==p and r['family']==f) for k,p in types],'Par dessin type, non multiplié; renvois inclus selon leur portée.'] for f in families]
typesheet=table('PL par type',['Famille',*[k+' / '+p.removeprefix('Logement type ') for k,p in types],'Portée'],type_rows)
typesheet.append(['RÉSERVE R01 : PL206 / PL208 répétés sur E104; aucune correction de repère déduite.'])
page_rows=[[k,len(d['records']),d['title'],d['scope'],f'S-1294-{k}-releve.jpg',f'S-1294-{k}-{d["type"]}.csv',f'STAM23A - E - Pour construction - 2024-10-08 - {d["page"]}.png'] for k,d in data.items()]
table('Pages reprise',['Feuille','Repères','Titre','Portée','Plan entier annoté','CSV','Source PNG'],page_rows)
# Ajuster les lignes des nouveaux onglets au texte; les cinq onglets historiques restent intacts.
from PIL import ImageFont
layout_font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',15)
for sheet in book:
 if sheet.title in original or sheet.title=='Accueil reprise':continue
 for row in sheet:
  needed=1
  for cell in row:
   if cell.value is None:continue
   width_px=max(40,sheet.column_dimensions[cell.column_letter].width*7-12)
   lines=0
   for para in str(cell.value).split('\n'):
    used=0;lines+=1
    for word in para.split(' '):
     size=layout_font.getlength(word+' ')
     if used and used+size>width_px:lines+=1;used=0
     if size>width_px:lines+=int(size//width_px);size%=width_px
     used+=size
   needed=max(needed,lines)
  sheet.row_dimensions[row[0].row].height=max(30 if row[0].row==1 else 24,min(400,needed*14+8))
  if str(row[0].value or '').startswith(('CARACTÉRISTIQUES','RENVOIS D’','RÉSERVE')):
   for cell in row:
    cell.fill=PatternFill('solid',fgColor='FFF0CA');cell.font=Font(bold=True,color='664200');cell.alignment=Alignment(wrap_text=True,vertical='top')
out=D/'S-1294-reprise-par-panneau.xlsx';book.save(out)
# Vérification directe du livrable, pas de tests simulant des sources.
reopened=load_workbook(out,read_only=True,data_only=False)
assert list(reopened['Relevés huit feuilles'].values)[1:]==[tuple(r) for r in expected]
for name,rows in original.items():assert list(reopened[name].values)==rows,name
for panel in sorted(set(r[0] for r in details[1:])):
 rows=[r for r in details[1:] if r[0]==panel];actual=list(reopened['P-'+panel].values)[1:1+len(rows)];assert actual==rows
with zipfile.ZipFile(out) as z:assert z.testzip() is None
assert len(seen)==11
verification={'workbook':out.name,'sha256':sha(out),'worksheets':len(reopened.sheetnames),'reprise_rows':len(expected),'baseline_schedule_rows':len(details)-1,'baseline_diagram_rows':len(diagram)-1,'baseline_sheets_preserved_cell_for_cell':list(original),'scan_files':len(scans),'scan_unique_sha256':len(seen),'sources':len(source_rows),'checks':f'Classeur relu depuis le fichier enregistré; {len(expected)} lignes comparées aux CSV; onglets existants et départs par panneau identiques; archive XLSX intègre. Aucun contrôle métier automatique revendiqué.'}
(D/'VERIFICATION-CLASSEUR-REPRISE.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2),'utf-8');print(json.dumps(verification,ensure_ascii=False))
