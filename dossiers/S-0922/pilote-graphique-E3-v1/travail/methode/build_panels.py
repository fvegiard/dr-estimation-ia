"""Transcribe visible panel schedule cells, preserving source status and doubts."""
import csv
import json
import sys
from collections import Counter
from pathlib import Path
ROOT = Path(__file__).resolve().parent
if ROOT.name == 'methode':
    ROOT = ROOT.parent
(ROOT / 'data').mkdir(parents=True, exist_ok=True)
sys.path.insert(0,str(Path(__file__).parent/'.venv/lib'))
import openpyxl
from openpyxl.styles import Font, PatternFill

rows=[]
def add(panel,circuits,amps,description,new=False,reserve=''):
    nums=[int(n) for n in circuits.split('/')]
    rows.append(dict(panel=panel,circuits=circuits,ampere=amps,poles=len(nums),quantity=1,
        description=description,protection_source='Non indiquée',sheet='E-2',
        status='NEUF (*)' if new else 'EXISTANT À RÉUTILISER',
        reserve=reserve,source='Cédule du 22 avril 2024; pôles selon cellules regroupées'))

p1odd={1:(15,'EM. LIGHTING + N.L.'),3:(15,'RMS'),5:(15,'SALES A. TRACK LIGHTING'),7:(15,'SALES A. TRACK LIGHTING'),9:(20,'SALES A. TRACK LIGHTING'),11:(15,'SALES A. TRACK LIGHTING'),13:(15,'SALES A. TRACK LIGHTING'),15:(20,'SALES A. TRACK LIGHTING'),17:(15,'SALES A. TRACK LIGHTING'),19:(15,'SALES A. TRACK LIGHTING'),21:(15,'SALES A. LIGHTING'),23:(20,'SIGNAGE'),25:(20,'SALES A. LIGHTING DISP.'),27:(15,'CORRIDOR LIGHTING'),47:(15,'FITTING RM MIRROR'),49:(15,'BLADE SIGN OUTLET')}
p1even={2:'OFFICE BOH LIGHTING',4:'BOH LIGHTING',6:'FITTING ROOM LIGHTING',10:'LED STRIP SALES AREA',12:'LED STRIP SALES AREA',14:'TYPE N',18:'VITRINE LED LIGHTING'}
for circuit in range(1,50):
    if circuit%2:
        amp,desc=p1odd.get(circuit,(20,'SPARE'))
    else:amp,desc=20,p1even.get(circuit,'SPARE')
    add('P1',str(circuit),amp,desc,circuit in {9,15})

p2odd={1:'CASH OUTLET',3:'CASH OUTLET',5:'TV OUTLET',7:'SALES AREA OUTLETS',9:'SALES AREA OUTLETS',11:'SALES AREA OUTLETS',13:'SALES AREA OUTLETS',15:'SALES AREA OUTLETS',17:'SALES AREA OUTLETS',19:'SALES AREA OUTLETS',21:'FLOOR FIXTURE OUTLET',23:'FLOOR FIXTURE OUTLET',25:'FLOOR FIXTURE OUTLET',27:'FLOOR FIXTURE OUTLET',29:'FLOOR FIXTURE OUTLET',31:'FLOOR FIXTURE OUTLET',33:'FLOOR FIXTURE OUTLET',35:'TV OUTLET',37:'TV OUTLET',39:'INITIATIVE SCREEN',41:'OFFICE DESK OUTLETS',43:'OFFICE OUTLET',45:'ROOM OUTLET',47:'AUTOMATIC DOOR OP.',49:'BOH OUTLET'}
for circuit,desc in p2odd.items():
    amp=20 if circuit in {25,27} else 30 if circuit==39 else 15
    add('P2',str(circuit),amp,desc,circuit==39)
for circuit,amps,description,new,reserve in [
    ('2',15,'BOH OUTLETS',False,''),('4',15,'ANTI-THEFT SYSTEM',False,''),
    ('6','','HANGING TICKER *',False,'* Calibre traversé par le nuage de révision; lecture apparente 40 A non retenue comme ferme.'),
    ('8',15,'HANGING TICKER ENTRANCE',False,''),
    ('10','','FLOOR FIXTURE OUTLET *',False,'* Calibre et libellé traversés par le nuage de révision; calibre non arrêté.'),
    ('12',15,'DOOR OPENER',False,''),
    ('14',20,'DUCT HEATER 01',True,'* Cellule traversée par le nuage : 20* lu; confirmer sur cédule non masquée.'),
    ('16',25,'DUCT HEATER 02',True,''),('18',15,'SPARE',False,''),('20',15,'SPARE',False,''),('22',15,'SPARE',False,''),
    ('24/26/28',15,'DUCT HEATER 03',True,''),('30/32',25,'WALL FAN HEATER "E"',True,''),
    ('34',15,'SPARE',False,''),('36',15,'SPARE',False,''),('38/40',30,'SPARE',False,''),
    ('42/44/46',25,'CEILING FAN HEATER "C"',True,'')]:
    add('P2',circuit,amps,description,new,reserve)
rows.sort(key=lambda r:(r['panel'],int(r['circuits'].split('/')[0])))
assert len(rows)==91
fields=list(rows[0])
for name in ['S-0922-breakers-detailed.csv','S-0922-E-2-breakers.csv']:
    with (ROOT/name).open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(rows)
wb=openpyxl.Workbook();notes=wb.active;notes.title='Lire avant usage'
for line in [
    'S-0922 — Disjoncteurs des cédules E-2 (22 avril 2024)',
    '91 appareils représentés : P1 = 49, P2 = 42. 8 marqués neufs dans la source; 83 existants à réutiliser.',
    'Pôles lus suivant les cellules fusionnées. Un appareil multipolaire compte pour 1.',
    'SPARE avec calibre inclus; SPACE sans calibre exclu (P1 : 50 à 68, P2 : 48 à 66 pairs et 51 à 65 impairs).',
    'P1 et P2 : 225 A; 120/208 V, 3 phases, 4 fils. Marque/modèle/ICC : EXISTING.',
    'Aucun principal ajouté aux départs. Cédules informatives : vérifier charges existantes et mettre à jour sur place.',
    'Attention : * source = neuf. La colonne réserve distingue les doutes du présent relevé.',
    'P2-6 et P2-10 : calibres masqués, laissés vides. P2-14 : lecture 20* à confirmer.',
    'Marque, série, protections spéciales et pouvoir de coupure non choisis.',
    'Le total ne constitue pas une liste d’achat de 91 appareils neufs.',
]:notes.append([line])
notes.column_dimensions['A'].width=145
for name,subset in [('Tous les disjoncteurs',rows),('P1',[r for r in rows if r['panel']=='P1']),('P2',[r for r in rows if r['panel']=='P2'])]:
    ws=wb.create_sheet(name);ws.append(fields)
    for r in subset:ws.append([r[f] for f in fields])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='2C78BD')
    for col,width in {'A':12,'B':16,'C':12,'D':10,'E':10,'F':35,'G':22,'H':12,'I':30,'J':100,'K':70}.items():ws.column_dimensions[col].width=width
summary=wb.create_sheet('Synthèse')
summary.append(['Panneau','Calibre A','Pôles','État source','Réserve','Quantité'])
counts=Counter((r['panel'],str(r['ampere']) or 'NON LISIBLE',r['poles'],r['status'],bool(r['reserve'])) for r in rows)
for key,n in sorted(counts.items()):summary.append(list(key)+[n])
summary.freeze_panes='A2'
for col in 'ABCDEF':summary.column_dimensions[col].width=24
wb.save(ROOT/'S-0922-breakers-by-panel.xlsx')
print(Counter(r['panel'] for r in rows),Counter(r['status'] for r in rows))

families={'PN':'Panneau existant (renvoi E-4)','SEC':'Sectionneur au schéma','TR':'Transformateur existant à suspendre','CT':'Contacteur éclairage','TC':'Minuterie','MS':'Interrupteur maître','COM':'Compteur existant','BJ':'Boîte au schéma'}
markers=[]
def mark(family,x,y,scope,model,reserve=''):
    markers.append(dict(id=f'E-2-{len(markers)+1:03}',family=family,x=x,y=y,scope=scope,model=model,reserve=reserve))
mark('COM',1154.5,1242.4,'CONSERVER','EXISTANT')
mark('BJ',1154.5,1265.7,'CONSERVER','EXISTANT')
for x,model,scope in [(1185.3,'100 A / F100 A','CONSERVER'),(1212.9,'100 A / F80 A','CONSERVER'),(1240.4,'30 A / F30 A','INSTALLER'),(1262.0,'200 A / F150 A','CONSERVER')]:mark('SEC',x,1254,scope,model)
mark('TR',1341.7,1228,'RENVOI_E-4','112,5 kVA; 600/120/208 V; 3 phases','* Transformateur existant à suspendre; support selon détail, aucun appareil neuf ajouté.')
for x in [1322.7,1360.4]:mark('PN',x,1281.5,'RENVOI_E-4','120/208 V; 225 A à la cédule')
mark('TC',1300.2,1274.3,'INSTALLER','INTERMATIC T7401B; 4 contacts 40 A; 120 V')
mark('CT',1300.2,1289.8,'INSTALLER','SQUARE D 8903-LG1200; 12 pôles 30 A')
mark('CT',1322.7,1305.3,'INSTALLER','SQUARE D 8903-LG60; 6 pôles 30 A')
mark('MS',1289.7,1289.8,'INSTALLER','Leviton WTG15/10 COVER ou équivalent selon E-1')
mark('MS',1322.7,1317.3,'INSTALLER','Leviton WTG15/10 COVER ou équivalent selon E-1')
families.update({'REP':'Répartiteur existant','DEM':'Démarreur au détail de principe','RA':'Relais au détail de principe','DG':'Détecteur de gaine au détail'})
mark('REP',1230,1278,'CONSERVER','400 A; 600 V; 3 phases; 4 fils')
mark('REP',1342,1254,'CONSERVER','400 A; 600 V inscrit; 3 phases; 4 fils','* Tension secondaire à clarifier : 600 V inscrit après transformateur 120/208 V.')
for family,x,y in [('DEM',351,1020.5),('RA',351,1048.2),('DG',454.3,1043)]:mark(family,x,y,'DÉTAIL DE PRINCIPE','MODÈLE NON PRÉCISÉ','* Symbole au détail seulement; aucun multiplicateur ni appareil ajouté au plan E-4.')
data=dict(title='Distribution et cédules — 22 avril 2024',legend_box=[1100,300,485,650],families=families,markers=markers,notes=[
    'CÉDULES : P1 = 49; P2 = 42 disjoncteurs. 91 appareils, dont 8 marqués neufs dans la source.',
    'Détail par circuit dans les CSV et le classeur. SPACE sans calibre exclu; SPARE avec calibre inclus.',
    'P2-6 et P2-10 : calibres masqués. P2-14 : 20* lu, à confirmer. Aucun calibre substitué.',
    'Pastilles : symboles du schéma uniquement. Les cédules ne sont pas des symboles d’appareils posés.',
    'Ne pas additionner les panneaux et le transformateur au plan E-4. Détail incendie : 3 symboles de principe, pas des appareils supplémentaires. Câblage non métré.',
    'Réserve : répartiteur secondaire noté 600 V au schéma après transformateur 120/208 V; à clarifier.',
    'RES / * : réserve du relevé; distinct du * = disjoncteur neuf dans la source.'
],verify_zones={'distribution':[1140,1160,1390,1335],'fire-detail':[295,980,515,1100],'P1':[438,20,735,495],'P2':[765,20,1050,495]})
(ROOT/'data/E-2.json').write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
