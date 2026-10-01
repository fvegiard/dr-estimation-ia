"""Export reviewed panel schedules as a detailed CSV and workbook."""
from pathlib import Path
import json,csv
from collections import Counter
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill
root=Path(__file__).resolve().parent
headers=['Panneau','Circuits','Ampères','Pôles','Quantité','Portée','Désignation','Feuille','Modèle','Réserve','X source px','Y source px']
rows=[]
for sheet in ['EX-M-PE01','EX-E-PE01']:
    data=json.loads((root/'evidence'/f'{sheet}.json').read_text())
    from PIL import Image
    scale=Image.open(root/data['source']).width/2000
    for m in data['markers']:
        rows.append([m['parent'],','.join(map(str,m['circuits'])),m['amps'],m['poles'],1,m['scope'],m['description'],sheet,'MODÈLE NON PRÉCISÉ',m.get('reserve',''),round(m['x']*scale,1),round(m['y']*scale,1)])
with (root/'S-0844-breakers-detailed.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.writer(f);w.writerow(headers);w.writerows(rows)
wb=Workbook();ws=wb.active;ws.title='Tous les disjoncteurs';ws.append(headers)
for row in rows:ws.append(row)
for panel in sorted({r[0] for r in rows}):
    ws=wb.create_sheet(panel);ws.append(headers)
    for row in rows:
        if row[0]==panel:ws.append(row)
ws=wb.create_sheet('Synthèse');ws.append(['Panneau','Feuille','Ampères','Pôles','Quantité'])
for key,count in sorted(Counter((r[0],r[7],r[2],r[3]) for r in rows).items()):ws.append([*key,count])
ws=wb.create_sheet('Réserves');ws.append(['Sujet','Réserve'])
for row in [
    ['Portée','235 appareils représentés aux panneaux nouveaux; 103 existants séparés. Aucun total d’achat global.'],
    ['Multipolaires','Un appareil par groupe de pôles liés. Espaces sans calibre exclus.'],
    ['C1-3','130 A / 3P aux circuits 2,4,6 : valeur source conservée, série à confirmer.'],
    ['Barres omnibus','0 A aux tableaux existants UE-1, C1-1, C3-1 et C4-1 : renseignement inutilisable, aucun remplacement arbitraire.'],
    ['Sélection','Marque, série, protections, pouvoir de coupure et principaux non précisés.'],
    ['Versions','EX-M-PE01 : 24 avril 2024 R06; EX-E-PE01 : 31 octobre 2023 R0. Versions antérieures non additionnées.']]:ws.append(row)
for ws in wb:
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='285577')
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width=min(65,max(12,max(len(str(c.value or '')) for c in column)+2))
dest=root/'S-0844-breakers-by-panel.xlsx';wb.save(dest)
check=load_workbook(dest,read_only=True,data_only=True)
assert check['Tous les disjoncteurs'].max_row==339
assert sum(r[4] for r in list(check['Tous les disjoncteurs'].values)[1:])==338
print('338 disjoncteurs; classeur rouvert et totaux contrôlés')
